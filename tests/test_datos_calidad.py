"""Controles de calidad previos a la optimización: la base válida pasa; cada defecto se detecta y bloquea."""
import datetime as dt
import subprocess
import sys

import openpyxl
import pytest

from conftest import BASE, CODIGO
from sagapi.datos import ErrorDatos, cargar_base


def test_base_final_sin_errores():
    d, inf = cargar_base(BASE)
    assert inf.ok and not inf.errores
    assert len(d.pedidos) == 96 and len(d.operarios) == 25
    assert d.procedencia["habilidades"] == "SIMULADO" and d.procedencia["pedidos"] == "SIMULADO"
    assert int(d.parametros["trabajadores_produccion"]) == 25 and int(d.parametros["turnos_por_dia"]) == 1
    assert d.calendario.h_dia == 8.0


def _copia_modificada(tmp_path, cambios):
    wb = openpyxl.load_workbook(BASE)
    for hoja, celda, valor in cambios:
        wb[hoja][celda] = valor
    ruta = tmp_path / "base_modificada.xlsx"
    wb.save(ruta)
    return ruta


def _col(hoja, nombre):
    wb = openpyxl.load_workbook(BASE, read_only=True)
    enc = [c.value for c in next(wb[hoja].iter_rows(max_row=1))]
    return openpyxl.utils.get_column_letter(enc.index(nombre) + 1)


CASOS = [
    ("Q01", "06_RUTAS", "min_por_unidad", "=1+1"),                 # fórmula en columna de entrada
    ("Q03", "20_PEDIDOS_SIM", "unidades", None),                   # obligatorio vacío
    ("Q04", "06_RUTAS", "min_fijos_lote", "no aplica"),            # texto de relleno en campo numérico
    ("Q05", "06_RUTAS", "trabajadores_requeridos", -1),            # fuera de rango
    ("Q07", "06_RUTAS", "recurso_id", "MAQ-INEXISTENTE"),          # integridad referencial
    ("Q10", "20_PEDIDOS_SIM", "fecha_compromiso", dt.datetime(2024, 12, 1)),   # compromiso antes de recepción
]


@pytest.mark.parametrize("codigo,hoja,columna,valor", CASOS)
def test_defecto_detectado_y_bloquea(tmp_path, codigo, hoja, columna, valor):
    ruta = _copia_modificada(tmp_path, [(hoja, f"{_col(hoja, columna)}2", valor)])
    d, inf = cargar_base(ruta, estricto=False)
    assert d is None, "con errores no se construye el modelo de datos"
    assert any(p.codigo == codigo and p.hoja == hoja for p in inf.errores), inf.texto(20)
    with pytest.raises(ErrorDatos):
        cargar_base(ruta)


def test_clave_duplicada(tmp_path):
    wb = openpyxl.load_workbook(BASE)
    ws = wb["20_PEDIDOS_SIM"]
    ws["A3"] = ws["A2"].value
    ruta = tmp_path / "dup.xlsx"
    wb.save(ruta)
    d, inf = cargar_base(ruta, estricto=False)
    assert any(p.codigo == "Q06" for p in inf.errores)


def test_cli_validar_codigos_de_salida(tmp_path):
    ok = subprocess.run([sys.executable, str(CODIGO / "Programa_principal.py"), "validar"], capture_output=True, text=True)
    assert ok.returncode == 0, ok.stdout + ok.stderr
    ruta = _copia_modificada(tmp_path, [("06_RUTAS", f"{_col('06_RUTAS', 'recurso_id')}2", "MAQ-X")])
    mal = subprocess.run([sys.executable, str(CODIGO / "Programa_principal.py"), "validar", "--base", str(ruta)],
                         capture_output=True, text=True)
    assert mal.returncode == 2
    assert "Q07" in mal.stdout


def test_cli_optimizar_con_base_invalida_se_detiene(tmp_path):
    ruta = _copia_modificada(tmp_path, [("06_RUTAS", f"{_col('06_RUTAS', 'min_por_unidad')}2", "=2*3")])
    r = subprocess.run([sys.executable, str(CODIGO / "Programa_principal.py"), "programar", "--base", str(ruta),
                        "--plan", "FIFO", "--reaccion", "CONV", "--escenario", "S00", "--sin-personal",
                        "--salida", str(tmp_path / "out")], capture_output=True, text=True)
    assert r.returncode == 2
    assert "no se optimiza" in r.stdout


def test_matriz_de_habilidades_real_reemplaza_a_la_simulada(tmp_path):
    """La matriz simulada se reemplaza llenando 16_REG_HABILIDADES y fijando fuente_habilidades = REAL.

    (Los registros de esta prueba son de PRUEBA: reproducen la matriz simulada para comprobar el mecanismo.)
    """
    d_sim, _ = cargar_base(BASE)
    wb = openpyxl.load_workbook(BASE)
    ws = wb["16_REG_HABILIDADES"]
    fila = 2
    for w in d_sim.operarios:
        for m in sorted(d_sim.habilidades[w]):
            ws.cell(fila, 1, w); ws.cell(fila, 2, m); ws.cell(fila, 3, 2); ws.cell(fila, 4, 1)
            ws.cell(fila, 5, "PRUEBA AUTOMÁTICA"); ws.cell(fila, 6, dt.datetime(2026, 1, 1))
            fila += 1
    wb["16_REG_HABILIDADES"].cell(fila, 1, d_sim.operarios[0]); wb["16_REG_HABILIDADES"].cell(fila, 2, "PLE-01")
    ws.cell(fila, 3, 1); ws.cell(fila, 4, 0)                         # nivel 1 / no opera solo: NO habilita
    ws.cell(fila, 5, "PRUEBA AUTOMÁTICA"); ws.cell(fila, 6, dt.datetime(2026, 1, 1))
    hp = wb["11_PARAMETROS"]
    for r in hp.iter_rows(min_row=2):
        if r[0].value == "fuente_habilidades":
            r[3].value = "REAL"
    ruta = tmp_path / "base_hab_real.xlsx"
    wb.save(ruta)
    d, inf = cargar_base(ruta)
    assert inf.ok
    assert d.procedencia["habilidades"] == "REAL"
    assert d.habilidades == {w: set(s) for w, s in d_sim.habilidades.items()}
