#!/usr/bin/env python3
"""Construye la base de datos SAGAPI v2.0 a partir del Excel v1.2 SIN modificar el original.

Pasos (``--paso``):
  migrar  original v1.2  ->  SAGAPI_BD_OPERATIVA_v2_0.xlsx  +  SAGAPI_BD_SIMULACION_v2_0.xlsx
  modelo  OPERATIVA + SIMULACION  ->  SAGAPI_BD_MODELO_v2_0.xlsx (esquema heredado que lee el
          programa Python 1.4.0 sin cambios de código; libro GENERADO, no editar a mano)
  todo    ambos pasos

Uso (desde la raíz del proyecto):
  python herramientas/construir_bd_v2.py --paso todo \
      --original Base_Datos_SAGAPI_25_TRABAJADORES_SIMULADOS_v1_2.xlsx --salida entrega_v2_0/base_datos

Reglas aplicadas: no se inventan datos; los textos de relleno en campos numéricos se convierten en
vacío; las fórmulas implícitas se hacen explícitas y se comprueba que reproducen los valores
originales; lo simulado se separa de lo operativo; cada fila declara su naturaleza.
"""
from __future__ import annotations

import argparse
import datetime as dt
import math
import os
import sys
from collections import OrderedDict, defaultdict

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bd_v2_comun import (Col, LISTA_NATURALEZA, VERSION_BD, a_fecha, entero, escribir_tabla,  # noqa: E402
                         hoja_diccionario, nombre_definido, num_o_none, validacion_decimal,
                         validacion_entero, validacion_lista)

ARCH_OPER = "SAGAPI_BD_OPERATIVA_v2_0.xlsx"
ARCH_SIM = "SAGAPI_BD_SIMULACION_v2_0.xlsx"
ARCH_MOD = "SAGAPI_BD_MODELO_v2_0.xlsx"

FAMILIA_ID = OrderedDict([("Racks industriales", "FAM-RAC"), ("Góndolas comerciales", "FAM-GON"),
                          ("Lockers metálicos", "FAM-LOC"), ("Otros productos metalmecánicos", "FAM-OTR")])
GRUPO = {"CHAPA": "CHAPA", "PERFIL": "PERFIL", "Perfil": "PERFIL", "HERRAJE": "HERRAJE", "Herraje": "HERRAJE",
         "OTRO": "OTRO", "Pintura": "PINTURA", "Preparación superficial": "PREPARACION_SUPERFICIAL",
         "Consumible": "CONSUMIBLE", "Soldadura": "SOLDADURA", "Fijación": "FIJACION", "Accesorio": "ACCESORIO",
         "Embalaje": "EMBALAJE", "Seguridad": "SEGURIDAD"}
UNIDAD = {"kg": "kg", "m": "m", "metro": "m", "unidad": "und", "litro": "l", "m³": "m3", "m²": "m2", "par": "par"}
GENERICOS = {"MAT-PR-1", "MAT-FIJ", "MAT-OTR"}
TIPO_RECURSO = {"CCH-01": "MAQUINA", "CPR-01": "MAQUINA", "PER-01": "MAQUINA", "PER-02": "MAQUINA",
                "PLE-01": "MAQUINA", "SOL-01": "MAQUINA", "SOL-02": "MAQUINA", "PIN-01": "EQUIPO_MANUAL",
                "HOR-01": "MAQUINA", "DES-01": "PUESTO_MANUAL", "ARM-01": "PUESTO_MANUAL", "MAN-01": "PUESTO_MANUAL",
                "PRE-01": "AREA", "CAL-01": "AREA", "EMP-01": "AREA", "SAL-PIN": "AREA", "PIN-EXT": "EXTERNO"}
TIPO_ASIG = {"PRINCIPAL": "PRINCIPAL", "ALTERNATIVA CON UTILLAJE COMPATIBLE": "ALT_UTILLAJE",
             "ALTERNATIVA MISMO PROCESO": "ALT_MISMO_PROCESO"}
TIPO_ASIG_INV = {v: k for k, v in TIPO_ASIG.items()}
ORIGEN_T = {"VALOR HEREDADO SIN CRONOMETRAJE": "HEREDADO_SIN_CRONOMETRAJE", "HIPOTESIS TECNICA": "HIPOTESIS_TECNICA"}
ORIGEN_T_INV = {v: k for k, v in ORIGEN_T.items()}
POLITICA = {"A": "A - reanudar", "B": "B - reiniciar lote"}
GEN_REC = {"GENERICO_POR_UNIDAD": ("(Genérico) operaciones POR_UNIDAD", "operaciones POR_UNIDAD"),
           "GENERICO_POR_LOTE": ("(Genérico) Curado térmico POR_LOTE", "Curado térmico POR_LOTE")}

# Parámetros: (id, nombre heredado leído por Python o descriptivo, valor_num, valor_txt, unidad, naturaleza,
#              nota, usado_por_python, libro)
PARAMETROS = [
    ("trabajadores_produccion", "Número total de trabajadores de producción", 25, None, "personas", "REAL",
     "Confirmado por la empresa (06C v1.2).", 0, "OPER"),
    ("turnos_por_dia", "Turnos diarios", 1, None, "turnos", "REAL", "Confirmado por la empresa.", 0, "OPER"),
    ("horas_productivas_dia", "Horas productivas por día", 8, None, "h/día", "REAL",
     "Jornada neta confirmada; el modelo la deriva de 'Jornada productiva'.", 0, "OPER"),
    ("jornada_productiva", "Jornada productiva", None, "08:00-12:00 y 13:00-17:00", "hora", "ESTIMADO",
     "Horario no confirmado por la empresa; define el eje de horas productivas.", 1, "OPER"),
    ("refrigerio", "Refrigerio", None, "12:00-13:00", "hora", "ESTIMADO", "Horario no confirmado.", 0, "OPER"),
    ("aplicar_feriados", "Aplicar feriados nacionales (1/0)", 0, None, "1/0", "CONFIGURACION",
     "0 = calendario v1.2 (feriados laborables). Práctica real de SAGAPI pendiente de confirmación.", 1, "OPER"),
    ("anio_base", "Año base del estudio", 2025, None, "año", "CONFIGURACION", "", 0, "OPER"),
    ("hora_liberacion", "Hora de liberación", None, "08:00", "hora", "ESTIMADO",
     "Fijada en el código (instancia.py); el valor aquí es informativo.", 0, "OPER"),
    ("hora_vencimiento", "Hora de vencimiento d_j", None, "17:00", "hora", "ESTIMADO",
     "Cierre del día comprometido o del último día productivo previo (calendario.py).", 0, "OPER"),
    ("densidad_acero", "Densidad del acero", 7850, None, "kg/m³", "REFERENCIA",
     "Constante física; reproduce exactamente la masa de chapa de v1.2.", 0, "OPER"),
    ("kg_por_m_perfil", "Masa lineal del perfil genérico", 1.5, None, "kg/m", "ESTIMADO",
     "Implícito en v1.2 (masa = piezas × longitud × 1,5); sección del perfil por definir.", 0, "OPER"),
    ("aprovechamiento_chapa", "Factor de aprovechamiento de chapa", 0.92, None, "fracción", "ESTIMADO",
     "Implícito en v1.2 (costo = masa × precio / 0,92); sin medición de merma.", 0, "OPER"),
    ("aprovechamiento_perfil", "Factor de aprovechamiento de perfil", 0.96, None, "fracción", "ESTIMADO",
     "Implícito en v1.2 (costo = masa × precio / 0,96).", 0, "OPER"),
    ("area_plancha_m2", "Área de plancha estándar 1,20 × 2,40 m", 2.88, None, "m²", "ESTIMADO",
     "Formato del material MAT-LAC2; no confirma que todas las chapas se compren en ese formato.", 0, "OPER"),
    ("costo_hora_hombre", "Costo hora-hombre", 13, None, "S/ por h", "ESTIMADO", "Supuesto económico sin planilla.",
     0, "OPER"),
    ("tarifa_electrica", "Tarifa eléctrica", 0.65, None, "S/ por kWh", "ESTIMADO", "Verificar con recibo eléctrico.",
     0, "OPER"),
    ("factor_emision_red", "Factor de emisión de la red", None, None, "kg CO2/kWh", "PENDIENTE",
     "Sin dato oficial cargado; el programa no calcula CO2.", 0, "OPER"),
    ("ventas_anuales", "Ventas anuales aproximadas", 2000000, None, "S/", "REAL", "Declarado por la empresa.", 0,
     "OPER"),
    ("uit_2025", "UIT 2025", 5350, None, "S/", "REFERENCIA", "Valor oficial 2025.", 0, "OPER"),
    ("ciclo_horno_min", "Ciclo horno", 40, None, "min por carga", "ESTIMADO",
     "Incluido en el tiempo fijo (45 min) del curado en 06_RUTAS.", 0, "OPER"),
    ("capacidad_horno_carga", "Capacidad horno por lote", 10, None, "piezas por carga (declarado)", "ESTIMADO",
     "Conflicto de unidades: el lote del modelo es de 10 UNIDADES de producto, no 10 piezas.", 0, "OPER"),
    # ---------------- configuración del modelo y de los experimentos (libro de simulación)
    ("maquina_averia", "Máquina de la avería", None, "Plegadora hidráulica de chapa metálica", "-", "CONFIGURACION",
     "Recurso crítico de los experimentos de Monte Carlo.", 1, "SIM"),
    ("beta_reasignacion", "Costo equivalente reasignación", 8, None, "h equivalentes", "CONFIGURACION",
     "Peso β de f2 (inestabilidad).", 1, "SIM"),
    ("replicas_mc", "Réplicas Monte Carlo de fallas de la plegadora", 40, None, "réplicas", "CONFIGURACION", "", 1,
     "SIM"),
    ("semilla_mc", "Semilla Monte Carlo", 2025200, None, "-", "CONFIGURACION", "", 1, "SIM"),
    ("regla_penalidad", "Regla de penalidad del modelo", None, "FIJA", "-", "CONFIGURACION",
     "FIJA = penalidad fija de 01_FAMILIAS si hay atraso; DIARIA_CON_TOPE solo en sensibilidad.", 0, "SIM"),
    ("penalidad_diaria_formula", "Penalidad diaria (sensibilidad)", None,
     "0,10·valor/(F·plazo); F=0,40 si plazo ≤ 60 d; 0,25 si > 60 d", "S/ por día", "ESTIMADO",
     "Referencia de contratación pública peruana; no es cláusula de SAGAPI. Fórmula fija en datos.py.", 0, "SIM"),
    ("tolerancia_lexicografica", "Tolerancia lexicográfica τ", 0.000001, None, "proporción", "CONFIGURACION",
     "Valor fijo en el código (optimizacion.py).", 0, "SIM"),
    ("tolerancia_verificacion", "Tolerancia de verificación", 0.000001, None, "h", "CONFIGURACION",
     "Valor fijo en el código (verificacion.py).", 0, "SIM"),
    ("rho_epsilon", "Coeficiente de aumentación ε (ρ)", 0.0001, None, "-", "CONFIGURACION", "", 0, "SIM"),
    ("limite_computo_s", "Límite de cómputo por fase", 60, None, "s", "CONFIGURACION",
     "Se fija con --tl en la línea de comandos.", 0, "SIM"),
    ("brecha_relativa", "Brecha relativa máxima", 0.00001, None, "proporción", "CONFIGURACION", "", 0, "SIM"),
    ("niveles_epsilon", "Niveles para epsilon", 7, None, "niveles", "CONFIGURACION", "", 0, "SIM"),
    ("solucionador", "Solucionador", None,
     "HiGHS vía scipy.optimize.milp (API pública); una resolución por fase lexicográfica; sin votación",
     "-", "CONFIGURACION", "Corrige el texto de v1.2 (cartera de 3 configuraciones), obsoleto en el código 1.4.0.",
     0, "SIM"),
]


# =============================================================================== lectura del original
def tabla(ws):
    filas = list(ws.iter_rows(values_only=True))
    enc = [str(c).strip() if c is not None else f"_c{j}" for j, c in enumerate(filas[0])]
    out = []
    for f in filas[1:]:
        if f is None or all(v is None for v in f):
            continue
        out.append({enc[j]: f[j] for j in range(min(len(enc), len(f)))})
    return out


def leer_original(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    t = {s: tabla(wb[s]) for s in wb.sheetnames if wb[s].max_row > 1}
    hab_ws = wb["06B_HABILIDADES"]
    filas = list(hab_ws.iter_rows(values_only=True))
    t["_06B_encabezado"] = filas[0]
    t["_06B_filas"] = filas[1:]
    return t


# =============================================================================== modelo de datos normalizado
def normalizar(o):
    """Devuelve un dict con las tablas normalizadas (listas de dict) de ambos libros fuente."""
    N = {}
    # ---------------- familias
    N["familias"] = [dict(familia_id=FAMILIA_ID[f["Familia"]], familia=f["Familia"],
                          unidades_2025=entero(f["Unidades agregadas 2025"]), ventas_2025=f["Ventas agregadas S/"],
                          penalidad_fija=f["Penalidad fija %"], penalidad_diaria=f["Penalidad diaria %"],
                          naturaleza="REAL") for f in o["12_PRECIOS"]]
    # ---------------- productos
    prod_id = OrderedDict()
    N["productos"] = []
    for k, f in enumerate(o["01_PRODUCTOS"], 1):
        pid = f"PRD-{k:03d}"
        prod_id[f["Nombre completo"]] = pid
        N["productos"].append(dict(producto_id=pid, familia_id=FAMILIA_ID[f["Familia"]], producto=f["Nombre completo"],
                                   tipo=f["Tipo"], alto_cm=num_o_none(f["Alto cm"]), ancho_cm=num_o_none(f["Ancho cm"]),
                                   fondo_cm=num_o_none(f["Fondo cm"]), puertas=num_o_none(f["Puertas"]),
                                   lote_guia=entero(f["Lote guía"]), naturaleza="ESTIMADO"))
    N["_prod_id"] = prod_id
    # ---------------- materiales
    mat_id = OrderedDict()
    N["materiales"] = []
    for f in o["02_MATERIALES"]:
        mid = f["Clave técnica auxiliar"]
        mat_id[f["Material o accesorio (nombre completo)"]] = mid
        precio = num_o_none(f["Costo unitario S/"])
        N["materiales"].append(dict(material_id=mid, material=f["Material o accesorio (nombre completo)"],
                                    grupo=GRUPO[f["Grupo"]], unidad_costo=UNIDAD[f["Unidad de costo"]],
                                    espesor_mm=num_o_none(f["Espesor mm"]), ancho_m=num_o_none(f["Ancho m"]),
                                    largo_m=num_o_none(f["Largo m"]), costo_unitario=precio,
                                    generico=1 if mid in GENERICOS else 0,
                                    naturaleza="ESTIMADO" if precio is not None else "PENDIENTE"))
    N["_mat_id"] = mat_id
    mat = {m["material_id"]: m for m in N["materiales"]}
    # ---------------- despiece
    N["despiece"] = []
    for k, f in enumerate(o["03_DESPIECE"], 1):
        mid = mat_id[f["Material"]]
        g = mat[mid]["grupo"]
        if g == "CHAPA":
            tipo = "CHAPA_AREA"
        elif g == "PERFIL":
            tipo = "PERFIL_LONGITUD"
        elif mat[mid]["costo_unitario"] is not None:
            tipo = "UNIDAD"
        else:
            tipo = "SIN_PRECIO"
        ancho, alto = num_o_none(f["Ancho pieza mm"]), num_o_none(f["Alto pieza mm"])
        lon = num_o_none(f["Longitud pieza m"])
        area = num_o_none(f["Área de chapa m²"])
        N["despiece"].append(dict(
            bom_id=f"BOM-{k:04d}", producto_id=prod_id[f["Producto"]], componente=f["Componente"], material_id=mid,
            tipo_calculo=tipo, piezas=entero(f["Piezas por unidad"]),
            ancho_mm=ancho, alto_mm=alto, longitud_m=(lon if lon not in (0, 0.0) else None),
            area_m2=(area if (area or 0) > 0 else None), area_formula=(ancho is not None and alto is not None),
            masa_v12=num_o_none(f["Masa kg / producto"]), costo_v12=num_o_none(f["Costo S/ / producto"]),
            planchas_v12=num_o_none(f["Planchas 1.20×2.40 equivalentes / producto"]), naturaleza="ESTIMADO"))
    # ---------------- recursos
    rec_id = OrderedDict()
    N["recursos"] = []
    for f in o["04_MAQUINAS"]:
        rid = f["Clave interna"]
        rec_id[f["Equipo o recurso completo"]] = rid
        N["recursos"].append(dict(recurso_id=rid, recurso=f["Equipo o recurso completo"], tipo_tecnico=f["Tipo técnico"],
                                  tipo_recurso=TIPO_RECURSO[rid], cantidad=entero(f["Cantidad de escenario"]),
                                  costo_uso=num_o_none(f["Costo de uso S//h"]),
                                  potencia_kw=num_o_none(f["Potencia nominal kW (referencia)"]),
                                  factor_carga=num_o_none(f["Factor de carga medio"]),
                                  grupo_operarios=f["Grupo de operarios"], critico=entero(f["Máquina crítica (1/0)"]),
                                  inventario_verificado=0, naturaleza="ESTIMADO",
                                  marca_ref=f["Marca comercial del escenario"], modelo_ref=f["Modelo comercial"],
                                  estatus_ref=f["Estatus físico"]))
    N["_rec_id"] = rec_id
    # ---------------- rutas
    N["rutas"] = []
    for f in o["05_RUTAS"]:
        pr = f["Precedencia"]
        N["rutas"].append(dict(producto_id=prod_id[f["Producto"]], paso=entero(f["Paso"]), operacion=f["Operación"],
                               recurso_id=rec_id[f["Equipo principal"]],
                               paso_predecesor=(None if str(pr).strip().lower() == "inicio" else entero(pr)),
                               min_fijos=round(num_o_none(f["Min fijos por lote"]), 10), min_unidad=round(num_o_none(f["Minutos por unidad"]), 10),
                               tratamiento=f["Tratamiento del tiempo"], min_ciclo_horno=num_o_none(f["Min ciclo horno (incluidos)"]),
                               origen_tiempo=ORIGEN_T[f["Origen de duración"]], naturaleza="ESTIMADO"))
    # ---------------- elegibilidad
    N["elegibilidad"] = []
    for f in o["13_ELEGIBILIDAD"]:
        tipo = TIPO_ASIG[f["Tipo de asignación"]]
        N["elegibilidad"].append(dict(producto_id=prod_id[f["Producto"]], paso=entero(f["Paso"]),
                                      recurso_id=rec_id[f["Máquina elegible"]], tipo_asignacion=tipo,
                                      elegible=entero(f["Elegible 1/0"]), factor_tiempo=num_o_none(f["Factor de tiempo relativo"]),
                                      naturaleza="CALCULADO" if tipo == "PRINCIPAL" else "ESTIMADO"))
    # ---------------- personal
    N["personal"] = [dict(operario_id=f["Operario (seudónimo)"], turno=entero(f["Turnos"]), hora_entrada=dt.time(8, 0),
                          hora_salida=dt.time(17, 0), descanso_h=num_o_none(f["Descanso h"]),
                          horas_netas=num_o_none(f["Horas netas / día"]), naturaleza="REAL",
                          especialidad=f["Especialidad"]) for f in o["06_OPERARIOS"]]
    for p, f in zip(N["personal"], o["06_OPERARIOS"]):
        assert f["Inicio"] == "08:00" and f["Fin"] == "17:00", "horario distinto del esperado"
    # ---------------- calendario
    N["calendario"] = [dict(fecha=a_fecha(f["Fecha"]), laborable_base=entero(f["Productivo (1/0)"]),
                            feriado=(f["Feriado nacional Perú"] or None), naturaleza="ESTIMADO") for f in o["07_CALENDARIO"]]
    # ---------------- parámetros
    N["parametros"] = [dict(zip(("parametro_id", "parametro", "valor_num", "valor_txt", "unidad", "naturaleza", "nota",
                                 "usado_python", "libro"), p)) for p in PARAMETROS]
    # ---------------- SIMULACIÓN: pedidos
    N["pedidos"] = []
    for f in o["08_PEDIDOS"]:
        assert FAMILIA_ID[f["Familia"]] == next(p["familia_id"] for p in N["productos"] if p["producto"] == f["Producto"])
        N["pedidos"].append(dict(pedido_id=f["Pedido"], producto_id=prod_id[f["Producto"]], unidades=entero(f["Unidades"]),
                                 lote=entero(f["Lote"]), recepcion=a_fecha(f["Recepción 2025"]),
                                 liberacion=a_fecha(f["Liberación 2025"]), compromiso=a_fecha(f["Compromiso 2025"]),
                                 naturaleza="SIMULADO"))
    # ---------------- habilidades (matriz simulada)
    enc = list(o["_06B_encabezado"])
    maqs = [c for c in enc[1:] if c and c not in ("Máquinas habilitadas", "Naturaleza")]
    N["hab_recursos"] = [rec_id[m] for m in maqs]
    N["habilidades"] = []
    for fila in o["_06B_filas"]:
        if not fila or not fila[0] or str(fila[0]).startswith(("Operarios habilitados", "=")):
            continue
        N["habilidades"].append(dict(operario_id=fila[0], valores=[entero(fila[j + 1]) for j in range(len(maqs))],
                                     naturaleza="SIMULADO"))
    # ---------------- eventos de indisponibilidad
    N["eventos"] = []
    for f in o["14_FALLAS_ESC"]:
        mant = str(f["Tipo de dato"]).upper().startswith("SUPUESTO")
        N["eventos"].append(dict(evento_id=f["Evento"], tipo_evento="MANTENIMIENTO_PLANIFICADO" if mant else "FALLA",
                                 recurso_id=rec_id[f["Equipo"]], inicio=a_fecha(f["Inicio"]),
                                 horas_reloj=num_o_none(f["Horas de reloj"]),
                                 horas_turno=num_o_none(f["Horas de turno afectadas"]), motivo=f["Motivo hipotético"],
                                 pedido_ref=f["Pedido potencialmente afectado"], paso_ref=entero(f["Paso"]),
                                 naturaleza="SIMULADO"))
    # ---------------- escenarios
    N["escenarios"] = []
    for f in o["18_ESCENARIOS"]:
        rec = str(f["Recuperación"] or "")
        pol = rec[:1] if rec[:1] in ("A", "B") else None
        N["escenarios"].append(dict(escenario_id=f["Escenario"], caso=f["Caso"],
                                    eventos=None if f["Evento"] in (None, "Ninguna") else f["Evento"],
                                    politica=pol, semilla=entero(f["Semilla"]),
                                    naturaleza="SIMULADO" if f["Escenario"] != "S00" else "CONFIGURACION"))
    # ---------------- recuperación
    N["recuperacion"] = []
    for f in o["15_RECUPERACION"]:
        p = str(f["Producto"])
        if p.startswith("(Genérico)"):
            ambito = "GENERICO_POR_LOTE" if "LOTE" in str(f["Operación"]).upper() else "GENERICO_POR_UNIDAD"
            pid, paso = None, None
        else:
            ambito, pid, paso = "PRODUCTO_PASO", prod_id[p], entero(f["Paso"])
        rep = f["Piezas a repetir por interrupción"]
        N["recuperacion"].append(dict(ambito=ambito, producto_id=pid, paso=paso, politica=str(f["Política"])[:1],
                                      conserva=entero(f["Conserva piezas"]), reinicia=entero(f["Reinicia lote"]),
                                      repite_setup=entero(f["Repite setup"]),
                                      piezas_repetir=(None if isinstance(rep, str) else entero(rep)),
                                      min_verificacion=num_o_none(f["Min de verificación tras reparación"]),
                                      operacion_v12=f["Operación"], naturaleza="ESTIMADO"))
    # ---------------- sensibilidad
    N["sensibilidad"] = [dict(factor=f["Factor"], niveles=f["Niveles"], aplicado_a=f["Aplicado a"], proposito=f["Propósito"],
                              naturaleza="CONFIGURACION") for f in o["27_SENSIBILIDAD"]]
    # ---------------- programa FIFO de referencia (calculado por regla sobre pedidos simulados)
    N["fifo"] = [dict(pedido_id=f["Pedido"], paso=entero(f["Paso"]), recurso_id=rec_id[f["Máquina"]],
                      inicio=a_fecha(f["Inicio"]), fin=a_fecha(f["Fin"]), duracion_h=f["Duración h"],
                      naturaleza="CALCULADO") for f in o["16_PROGRAMA_BASE"]]
    return N


# =============================================================================== cálculos espejo
def calc_despiece(N, par):
    """Masa, costo, área y planchas con las mismas fórmulas que se escriben en Excel."""
    mat = {m["material_id"]: m for m in N["materiales"]}
    for b in N["despiece"]:
        m = mat[b["material_id"]]
        area = b["ancho_mm"] * b["alto_mm"] / 1e6 if b["area_formula"] else b["area_m2"]
        b["area_calc"] = area
        if b["tipo_calculo"] == "CHAPA_AREA":
            masa = b["piezas"] * area * m["espesor_mm"] / 1000 * par["densidad_acero"]
            costo = masa * m["costo_unitario"] / par["aprovechamiento_chapa"]
        elif b["tipo_calculo"] == "PERFIL_LONGITUD":
            masa = b["piezas"] * b["longitud_m"] * par["kg_por_m_perfil"]
            costo = masa * m["costo_unitario"] / par["aprovechamiento_perfil"]
        elif b["tipo_calculo"] == "UNIDAD":
            masa, costo = None, b["piezas"] * m["costo_unitario"]
        else:
            masa, costo = None, None
        b["masa_calc"], b["costo_calc"] = masa, costo
        b["planchas_calc"] = b["piezas"] * area / par["area_plancha_m2"] if area else None


def comprobar_reproduccion(N):
    """Las fórmulas explícitas deben reproducir los valores v1.2 (salvo las correcciones documentadas)."""
    errores = []
    for b in N["despiece"]:
        for clave, v12 in (("masa_calc", b["masa_v12"]), ("costo_calc", b["costo_v12"])):
            if v12 is None:
                continue
            if b[clave] is None or abs(b[clave] - v12) > 1e-6 * max(1, abs(v12)):
                errores.append((b["bom_id"], clave, b[clave], v12))
        if b["area_formula"] and abs(b["area_calc"] - (b["area_m2"] or 0)) > 1e-9:
            errores.append((b["bom_id"], "area", b["area_calc"], b["area_m2"]))
    if errores:
        raise SystemExit(f"Las fórmulas explícitas no reproducen v1.2: {errores[:5]}")


# =============================================================================== libro OPERATIVO
def construir_operativa(N, ruta):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    espec = []
    nprod, nmat, nbom = len(N["productos"]), len(N["materiales"]), len(N["despiece"])
    nrec, nrut, nele = len(N["recursos"]), len(N["rutas"]), len(N["elegibilidad"])
    ncal = len(N["calendario"])
    R = lambda hoja, col, n: f"'{hoja}'!${col}$2:${col}${n + 1}"  # noqa: E731

    # ---------------- 11_PARAMETROS (primero: otras hojas usan sus nombres definidos)
    pars = [p for p in N["parametros"] if p["libro"] == "OPER"]
    ws = wb.create_sheet("11_PARAMETROS")
    cols = [Col("parametro_id", "txt", 26, "Clave única del parámetro (nombre definido p_<id>)."),
            Col("parametro", "txt", 44, "Nombre descriptivo; coincide con el nombre que lee Python cuando aplica."),
            Col("valor_num", "gen", 14, "Valor numérico (vacío si el parámetro es textual o pendiente).", oblig=0),
            Col("valor_txt", "txt", 40, "Valor textual (horarios, reglas).", oblig=0),
            Col("unidad", "txt", 18, "Unidad de medida."), Col("naturaleza", "txt", 15, "Código de 00_CODIGOS."),
            Col("nota", "txt", 70, "Límite o fuente.", oblig=0),
            Col("usado_por_python", "bin", 12, "1 si Optimizacion_MILP_SAGAPI.py lo lee.")]
    filas = []
    for p in pars:
        vn = p["valor_num"]
        if p["parametro_id"] == "categoria_empresarial":
            vn = None
        filas.append([p["parametro_id"], p["parametro"], vn, p["valor_txt"], p["unidad"], p["naturaleza"], p["nota"],
                      p["usado_python"]])
    fila_ventas = [p["parametro_id"] for p in pars].index("ventas_anuales") + 2
    fila_uit = [p["parametro_id"] for p in pars].index("uit_2025") + 2
    filas.append(["categoria_empresarial", "Categoría empresarial (Ley 30056)", None,
                  f'=IF(C{fila_ventas}/C{fila_uit}<=150,"Microempresa",IF(C{fila_ventas}/C{fila_uit}<=1700,'
                  f'"Pequeña empresa",IF(C{fila_ventas}/C{fila_uit}<=2300,"Mediana empresa","Gran empresa")))',
                  "-", "CALCULADO", "Ventas anuales / UIT según Ley 30056.", 0])
    escribir_tabla(ws, cols, filas, tab="oper")
    for i, p in enumerate(pars, 2):
        if p["valor_num"] is not None:
            nombre_definido(wb, f"p_{p['parametro_id']}", "11_PARAMETROS", f"C{i}")
    validacion_lista(ws, LISTA_NATURALEZA, f"F2:F{len(filas) + 1}")
    validacion_entero(ws, f"H2:H{len(filas) + 1}", 0, 1)
    espec.append(("11_PARAMETROS", cols))

    # ---------------- 01_FAMILIAS
    ws = wb.create_sheet("01_FAMILIAS", 0)
    cols = [Col("familia_id", "txt", 12, "Clave de familia.", nat="REAL"),
            Col("familia", "txt", 32, "Familia productiva declarada por la empresa.", nat="REAL"),
            Col("unidades_vendidas_2025", "int", 14, "Unidades agregadas 2025 del archivo de la empresa.", "und", nat="REAL",
                python="datos.leer → 12_PRECIOS (precio medio)"),
            Col("ventas_2025_soles", "soles", 16, "Ventas agregadas 2025 del archivo de la empresa.", "S/", nat="REAL",
                python="datos.leer → 12_PRECIOS (precio medio)"),
            Col("precio_promedio_soles", "soles", 16, "= ventas / unidades.", "S/ por und", nat="CALCULADO"),
            Col("penalidad_fija_frac", "pct", 14, "Fracción del valor del pedido si se entrega tarde.", "fracción",
                nat="ESTIMADO", python="Objetivo f1 (P_j)"),
            Col("penalidad_diaria_frac", "pct", 14, "Fracción diaria (0 = no se usa).", "fracción/día", nat="ESTIMADO",
                python="datos.leer (no interviene en f1)"),
            Col("naturaleza", "txt", 13, "Naturaleza principal de la fila.")]
    filas = [[f["familia_id"], f["familia"], f["unidades_2025"], f["ventas_2025"], f"=D{i}/C{i}", f["penalidad_fija"],
              f["penalidad_diaria"], f["naturaleza"]] for i, f in enumerate(N["familias"], 2)]
    escribir_tabla(ws, cols, filas, tab="oper")
    espec.insert(0, ("01_FAMILIAS", cols))

    # ---------------- 02_PRODUCTOS
    ws = wb.create_sheet("02_PRODUCTOS", 1)
    cols = [Col("producto_id", "txt", 11, "Clave de producto."), Col("familia_id", "txt", 11, "Familia (01_FAMILIAS)."),
            Col("producto", "txt", 36, "Nombre del producto (clave que usa Python).", python="datos.leer → 01_PRODUCTOS"),
            Col("tipo", "txt", 24, "Tipo comercial."),
            Col("alto_cm", "dec1", 9, "Alto exterior.", "cm", nat="ESTIMADO"),
            Col("ancho_cm", "dec1", 9, "Ancho exterior.", "cm", nat="ESTIMADO"),
            Col("fondo_cm", "dec1", 9, "Fondo exterior.", "cm", nat="ESTIMADO"),
            Col("puertas", "int", 9, "Número de puertas (vacío = no aplica).", "und", oblig=0),
            Col("lote_guia_und", "int", 11, "Tamaño de lote de fabricación.", "und", nat="ESTIMADO",
                python="Lote por defecto del pedido; duración (ROUNDUP(q/lote)·setup)"),
            Col("n_componentes", "int", 12, "Líneas de despiece (fórmula).", nat="CALCULADO"),
            Col("n_pasos_ruta", "int", 11, "Pasos de ruta (fórmula).", nat="CALCULADO"),
            Col("dimensiones_genericas", "bin", 13, "1 si ≥5 productos comparten exactamente las mismas dimensiones "
                "(indicio de plantilla, no de medición).", nat="CALCULADO"),
            Col("naturaleza", "txt", 12, "Naturaleza principal de la fila.")]
    filas = []
    for i, p in enumerate(N["productos"], 2):
        filas.append([p["producto_id"], p["familia_id"], p["producto"], p["tipo"], p["alto_cm"], p["ancho_cm"],
                      p["fondo_cm"], p["puertas"], p["lote_guia"],
                      f"=COUNTIF({R('04_DESPIECE', 'B', nbom)},A{i})", f"=COUNTIF({R('06_RUTAS', 'A', nrut)},A{i})",
                      f"=IF(COUNTIFS($E$2:$E${nprod + 1},E{i},$F$2:$F${nprod + 1},F{i},$G$2:$G${nprod + 1},G{i})>=5,1,0)",
                      p["naturaleza"]])
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_lista(ws, f"='01_FAMILIAS'!$A$2:$A${len(N['familias']) + 1}", f"B2:B{nprod + 1}")
    validacion_entero(ws, f"I2:I{nprod + 1}", 1)
    validacion_lista(ws, LISTA_NATURALEZA, f"M2:M{nprod + 1}")
    espec.insert(1, ("02_PRODUCTOS", cols))

    # ---------------- 03_MATERIALES
    ws = wb.create_sheet("03_MATERIALES", 2)
    cols = [Col("material_id", "txt", 12, "Clave técnica del material."), Col("material", "txt", 52, "Descripción."),
            Col("grupo", "txt", 22, "Grupo normalizado."), Col("unidad_costo", "txt", 10, "Unidad del costo unitario."),
            Col("espesor_mm", "dec2", 10, "Espesor nominal (vacío = no especificado).", "mm", oblig=0),
            Col("ancho_m", "dec2", 9, "Ancho del formato de compra.", "m", oblig=0),
            Col("largo_m", "dec2", 9, "Largo del formato de compra.", "m", oblig=0),
            Col("costo_unitario_soles", "soles", 14, "Costo por unidad de costo (vacío = sin cotización).", "S/", oblig=0,
                nat="ESTIMADO"),
            Col("material_generico", "bin", 12, "1 = material genérico (sección o tipo por definir)."),
            Col("naturaleza", "txt", 12, "ESTIMADO si tiene costo; PENDIENTE si falta el costo.")]
    filas = [[m["material_id"], m["material"], m["grupo"], m["unidad_costo"], m["espesor_mm"], m["ancho_m"], m["largo_m"],
              m["costo_unitario"], m["generico"], m["naturaleza"]] for m in N["materiales"]]
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_decimal(ws, f"H2:H{nmat + 1}", 0)
    validacion_lista(ws, LISTA_NATURALEZA, f"J2:J{nmat + 1}")
    espec.insert(2, ("03_MATERIALES", cols))

    # ---------------- 04_DESPIECE
    ws = wb.create_sheet("04_DESPIECE", 3)
    MID, PRE, ESP = R("03_MATERIALES", "A", nmat), R("03_MATERIALES", "H", nmat), R("03_MATERIALES", "E", nmat)
    cols = [Col("bom_id", "txt", 11, "Clave de línea de despiece."), Col("producto_id", "txt", 11, "Producto."),
            Col("componente", "txt", 30, "Componente."), Col("material_id", "txt", 12, "Material (03_MATERIALES)."),
            Col("tipo_calculo", "txt", 17, "CHAPA_AREA | PERFIL_LONGITUD | UNIDAD | SIN_PRECIO."),
            Col("piezas_por_unidad", "int", 10, "Piezas por unidad de producto.", "pzas", nat="ESTIMADO"),
            Col("ancho_pieza_mm", "dec1", 10, "Desarrollo de la pieza de chapa.", "mm", oblig=0),
            Col("alto_pieza_mm", "dec1", 10, "Desarrollo de la pieza de chapa.", "mm", oblig=0),
            Col("longitud_pieza_m", "dec3", 10, "Longitud de perfil por pieza.", "m", oblig=0),
            Col("area_chapa_m2_pieza", "dec6", 12, "Área por PIEZA (fórmula ancho×alto si hay medidas).", "m²", oblig=0),
            Col("masa_kg_producto", "dec4", 12, "Masa por unidad de producto (fórmula; vacío si no es calculable).", "kg",
                oblig=0, nat="CALCULADO", python="03_DESPIECE → masa de chapa (reprocesos)"),
            Col("costo_soles_producto", "soles", 14, "Costo de la línea por unidad de producto (fórmula).", "S/", oblig=0,
                nat="CALCULADO"),
            Col("planchas_equiv_producto", "dec4", 12, "Piezas × área / área de plancha (corregido en v2).", "planchas",
                oblig=0, nat="CALCULADO"),
            Col("naturaleza", "txt", 12, "Naturaleza de la fila.")]
    filas = []
    for i, b in enumerate(N["despiece"], 2):
        precio = f"INDEX({PRE},MATCH(D{i},{MID},0))"
        filas.append([
            b["bom_id"], b["producto_id"], b["componente"], b["material_id"], b["tipo_calculo"], b["piezas"],
            b["ancho_mm"], b["alto_mm"], b["longitud_m"],
            (f"=G{i}*H{i}/1000000" if b["area_formula"] else b["area_m2"]),
            f'=IF(E{i}="CHAPA_AREA",F{i}*J{i}*INDEX({ESP},MATCH(D{i},{MID},0))/1000*p_densidad_acero,'
            f'IF(E{i}="PERFIL_LONGITUD",F{i}*I{i}*p_kg_por_m_perfil,""))',
            f'=IF(E{i}="CHAPA_AREA",K{i}*{precio}/p_aprovechamiento_chapa,IF(E{i}="PERFIL_LONGITUD",'
            f'K{i}*{precio}/p_aprovechamiento_perfil,IF(E{i}="UNIDAD",F{i}*{precio},"")))',
            f'=IF(N(J{i})>0,F{i}*J{i}/p_area_plancha_m2,"")', b["naturaleza"]])
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_lista(ws, f"='02_PRODUCTOS'!$A$2:$A${nprod + 1}", f"B2:B{nbom + 1}")
    validacion_lista(ws, f"='03_MATERIALES'!$A$2:$A${nmat + 1}", f"D2:D{nbom + 1}")
    validacion_lista(ws, '"CHAPA_AREA,PERFIL_LONGITUD,UNIDAD,SIN_PRECIO"', f"E2:E{nbom + 1}")
    validacion_entero(ws, f"F2:F{nbom + 1}", 1)
    espec.insert(3, ("04_DESPIECE", cols))

    # ---------------- 05_RECURSOS
    ws = wb.create_sheet("05_RECURSOS", 4)
    cols = [Col("recurso_id", "txt", 10, "Clave del recurso."),
            Col("recurso", "txt", 46, "Nombre (clave que usa Python).", python="datos.leer → 04_MAQUINAS"),
            Col("tipo_tecnico", "txt", 40, "Descripción técnica."),
            Col("tipo_recurso", "txt", 15, "MAQUINA | EQUIPO_MANUAL | PUESTO_MANUAL | AREA | EXTERNO.", nat="ESTIMADO"),
            Col("cantidad", "int", 9, "Unidades del recurso en el escenario.", "und", nat="ESTIMADO",
                python="Validación: recurso disponible si ≥ 1 (el MILP trata cada recurso con capacidad 1)"),
            Col("costo_uso_soles_h", "soles", 12, "Tarifa de uso (vacío = sin dato).", "S/ por h", oblig=0, nat="ESTIMADO"),
            Col("potencia_kw_ref", "dec2", 10, "Potencia de catálogo.", "kW", nat="ESTIMADO", python="Energía (indicador)"),
            Col("factor_carga", "dec2", 9, "Factor medio de carga eléctrica.", "fracción", nat="ESTIMADO",
                python="Energía (indicador)"),
            Col("grupo_operarios", "txt", 18, "Grupo de personal.", nat="ESTIMADO"),
            Col("recurso_critico", "bin", 10, "1 = recurso de estudio (plegadora)."),
            Col("inventario_verificado", "bin", 12, "1 solo cuando la empresa confirme la existencia física.", nat="PENDIENTE"),
            Col("naturaleza", "txt", 12, "Naturaleza de la fila.")]
    filas = [[r["recurso_id"], r["recurso"], r["tipo_tecnico"], r["tipo_recurso"], r["cantidad"], r["costo_uso"],
              r["potencia_kw"], r["factor_carga"], r["grupo_operarios"], r["critico"], r["inventario_verificado"],
              r["naturaleza"]] for r in N["recursos"]]
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_lista(ws, '"MAQUINA,EQUIPO_MANUAL,PUESTO_MANUAL,AREA,EXTERNO"', f"D2:D{nrec + 1}")
    validacion_entero(ws, f"E2:E{nrec + 1}", 0)
    validacion_entero(ws, f"J2:K{nrec + 1}", 0, 1)
    espec.insert(4, ("05_RECURSOS", cols))

    # ---------------- 06_RUTAS
    ws = wb.create_sheet("06_RUTAS", 5)
    cols = [Col("producto_id", "txt", 11, "Producto."), Col("paso", "int", 6, "Número de paso (1..n)."),
            Col("operacion", "txt", 32, "Operación."), Col("recurso_id", "txt", 10, "Recurso principal."),
            Col("paso_predecesor", "int", 10, "Paso anterior (vacío = primera operación).", oblig=0),
            Col("min_fijos_lote", "dec2", 10, "Preparación por lote.", "min por lote", nat="ESTIMADO",
                python="p_om = (ROUNDUP(q/lote)·setup + q·u)/60"),
            Col("min_por_unidad", "dec4", 10, "Tiempo por unidad de producto.", "min por und", nat="ESTIMADO",
                python="p_om"),
            Col("tratamiento_tiempo", "txt", 13, "POR_UNIDAD | POR_LOTE.", python="Recuperación tras avería"),
            Col("min_ciclo_horno_incluido", "dec1", 11, "Minutos de ciclo de horno ya incluidos en min_fijos_lote.", "min",
                oblig=0),
            Col("origen_tiempo", "txt", 26, "HEREDADO_SIN_CRONOMETRAJE | HIPOTESIS_TECNICA."),
            Col("naturaleza", "txt", 12, "Naturaleza de la fila.")]
    filas = [[r["producto_id"], r["paso"], r["operacion"], r["recurso_id"], r["paso_predecesor"], r["min_fijos"],
              r["min_unidad"], r["tratamiento"], r["min_ciclo_horno"], r["origen_tiempo"], r["naturaleza"]]
             for r in N["rutas"]]
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_lista(ws, f"='02_PRODUCTOS'!$A$2:$A${nprod + 1}", f"A2:A{nrut + 1}")
    validacion_lista(ws, f"='05_RECURSOS'!$A$2:$A${nrec + 1}", f"D2:D{nrut + 1}")
    validacion_lista(ws, '"POR_UNIDAD,POR_LOTE"', f"H2:H{nrut + 1}")
    validacion_lista(ws, '"HEREDADO_SIN_CRONOMETRAJE,HIPOTESIS_TECNICA,CRONOMETRADO"', f"J2:J{nrut + 1}")
    validacion_decimal(ws, f"F2:G{nrut + 1}", 0)
    espec.insert(5, ("06_RUTAS", cols))

    # ---------------- 07_ELEGIBILIDAD
    ws = wb.create_sheet("07_ELEGIBILIDAD", 6)
    cols = [Col("producto_id", "txt", 11, "Producto."), Col("paso", "int", 6, "Paso."),
            Col("recurso_id", "txt", 10, "Recurso elegible.", python="Conjunto M_o del MILP"),
            Col("tipo_asignacion", "txt", 18, "PRINCIPAL | ALT_UTILLAJE | ALT_MISMO_PROCESO."),
            Col("elegible", "bin", 8, "1 = elegible."),
            Col("factor_tiempo", "dec2", 9, "Multiplicador del tiempo de ruta en ese recurso.", "×", nat="ESTIMADO",
                python="p_om en máquina alternativa"),
            Col("naturaleza", "txt", 12, "CALCULADO (principal) / ESTIMADO (alternativa).")]
    filas = [[e["producto_id"], e["paso"], e["recurso_id"], e["tipo_asignacion"], e["elegible"], e["factor_tiempo"],
              e["naturaleza"]] for e in N["elegibilidad"]]
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_lista(ws, f"='05_RECURSOS'!$A$2:$A${nrec + 1}", f"C2:C{nele + 1}")
    validacion_lista(ws, '"PRINCIPAL,ALT_UTILLAJE,ALT_MISMO_PROCESO"', f"D2:D{nele + 1}")
    validacion_entero(ws, f"E2:E{nele + 1}", 0, 1)
    espec.insert(6, ("07_ELEGIBILIDAD", cols))

    # ---------------- 08_PERSONAL
    ws = wb.create_sheet("08_PERSONAL", 7)
    cols = [Col("operario_id", "txt", 11, "Seudónimo del trabajador (25 confirmados).", nat="REAL",
                python="06_OPERARIOS (lista de trabajadores)"),
            Col("turno", "int", 7, "Turno asignado.", nat="REAL"),
            Col("hora_entrada", "hora", 10, "Entrada.", nat="ESTIMADO"), Col("hora_salida", "hora", 10, "Salida.", nat="ESTIMADO"),
            Col("descanso_h", "dec1", 9, "Descanso no productivo.", "h", nat="ESTIMADO"),
            Col("horas_netas_dia", "dec1", 10, "Horas netas confirmadas.", "h", nat="REAL"),
            Col("control_horario_h", "dec2", 11, "(salida − entrada) − descanso; debe igualar horas netas.", "h",
                nat="CALCULADO"),
            Col("naturaleza", "txt", 12, "Naturaleza principal (existencia y jornada).")]
    filas = [[p["operario_id"], p["turno"], p["hora_entrada"], p["hora_salida"], p["descanso_h"], p["horas_netas"],
              f"=(D{i}-C{i})*24-E{i}", p["naturaleza"]] for i, p in enumerate(N["personal"], 2)]
    escribir_tabla(ws, cols, filas, tab="oper")
    espec.insert(7, ("08_PERSONAL", cols))

    # ---------------- 09_CALENDARIO
    ws = wb.create_sheet("09_CALENDARIO", 8)
    cols = [Col("fecha", "fecha", 12, "Fecha.", python="07_CALENDARIO"),
            Col("dia_semana", "int", 9, "1 = lunes … 7 = domingo (fórmula).", nat="CALCULADO"),
            Col("laborable_base", "bin", 10, "Día laborable del calendario base (lunes-viernes).", nat="ESTIMADO",
                python="Productivo (1/0)"),
            Col("feriado_nacional", "txt", 30, "Nombre del feriado nacional (vacío = no feriado).", oblig=0,
                nat="REFERENCIA", python="Feriado nacional Perú"),
            Col("productivo_efectivo", "bin", 11, "Fórmula: aplica feriados si p_aplicar_feriados = 1.", nat="CALCULADO"),
            Col("horas_productivas", "dec1", 10, "productivo_efectivo × horas por día.", "h", nat="CALCULADO"),
            Col("naturaleza", "txt", 12, "Naturaleza principal.")]
    filas = [[c["fecha"], f"=WEEKDAY(A{i},2)", c["laborable_base"], c["feriado"],
              f'=IF(AND(p_aplicar_feriados=1,D{i}<>""),0,C{i})', f"=E{i}*p_horas_productivas_dia", c["naturaleza"]]
             for i, c in enumerate(N["calendario"], 2)]
    escribir_tabla(ws, cols, filas, tab="oper")
    validacion_entero(ws, f"C2:C{ncal + 1}", 0, 1)
    espec.insert(8, ("09_CALENDARIO", cols))

    # ---------------- 10_CAPACIDAD_MENSUAL
    ws = wb.create_sheet("10_CAPACIDAD_MENSUAL", 9)
    cols = [Col("recurso_id", "txt", 10, "Recurso."), Col("mes", "int", 6, "Mes del año base."),
            Col("dias_productivos", "int", 10, "Días productivos efectivos del mes (fórmula).", "días", nat="CALCULADO"),
            Col("horas_dia", "dec1", 8, "Horas productivas por día.", "h", nat="REAL"),
            Col("cantidad", "int", 8, "Unidades del recurso.", nat="ESTIMADO"),
            Col("horas_disponibles", "dec1", 11, "días × horas × cantidad (sin descontar paradas: no hay historial).", "h",
                nat="CALCULADO"),
            Col("naturaleza", "txt", 12, "")]
    CF, CE = R("09_CALENDARIO", "A", ncal), R("09_CALENDARIO", "E", ncal)
    filas, i = [], 2
    for r in N["recursos"]:
        for mes in range(1, 13):
            filas.append([r["recurso_id"], mes,
                          f"=SUMPRODUCT((MONTH({CF})=B{i})*(YEAR({CF})=p_anio_base)*{CE})", "=p_horas_productivas_dia",
                          f"=INDEX({R('05_RECURSOS', 'E', nrec)},MATCH(A{i},{R('05_RECURSOS', 'A', nrec)},0))",
                          f"=C{i}*D{i}*E{i}", "CALCULADO"])
            i += 1
    escribir_tabla(ws, cols, filas, tab="oper")
    espec.insert(9, ("10_CAPACIDAD_MENSUAL", cols))

    # ---------------- registros de captura (vacíos válidos)
    regs = [
        ("12_REG_ENTREGAS", "Pedidos y entregas reales (reemplazará a los pedidos simulados).", [
            Col("pedido_real_id", "txt", 14, "Código del pedido/OT de SAGAPI."), Col("producto_id", "txt", 11, "Producto."),
            Col("cliente_codigo", "txt", 12, "Código anónimo del cliente.", oblig=0),
            Col("unidades", "int", 9, "Unidades pedidas.", "und"),
            Col("fecha_recepcion", "fecha", 12, "AAAA-MM-DD."), Col("fecha_liberacion", "fecha", 12, "Fecha en que se puede iniciar (material disponible)."),
            Col("fecha_compromiso", "fecha", 12, "Fecha pactada con el cliente."),
            Col("fecha_entrega", "fecha", 12, "Fecha real de entrega.", oblig=0),
            Col("unidades_conformes", "int", 10, "", "und", oblig=0), Col("unidades_no_conformes", "int", 10, "", "und", oblig=0),
            Col("precio_venta_soles", "soles", 13, "Precio total facturado.", "S/", oblig=0),
            Col("penalidad_contractual_frac", "pct", 12, "Penalidad pactada (si existe).", oblig=0),
            Col("causa_atraso", "txt", 24, "", oblig=0), Col("fuente_documento", "txt", 24, "Orden de compra, guía, factura.")]),
        ("13_REG_PARADAS", "Paradas reales de máquinas (calibra fallas y MTBF/MTTR).", [
            Col("parada_id", "txt", 10, "Correlativo."), Col("recurso_id", "txt", 10, "Recurso (05_RECURSOS)."),
            Col("inicio", "fechahora", 18, "AAAA-MM-DD hh:mm."), Col("fin", "fechahora", 18, "AAAA-MM-DD hh:mm."),
            Col("tipo_parada", "txt", 18, "FALLA | MANTENIMIENTO | FALTA_MATERIAL | FALTA_PERSONAL | OTRO."),
            Col("causa", "txt", 30, ""), Col("pedido_afectado", "txt", 12, "", oblig=0),
            Col("piezas_perdidas", "int", 9, "", "pzas", oblig=0), Col("registrado_por", "txt", 14, ""),
            Col("fuente_documento", "txt", 22, "Bitácora, orden de trabajo de mantenimiento.")]),
        ("14_REG_CALIDAD", "Inspecciones y no conformidades reales.", [
            Col("inspeccion_id", "txt", 11, ""), Col("pedido_real_id", "txt", 14, ""), Col("producto_id", "txt", 11, ""),
            Col("fecha", "fecha", 12, "AAAA-MM-DD."), Col("cantidad_inspeccionada", "int", 11, "", "und"),
            Col("conformes", "int", 9, "", "und"), Col("no_conformes", "int", 9, "", "und"),
            Col("defecto", "txt", 24, "", oblig=0), Col("reproceso_h", "dec2", 10, "", "h", oblig=0),
            Col("material_perdido_soles", "soles", 12, "", "S/", oblig=0)]),
        ("15_REG_TIEMPOS", "Cronometraje de operaciones (valida 06_RUTAS).", [
            Col("medicion_id", "txt", 11, ""), Col("producto_id", "txt", 11, ""), Col("paso", "int", 6, ""),
            Col("recurso_id", "txt", 10, ""), Col("fecha", "fecha", 12, "AAAA-MM-DD."),
            Col("unidades_medidas", "int", 9, "", "und"), Col("min_preparacion_obs", "dec2", 11, "", "min"),
            Col("min_total_obs", "dec2", 11, "", "min"), Col("operarios_en_operacion", "int", 10, "", "personas"),
            Col("metodo", "txt", 14, "CRONOMETRO | VIDEO | REGISTRO_OT."), Col("observador", "txt", 14, "")]),
        ("16_REG_HABILIDADES", "Habilidades reales validadas por el jefe de planta (reemplaza la matriz simulada).", [
            Col("operario_id", "txt", 11, "08_PERSONAL."), Col("recurso_id", "txt", 10, "05_RECURSOS."),
            Col("nivel_0_3", "int", 9, "0 no opera; 1 con supervisión; 2 autónomo; 3 instructor."),
            Col("opera_solo_1_0", "bin", 9, ""), Col("validado_por", "txt", 16, ""),
            Col("fecha_validacion", "fecha", 12, "AAAA-MM-DD.")]),
    ]
    for nombre, desc, cols in regs:
        ws = wb.create_sheet(nombre)
        escribir_tabla(ws, cols, [], tab="reg")
        for j, c in enumerate(cols, 1):
            if c.nombre in ("producto_id",):
                validacion_lista(ws, f"='02_PRODUCTOS'!$A$2:$A${nprod + 1}", f"{openpyxl.utils.get_column_letter(j)}2:"
                                 f"{openpyxl.utils.get_column_letter(j)}1000")
            if c.nombre == "recurso_id":
                validacion_lista(ws, f"='05_RECURSOS'!$A$2:$A${nrec + 1}", f"{openpyxl.utils.get_column_letter(j)}2:"
                                 f"{openpyxl.utils.get_column_letter(j)}1000")
            if c.nombre == "operario_id":
                validacion_lista(ws, f"='08_PERSONAL'!$A$2:$A${len(N['personal']) + 1}",
                                 f"{openpyxl.utils.get_column_letter(j)}2:{openpyxl.utils.get_column_letter(j)}1000")
            if c.nombre == "tipo_parada":
                validacion_lista(ws, '"FALLA,MANTENIMIENTO,FALTA_MATERIAL,FALTA_PERSONAL,OTRO"', "E2:E1000")
            if c.nombre == "nivel_0_3":
                validacion_entero(ws, "C2:C1000", 0, 3)
        espec.append((nombre, cols))
    hoja_diccionario(wb, espec, "SAGAPI — Base de datos OPERATIVA v2.0 (datos reales, de referencia, calculados y "
                                "estimados; sin datos simulados).")
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "SAGAPI BD OPERATIVA v2.0"
    wb.properties.creator = "SAGAPI"
    wb.save(ruta)


# =============================================================================== libro SIMULACIÓN
def construir_simulacion(N, ruta):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    espec = []
    nped, nev = len(N["pedidos"]), len(N["eventos"])

    ws = wb.create_sheet("01_PEDIDOS_SIM")
    cols = [Col("pedido_id", "txt", 9, "Pedido simulado."), Col("producto_id", "txt", 11, "Producto (02_PRODUCTOS)."),
            Col("unidades", "int", 9, "Unidades.", "und", python="q_j"),
            Col("lote_und", "int", 8, "Tamaño de lote.", "und", python="ROUNDUP(q/lote) preparaciones"),
            Col("fecha_recepcion", "fecha", 12, "Recepción."),
            Col("fecha_liberacion", "fecha", 12, "Liberación (r_j a las 08:00).", python="r_j"),
            Col("fecha_compromiso", "fecha", 12, "Compromiso (d_j al cierre del día).", python="d_j"),
            Col("plazo_dias", "int", 8, "compromiso − recepción.", "días", nat="CALCULADO"),
            Col("liberacion_fin_semana", "bin", 11, "1 si la liberación cae sábado o domingo.", nat="CALCULADO"),
            Col("compromiso_fin_semana", "bin", 11, "1 si el compromiso cae sábado o domingo.", nat="CALCULADO"),
            Col("naturaleza", "txt", 11, "SIMULADO.")]
    filas = [[p["pedido_id"], p["producto_id"], p["unidades"], p["lote"], p["recepcion"], p["liberacion"], p["compromiso"],
              f"=G{i}-E{i}", f"=IF(WEEKDAY(F{i},2)>5,1,0)", f"=IF(WEEKDAY(G{i},2)>5,1,0)", p["naturaleza"]]
             for i, p in enumerate(N["pedidos"], 2)]
    escribir_tabla(ws, cols, filas, tab="sim")
    validacion_lista(ws, f"='99_LISTAS'!$A$2:$A${len(N['productos']) + 1}", f"B2:B{nped + 1}")
    validacion_entero(ws, f"C2:D{nped + 1}", 1)
    espec.append(("01_PEDIDOS_SIM", cols))

    ws = wb.create_sheet("02_PERFILES_SIM")
    cols = [Col("operario_id", "txt", 11, "Trabajador (08_PERSONAL)."),
            Col("especialidad_simulada", "txt", 24, "Perfil primario asignado sin validación."),
            Col("naturaleza", "txt", 11, "SIMULADO.")]
    escribir_tabla(ws, cols, [[p["operario_id"], p["especialidad"], "SIMULADO"] for p in N["personal"]], tab="sim")
    espec.append(("02_PERFILES_SIM", cols))

    ws = wb.create_sheet("03_HABILIDADES_SIM")
    cols = [Col("operario_id", "txt", 11, "Trabajador.")] + \
           [Col(r, "bin", 8, f"1 = habilitado (simulado) para {r}.") for r in N["hab_recursos"]] + \
           [Col("naturaleza", "txt", 11, "SIMULADO.")]
    escribir_tabla(ws, cols, [[h["operario_id"]] + h["valores"] + [h["naturaleza"]] for h in N["habilidades"]], tab="sim")
    validacion_entero(ws, f"B2:{openpyxl.utils.get_column_letter(len(N['hab_recursos']) + 1)}{len(N['habilidades']) + 1}",
                      0, 1, permitir_vacio=False)
    espec.append(("03_HABILIDADES_SIM", cols))

    ws = wb.create_sheet("04_EVENTOS_SIM")
    cols = [Col("evento_id", "txt", 9, "Evento."), Col("tipo_evento", "txt", 24, "FALLA | MANTENIMIENTO_PLANIFICADO."),
            Col("recurso_id", "txt", 10, "Recurso indisponible."),
            Col("inicio", "fechahora", 19, "Inicio de la indisponibilidad (reloj).", python="Ventana (m, a, b)"),
            Col("horas_reloj", "dec2", 9, "Duración en horas de reloj.", "h", python="b = inicio + horas"),
            Col("fin", "fechahora", 19, "inicio + horas_reloj/24.", nat="CALCULADO"),
            Col("horas_turno_afectadas", "dec4", 11, "Horas productivas cubiertas (calendario).", "h", nat="CALCULADO",
                python="Comprobación en 'validar'"),
            Col("motivo", "txt", 40, "Motivo hipotético."), Col("pedido_ref", "txt", 9, "Pedido en curso en el FIFO."),
            Col("paso_ref", "int", 7, "Paso en curso en el FIFO."), Col("naturaleza", "txt", 11, "SIMULADO.")]
    filas = [[e["evento_id"], e["tipo_evento"], e["recurso_id"], e["inicio"], e["horas_reloj"], f"=D{i}+E{i}/24",
              e["horas_turno"], e["motivo"], e["pedido_ref"], e["paso_ref"], e["naturaleza"]]
             for i, e in enumerate(N["eventos"], 2)]
    escribir_tabla(ws, cols, filas, tab="sim")
    validacion_lista(ws, '"FALLA,MANTENIMIENTO_PLANIFICADO"', f"B2:B{nev + 1}")
    validacion_lista(ws, f"='99_LISTAS'!$B$2:$B${len(N['recursos']) + 1}", f"C2:C{nev + 1}")
    validacion_decimal(ws, f"E2:E{nev + 1}", 0)
    espec.append(("04_EVENTOS_SIM", cols))

    ws = wb.create_sheet("05_ESCENARIOS_SIM")
    cols = [Col("escenario_id", "txt", 10, "Escenario."), Col("caso", "txt", 46, "Descripción breve."),
            Col("eventos", "txt", 44, "Eventos separados por '+' (vacío = sin falla).", oblig=0),
            Col("politica_recuperacion", "txt", 10, "A = reanudar; B = reiniciar lote; vacío = sin falla.", oblig=0),
            Col("semilla", "int", 10, "Semilla."), Col("naturaleza", "txt", 13, "")]
    filas = [[s["escenario_id"], s["caso"], s["eventos"], s["politica"], s["semilla"], s["naturaleza"]]
             for s in N["escenarios"]]
    escribir_tabla(ws, cols, filas, tab="sim")
    validacion_lista(ws, '"A,B"', f"D2:D{len(filas) + 1}")
    espec.append(("05_ESCENARIOS_SIM", cols))

    ws = wb.create_sheet("06_RECUPERACION_SIM")
    cols = [Col("ambito", "txt", 20, "PRODUCTO_PASO | GENERICO_POR_UNIDAD | GENERICO_POR_LOTE."),
            Col("producto_id", "txt", 11, "Producto (solo PRODUCTO_PASO).", oblig=0),
            Col("paso", "int", 6, "Paso (solo PRODUCTO_PASO).", oblig=0), Col("politica", "txt", 8, "A | B."),
            Col("conserva_piezas", "bin", 9, ""), Col("reinicia_lote", "bin", 9, ""), Col("repite_setup", "bin", 9, ""),
            Col("piezas_a_repetir", "int", 9, "Piezas que se repiten (vacío si se reinicia el lote completo).", "pzas",
                oblig=0),
            Col("min_verificacion", "dec1", 10, "Verificación tras reparación.", "min"),
            Col("naturaleza", "txt", 11, "ESTIMADO (hipótesis técnica sin validar).")]
    filas = [[r["ambito"], r["producto_id"], r["paso"], r["politica"], r["conserva"], r["reinicia"], r["repite_setup"],
              r["piezas_repetir"], r["min_verificacion"], r["naturaleza"]] for r in N["recuperacion"]]
    escribir_tabla(ws, cols, filas, tab="sim")
    validacion_lista(ws, '"PRODUCTO_PASO,GENERICO_POR_UNIDAD,GENERICO_POR_LOTE"', f"A2:A{len(filas) + 1}")
    validacion_lista(ws, '"A,B"', f"D2:D{len(filas) + 1}")
    espec.append(("06_RECUPERACION_SIM", cols))

    ws = wb.create_sheet("07_SENSIBILIDAD_SIM")
    cols = [Col("factor", "txt", 36, ""), Col("niveles", "txt", 22, "Separados por ';'."), Col("aplicado_a", "txt", 26, ""),
            Col("proposito", "txt", 46, ""), Col("naturaleza", "txt", 14, "")]
    escribir_tabla(ws, cols, [[s["factor"], s["niveles"], s["aplicado_a"], s["proposito"], s["naturaleza"]]
                              for s in N["sensibilidad"]], tab="sim")
    espec.append(("07_SENSIBILIDAD_SIM", cols))

    ws = wb.create_sheet("08_CONFIG_MODELO")
    cols = [Col("parametro_id", "txt", 24, ""), Col("parametro", "txt", 44, "Nombre que lee Python cuando aplica."),
            Col("valor_num", "gen", 12, "", oblig=0), Col("valor_txt", "txt", 46, "", oblig=0), Col("unidad", "txt", 14, ""),
            Col("naturaleza", "txt", 15, ""), Col("nota", "txt", 60, "", oblig=0), Col("usado_por_python", "bin", 10, "")]
    pars = [p for p in N["parametros"] if p["libro"] == "SIM"]
    escribir_tabla(ws, cols, [[p["parametro_id"], p["parametro"], p["valor_num"], p["valor_txt"], p["unidad"],
                               p["naturaleza"], p["nota"], p["usado_python"]] for p in pars], tab="sim")
    espec.append(("08_CONFIG_MODELO", cols))

    ws = wb.create_sheet("09_PROGRAMA_FIFO_REF")
    cols = [Col("pedido_id", "txt", 9, ""), Col("paso", "int", 6, ""), Col("recurso_id", "txt", 10, ""),
            Col("inicio", "fechahora", 19, ""), Col("fin", "fechahora", 19, ""),
            Col("duracion_h", "dec6", 11, "Horas productivas.", "h"),
            Col("naturaleza", "txt", 12, "CALCULADO por regla FIFO sobre pedidos simulados (Python lo reproduce).")]
    escribir_tabla(ws, cols, [[f["pedido_id"], f["paso"], f["recurso_id"], f["inicio"], f["fin"], f["duracion_h"],
                               f["naturaleza"]] for f in N["fifo"]], tab="der")
    espec.append(("09_PROGRAMA_FIFO_REF", cols))

    ws = wb.create_sheet("99_LISTAS")
    cols = [Col("producto_id", "txt", 11, "Copia de referencia de 02_PRODUCTOS (solo para validación de datos).", oblig=0),
            Col("recurso_id", "txt", 10, "Copia de referencia de 05_RECURSOS.", oblig=0),
            Col("operario_id", "txt", 11, "Copia de referencia de 08_PERSONAL.", oblig=0)]
    n = max(len(N["productos"]), len(N["recursos"]), len(N["personal"]))
    filas = [[N["productos"][k]["producto_id"] if k < len(N["productos"]) else None,
              N["recursos"][k]["recurso_id"] if k < len(N["recursos"]) else None,
              N["personal"][k]["operario_id"] if k < len(N["personal"]) else None] for k in range(n)]
    escribir_tabla(ws, cols, filas, tab="dicc")
    espec.append(("99_LISTAS", cols))
    hoja_diccionario(wb, espec, "SAGAPI — Base complementaria de SIMULACIÓN v2.0 (pedidos, fallas, escenarios y "
                                "habilidades artificiales para probar el modelo; NO son registros de SAGAPI).")
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "SAGAPI BD SIMULACION v2.0"
    wb.save(ruta)


# =============================================================================== lectura de los libros fuente
def _leer_fuente(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=False)
    out = {}
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        enc = list(filas[0])
        out[ws.title] = [dict(zip(enc, f)) for f in filas[1:] if f and any(v is not None for v in f)]
        out["_enc_" + ws.title] = enc
    return out


def leer_fuentes(dir_salida):
    """Reconstruye N (estructura normalizada) desde OPERATIVA + SIMULACION, para exportar el modelo."""
    O = _leer_fuente(os.path.join(dir_salida, ARCH_OPER))
    S = _leer_fuente(os.path.join(dir_salida, ARCH_SIM))
    lit = lambda v: None if isinstance(v, str) and v.startswith("=") else v  # noqa: E731
    N = {}
    N["familias"] = [dict(familia_id=f["familia_id"], familia=f["familia"], unidades_2025=f["unidades_vendidas_2025"],
                          ventas_2025=f["ventas_2025_soles"], penalidad_fija=f["penalidad_fija_frac"],
                          penalidad_diaria=f["penalidad_diaria_frac"], naturaleza=f["naturaleza"]) for f in O["01_FAMILIAS"]]
    N["productos"] = [dict(producto_id=f["producto_id"], familia_id=f["familia_id"], producto=f["producto"], tipo=f["tipo"],
                           alto_cm=f["alto_cm"], ancho_cm=f["ancho_cm"], fondo_cm=f["fondo_cm"], puertas=f["puertas"],
                           lote_guia=f["lote_guia_und"], naturaleza=f["naturaleza"]) for f in O["02_PRODUCTOS"]]
    N["materiales"] = [dict(material_id=f["material_id"], material=f["material"], grupo=f["grupo"],
                            unidad_costo=f["unidad_costo"], espesor_mm=f["espesor_mm"], ancho_m=f["ancho_m"],
                            largo_m=f["largo_m"], costo_unitario=f["costo_unitario_soles"],
                            generico=f["material_generico"], naturaleza=f["naturaleza"]) for f in O["03_MATERIALES"]]
    N["despiece"] = []
    for f in O["04_DESPIECE"]:
        af = isinstance(f["area_chapa_m2_pieza"], str) and f["area_chapa_m2_pieza"].startswith("=")
        N["despiece"].append(dict(bom_id=f["bom_id"], producto_id=f["producto_id"], componente=f["componente"],
                                  material_id=f["material_id"], tipo_calculo=f["tipo_calculo"],
                                  piezas=f["piezas_por_unidad"], ancho_mm=f["ancho_pieza_mm"], alto_mm=f["alto_pieza_mm"],
                                  longitud_m=f["longitud_pieza_m"], area_m2=lit(f["area_chapa_m2_pieza"]),
                                  area_formula=af, naturaleza=f["naturaleza"]))
    N["recursos"] = [dict(recurso_id=f["recurso_id"], recurso=f["recurso"], tipo_tecnico=f["tipo_tecnico"],
                          tipo_recurso=f["tipo_recurso"], cantidad=f["cantidad"], costo_uso=f["costo_uso_soles_h"],
                          potencia_kw=f["potencia_kw_ref"], factor_carga=f["factor_carga"],
                          grupo_operarios=f["grupo_operarios"], critico=f["recurso_critico"],
                          naturaleza=f["naturaleza"]) for f in O["05_RECURSOS"]]
    N["rutas"] = [dict(producto_id=f["producto_id"], paso=f["paso"], operacion=f["operacion"], recurso_id=f["recurso_id"],
                       paso_predecesor=f["paso_predecesor"], min_fijos=f["min_fijos_lote"], min_unidad=f["min_por_unidad"],
                       tratamiento=f["tratamiento_tiempo"], min_ciclo_horno=f["min_ciclo_horno_incluido"],
                       origen_tiempo=f["origen_tiempo"], naturaleza=f["naturaleza"]) for f in O["06_RUTAS"]]
    N["elegibilidad"] = [dict(producto_id=f["producto_id"], paso=f["paso"], recurso_id=f["recurso_id"],
                              tipo_asignacion=f["tipo_asignacion"], elegible=f["elegible"], factor_tiempo=f["factor_tiempo"],
                              naturaleza=f["naturaleza"]) for f in O["07_ELEGIBILIDAD"]]
    esp = {f["operario_id"]: f["especialidad_simulada"] for f in S["02_PERFILES_SIM"]}
    N["personal"] = [dict(operario_id=f["operario_id"], turno=f["turno"], horas_netas=f["horas_netas_dia"],
                          especialidad=esp.get(f["operario_id"]), naturaleza=f["naturaleza"]) for f in O["08_PERSONAL"]]
    N["calendario"] = [dict(fecha=f["fecha"], laborable_base=f["laborable_base"], feriado=f["feriado_nacional"],
                            naturaleza=f["naturaleza"]) for f in O["09_CALENDARIO"]]
    N["parametros"] = []
    for libro, hoja in (("OPER", O["11_PARAMETROS"]), ("SIM", S["08_CONFIG_MODELO"])):
        for f in hoja:
            N["parametros"].append(dict(parametro_id=f["parametro_id"], parametro=f["parametro"], valor_num=f["valor_num"],
                                        valor_txt=f["valor_txt"], unidad=f["unidad"], naturaleza=f["naturaleza"],
                                        nota=f["nota"], usado_python=f["usado_por_python"], libro=libro))
    N["pedidos"] = [dict(pedido_id=f["pedido_id"], producto_id=f["producto_id"], unidades=f["unidades"], lote=f["lote_und"],
                         recepcion=f["fecha_recepcion"], liberacion=f["fecha_liberacion"],
                         compromiso=f["fecha_compromiso"], naturaleza=f["naturaleza"]) for f in S["01_PEDIDOS_SIM"]]
    enc = S["_enc_03_HABILIDADES_SIM"]
    N["hab_recursos"] = [c for c in enc[1:] if c and c != "naturaleza"]
    N["habilidades"] = [dict(operario_id=f["operario_id"], valores=[f[r] for r in N["hab_recursos"]],
                             naturaleza=f["naturaleza"]) for f in S["03_HABILIDADES_SIM"]]
    N["eventos"] = [dict(evento_id=f["evento_id"], tipo_evento=f["tipo_evento"], recurso_id=f["recurso_id"],
                         inicio=f["inicio"], horas_reloj=f["horas_reloj"], horas_turno=f["horas_turno_afectadas"],
                         motivo=f["motivo"], pedido_ref=f["pedido_ref"], paso_ref=f["paso_ref"],
                         naturaleza=f["naturaleza"]) for f in S["04_EVENTOS_SIM"]]
    N["escenarios"] = [dict(escenario_id=f["escenario_id"], caso=f["caso"], eventos=f["eventos"],
                            politica=f["politica_recuperacion"], semilla=f["semilla"], naturaleza=f["naturaleza"])
                       for f in S["05_ESCENARIOS_SIM"]]
    N["recuperacion"] = [dict(ambito=f["ambito"], producto_id=f["producto_id"], paso=f["paso"], politica=f["politica"],
                              conserva=f["conserva_piezas"], reinicia=f["reinicia_lote"], repite_setup=f["repite_setup"],
                              piezas_repetir=f["piezas_a_repetir"], min_verificacion=f["min_verificacion"],
                              naturaleza=f["naturaleza"]) for f in S["06_RECUPERACION_SIM"]]
    N["sensibilidad"] = [dict(factor=f["factor"], niveles=f["niveles"], aplicado_a=f["aplicado_a"], proposito=f["proposito"],
                              naturaleza=f["naturaleza"]) for f in S["07_SENSIBILIDAD_SIM"]]
    N["fifo"] = [dict(pedido_id=f["pedido_id"], paso=f["paso"], recurso_id=f["recurso_id"], inicio=f["inicio"], fin=f["fin"],
                      duracion_h=f["duracion_h"], naturaleza=f["naturaleza"]) for f in S["09_PROGRAMA_FIFO_REF"]]
    return N


# =============================================================================== libro MODELO (compatible)
def construir_modelo(N, ruta):
    """Esquema heredado (nombres de hoja y encabezado exactos) que lee Optimizacion_MILP_SAGAPI.py 1.4.0.

    Las columnas que lee Python se escriben como VALORES (Python no evalúa fórmulas). Las columnas de
    análisis que Python no lee son fórmulas vivas.
    """
    par = {p["parametro_id"]: p["valor_num"] for p in N["parametros"]}
    calc_despiece(N, par)
    fam = {f["familia_id"]: f for f in N["familias"]}
    prod = {p["producto_id"]: p for p in N["productos"]}
    rec = {r["recurso_id"]: r for r in N["recursos"]}
    mat = {m["material_id"]: m for m in N["materiales"]}
    pnom = lambda pid: prod[pid]["producto"]  # noqa: E731
    rnom = lambda rid: rec[rid]["recurso"]  # noqa: E731
    pfam = lambda pid: fam[prod[pid]["familia_id"]]["familia"]  # noqa: E731
    ruta_op = {(r["producto_id"], r["paso"]): r for r in N["rutas"]}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n = {k: len(v) for k, v in N.items() if isinstance(v, list)}
    R = lambda hoja, col, k: f"'{hoja}'!${col}$2:${col}${k + 1}"  # noqa: E731
    indice = []

    def hoja(nombre, rol, origen, tab):
        indice.append([nombre, rol, origen])
        ws = wb.create_sheet(nombre)
        ws.sheet_properties.tabColor = {"oper": "2E7D32", "sim": "E65100", "der": "595959", "ctl": "B8860B",
                                        "dicc": "7F7F7F"}[tab]
        return ws

    # ---------------- 19_PARAMETROS (primero, para los nombres definidos)
    ws = hoja("19_PARAMETROS", "ENTRADA_PYTHON", "OPERATIVA 11_PARAMETROS + SIMULACION 08_CONFIG_MODELO", "oper")
    cols = [Col("Parámetro", "txt", 46), Col("Valor", "gen", 40), Col("Unidad", "txt", 16), Col("Naturaleza", "txt", 15),
            Col("Nota", "txt", 60), Col("Leído por Python (1/0)", "bin", 12)]
    filas = []
    for p in N["parametros"]:
        if p["parametro_id"] == "categoria_empresarial":
            continue
        v = p["valor_num"] if p["valor_num"] is not None else p["valor_txt"]
        filas.append([p["parametro"], v, p["unidad"], p["naturaleza"], p["nota"], p["usado_python"]])
    escribir_tabla(ws, cols, filas, comentarios=False)
    pos = {p[0]: i for i, p in enumerate(filas, 2)}
    for pid, nombre in (("horas_productivas_dia", "Horas productivas por día"),
                        ("aplicar_feriados", "Aplicar feriados nacionales (1/0)"), ("costo_hora_hombre", "Costo hora-hombre"),
                        ("anio_base", "Año base del estudio")):
        nombre_definido(wb, f"p_{pid}", "19_PARAMETROS", f"B{pos[nombre]}")

    # ---------------- 01_PRODUCTOS
    ws = hoja("01_PRODUCTOS", "ENTRADA_PYTHON", "OPERATIVA 02_PRODUCTOS", "oper")
    cols = [Col("Familia", "txt", 28), Col("Nombre completo", "txt", 36), Col("Tipo", "txt", 22), Col("Alto cm", "dec1", 8),
            Col("Ancho cm", "dec1", 8), Col("Fondo cm", "dec1", 8), Col("Puertas", "int", 8), Col("Lote guía", "int", 8),
            Col("Componentes", "int", 11), Col("Pasos de ruta", "int", 10), Col("Familia válida (1/0)", "bin", 10),
            Col("Dimensiones genéricas (1/0)", "bin", 12), Col("Naturaleza", "txt", 12)]
    np_ = n["productos"]
    filas = [[pfam(p["producto_id"]), p["producto"], p["tipo"], p["alto_cm"], p["ancho_cm"], p["fondo_cm"], p["puertas"],
              p["lote_guia"], f"=COUNTIF({R('03_DESPIECE', 'A', n['despiece'])},B{i})",
              f"=COUNTIF({R('05_RUTAS', 'A', n['rutas'])},B{i})",
              f"=IF(COUNTIF({R('12_PRECIOS', 'A', n['familias'])},A{i})>0,1,0)",
              f"=IF(COUNTIFS($D$2:$D${np_ + 1},D{i},$E$2:$E${np_ + 1},E{i},$F$2:$F${np_ + 1},F{i})>=5,1,0)",
              p["naturaleza"]] for i, p in enumerate(N["productos"], 2)]
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 02_MATERIALES
    ws = hoja("02_MATERIALES", "REFERENCIA", "OPERATIVA 03_MATERIALES", "oper")
    cols = [Col("Material o accesorio (nombre completo)", "txt", 50), Col("Clave técnica auxiliar", "txt", 12),
            Col("Grupo", "txt", 20), Col("Unidad de costo", "txt", 9), Col("Espesor mm", "dec2", 9), Col("Ancho m", "dec2", 8),
            Col("Largo m", "dec2", 8), Col("Costo unitario S/", "soles", 12), Col("Material genérico (1/0)", "bin", 10),
            Col("Naturaleza", "txt", 12)]
    escribir_tabla(ws, cols, [[m["material"], m["material_id"], m["grupo"], m["unidad_costo"], m["espesor_mm"], m["ancho_m"],
                               m["largo_m"], m["costo_unitario"], m["generico"], m["naturaleza"]] for m in N["materiales"]],
                   comentarios=False)

    # ---------------- 03_DESPIECE (masa leída por Python → valores)
    ws = hoja("03_DESPIECE", "ENTRADA_PYTHON (opcional)", "OPERATIVA 04_DESPIECE (valores materializados)", "oper")
    cols = [Col("Producto", "txt", 32), Col("Componente", "txt", 28), Col("Material", "txt", 44), Col("Piezas por unidad", "int", 9),
            Col("Ancho pieza mm", "dec1", 9), Col("Alto pieza mm", "dec1", 9), Col("Longitud pieza m", "dec3", 9),
            Col("Área de chapa m²", "dec6", 11), Col("Masa kg / producto", "dec4", 11), Col("Costo S/ / producto", "soles", 12),
            Col("Planchas 1.20×2.40 equivalentes / producto", "dec4", 12), Col("Costo disponible (1/0)", "bin", 10),
            Col("Tipo de cálculo", "txt", 16), Col("Naturaleza", "txt", 12)]
    escribir_tabla(ws, cols, [[pnom(b["producto_id"]), b["componente"], mat[b["material_id"]]["material"], b["piezas"],
                               b["ancho_mm"], b["alto_mm"], b["longitud_m"], b["area_calc"],
                               None if b["masa_calc"] is None else round(b["masa_calc"], 10),
                               None if b["costo_calc"] is None else round(b["costo_calc"], 10),
                               None if b["planchas_calc"] is None else round(b["planchas_calc"], 10),
                               0 if b["costo_calc"] is None else 1, b["tipo_calculo"], b["naturaleza"]]
                              for b in N["despiece"]], comentarios=False)

    # ---------------- 04_MAQUINAS
    ws = hoja("04_MAQUINAS", "ENTRADA_PYTHON", "OPERATIVA 05_RECURSOS", "oper")
    cols = [Col("Equipo o recurso completo", "txt", 44), Col("Clave interna", "txt", 10), Col("Tipo técnico", "txt", 36),
            Col("Tipo de recurso", "txt", 15), Col("Cantidad de escenario", "int", 9), Col("Costo de uso S//h", "soles", 11),
            Col("Potencia nominal kW (referencia)", "dec2", 10), Col("Factor de carga medio", "dec2", 9),
            Col("Grupo de operarios", "txt", 18), Col("Máquina crítica (1/0)", "bin", 9),
            Col("Operarios habilitados (simulado)", "int", 11), Col("Horas requeridas 2025 (calculado)", "dec1", 12),
            Col("Utilización anual 2025", "pct", 11), Col("Usado en rutas (1/0)", "bin", 9), Col("Naturaleza", "txt", 12)]
    nh = len(N["habilidades"])
    from openpyxl.utils import get_column_letter as GL
    ult_hab = GL(len(N["hab_recursos"]) + 1)
    filas = []
    for i, r in enumerate(N["recursos"], 2):
        filas.append([r["recurso"], r["recurso_id"], r["tipo_tecnico"], r["tipo_recurso"], r["cantidad"], r["costo_uso"],
                      r["potencia_kw"], r["factor_carga"], r["grupo_operarios"], r["critico"],
                      f"=SUMPRODUCT(('06B_HABILIDADES'!$B$1:${ult_hab}$1=A{i})*('06B_HABILIDADES'!$B$2:${ult_hab}${nh + 1}))",
                      f"=SUMIF({R('09_CARGA_ORDEN', 'E', 1152)},A{i},{R('09_CARGA_ORDEN', 'J', 1152)})",
                      f'=IF(N(E{i})=0,"",L{i}/(SUMPRODUCT((YEAR({R("07_CALENDARIO", "A", n["calendario"])})=p_anio_base)*'
                      f'{R("07_CALENDARIO", "D", n["calendario"])})*p_horas_productivas_dia*E{i}))',
                      f"=IF(COUNTIF({R('05_RUTAS', 'E', n['rutas'])},A{i})>0,1,0)", r["naturaleza"]])
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 05_RUTAS
    ws = hoja("05_RUTAS", "ENTRADA_PYTHON", "OPERATIVA 06_RUTAS", "oper")
    cols = [Col("Producto", "txt", 32), Col("Familia", "txt", 26), Col("Paso", "int", 6), Col("Operación", "txt", 30),
            Col("Equipo principal", "txt", 44), Col("Min fijos por lote", "dec2", 9), Col("Minutos por unidad", "dec4", 10),
            Col("Precedencia", "gen", 10), Col("Tratamiento del tiempo", "txt", 12), Col("Min ciclo horno (incluidos)", "dec1", 10),
            Col("Origen de duración", "txt", 30), Col("Equipo válido (1/0)", "bin", 9), Col("Naturaleza", "txt", 12)]
    escribir_tabla(ws, cols, [[pnom(r["producto_id"]), pfam(r["producto_id"]), r["paso"], r["operacion"], rnom(r["recurso_id"]),
                               r["min_fijos"], r["min_unidad"], "Inicio" if r["paso_predecesor"] is None else r["paso_predecesor"],
                               r["tratamiento"], r["min_ciclo_horno"], ORIGEN_T_INV[r["origen_tiempo"]],
                               f"=IF(COUNTIF({R('04_MAQUINAS', 'A', n['recursos'])},E{i})>0,1,0)", r["naturaleza"]]
                              for i, r in enumerate(N["rutas"], 2)], comentarios=False)

    # ---------------- 06_OPERARIOS / 06B_HABILIDADES
    ws = hoja("06_OPERARIOS", "ENTRADA_PYTHON", "OPERATIVA 08_PERSONAL + SIMULACION 02_PERFILES_SIM", "oper")
    cols = [Col("Operario (seudónimo)", "txt", 14), Col("Especialidad (simulada)", "txt", 24), Col("Turnos", "int", 8),
            Col("Horas netas / día", "dec1", 9), Col("Naturaleza", "txt", 30)]
    escribir_tabla(ws, cols, [[p["operario_id"], p["especialidad"], p["turno"], p["horas_netas"],
                               "REAL (cantidad y jornada); especialidad SIMULADA"] for p in N["personal"]], comentarios=False)
    ws = hoja("06B_HABILIDADES", "ENTRADA_PYTHON", "SIMULACION 03_HABILIDADES_SIM", "sim")
    cols = [Col("Operario (seudónimo)", "txt", 14)] + [Col(rnom(r), "bin", 8) for r in N["hab_recursos"]] + \
           [Col("Naturaleza", "txt", 11)]
    escribir_tabla(ws, cols, [[h["operario_id"]] + list(h["valores"]) + [h["naturaleza"]] for h in N["habilidades"]],
                   comentarios=False)
    from openpyxl.styles import Alignment
    for j in range(2, len(N["hab_recursos"]) + 2):
        ws.cell(1, j).alignment = Alignment(text_rotation=90, wrap_text=True, vertical="bottom")
    ws.row_dimensions[1].height = 170

    # ---------------- 07_CALENDARIO
    ws = hoja("07_CALENDARIO", "ENTRADA_PYTHON", "OPERATIVA 09_CALENDARIO", "oper")
    cols = [Col("Fecha", "fecha", 12), Col("Productivo (1/0)", "bin", 10), Col("Feriado nacional Perú", "txt", 30),
            Col("Productivo efectivo (1/0)", "bin", 11), Col("Naturaleza", "txt", 12)]
    escribir_tabla(ws, cols, [[c["fecha"], c["laborable_base"], c["feriado"],
                               f'=IF(AND(p_aplicar_feriados=1,C{i}<>""),0,B{i})', c["naturaleza"]]
                              for i, c in enumerate(N["calendario"], 2)], comentarios=False)

    # ---------------- 08_PEDIDOS
    ws = hoja("08_PEDIDOS", "ENTRADA_PYTHON", "SIMULACION 01_PEDIDOS_SIM", "sim")
    nc, ncal = n["pedidos"], n["calendario"]
    cols = [Col("Pedido", "txt", 8), Col("Producto", "txt", 32), Col("Familia", "txt", 26), Col("Unidades", "int", 8),
            Col("Recepción 2025", "fecha", 11), Col("Liberación 2025", "fecha", 11), Col("Compromiso 2025", "fecha", 11),
            Col("Lote", "int", 6), Col("Clasificación", "txt", 11), Col("Precio familiar S/", "soles", 12),
            Col("Valor pedido S/", "soles", 13), Col("Penalidad fija S/", "soles", 12), Col("Plazo días", "int", 8),
            Col("Horas de proceso (ruta principal)", "dec2", 12), Col("Horas productivas disponibles r→d", "dec1", 12),
            Col("Ruta excede ventana (1/0)", "bin", 10), Col("Fecha en fin de semana (1/0)", "bin", 10),
            Col("Producto válido (1/0)", "bin", 9), Col("Naturaleza", "txt", 11)]
    CA, CD = R("07_CALENDARIO", "A", ncal), R("07_CALENDARIO", "D", ncal)
    filas = []
    for i, p in enumerate(N["pedidos"], 2):
        filas.append([p["pedido_id"], pnom(p["producto_id"]), pfam(p["producto_id"]), p["unidades"], p["recepcion"],
                      p["liberacion"], p["compromiso"], p["lote"], p["naturaleza"],
                      f"=INDEX({R('12_PRECIOS', 'D', n['familias'])},MATCH(C{i},{R('12_PRECIOS', 'A', n['familias'])},0))",
                      f"=J{i}*D{i}",
                      f"=K{i}*INDEX({R('12_PRECIOS', 'E', n['familias'])},MATCH(C{i},{R('12_PRECIOS', 'A', n['familias'])},0))",
                      f"=G{i}-E{i}", f"=SUMIF({R('09_CARGA_ORDEN', 'A', 1152)},A{i},{R('09_CARGA_ORDEN', 'J', 1152)})",
                      f"=SUMPRODUCT(({CA}>=F{i})*({CA}<=G{i})*{CD})*p_horas_productivas_dia",
                      f"=IF(N{i}>O{i}+0.000001,1,0)", f"=IF(OR(WEEKDAY(F{i},2)>5,WEEKDAY(G{i},2)>5),1,0)",
                      f"=IF(COUNTIF({R('01_PRODUCTOS', 'B', n['productos'])},B{i})>0,1,0)", "SIMULADO"])
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 09_CARGA_ORDEN (derivada)
    ws = hoja("09_CARGA_ORDEN", "DERIVADA (fórmulas)", "08_PEDIDOS × 05_RUTAS", "der")
    cols = [Col("Pedido", "txt", 8), Col("Producto", "txt", 32), Col("Paso", "int", 6), Col("Operación", "txt", 28),
            Col("Equipo", "txt", 40), Col("Cantidad", "int", 8), Col("Lote", "int", 6), Col("Min fijo/lote", "dec2", 9),
            Col("Min unidad", "dec4", 9), Col("Horas requeridas", "dec4", 11), Col("Mes", "int", 5),
            Col("Costo máquina S/ (tarifa 04)", "soles", 12), Col("Energía kWh (04)", "dec2", 10), Col("Naturaleza", "txt", 12)]
    RP, RR = "08_PEDIDOS", "05_RUTAS"
    nr = n["rutas"]
    filas, i = [], 2
    for p in N["pedidos"]:
        pasos = sorted([r for r in N["rutas"] if r["producto_id"] == p["producto_id"]], key=lambda r: r["paso"])
        for r in pasos:
            mq = R("04_MAQUINAS", "A", n["recursos"])
            filas.append([p["pedido_id"], pnom(p["producto_id"]), r["paso"], r["operacion"], rnom(r["recurso_id"]),
                          f"=INDEX({R(RP, 'D', nc)},MATCH(A{i},{R(RP, 'A', nc)},0))",
                          f"=INDEX({R(RP, 'H', nc)},MATCH(A{i},{R(RP, 'A', nc)},0))",
                          f"=SUMIFS({R(RR, 'F', nr)},{R(RR, 'A', nr)},B{i},{R(RR, 'C', nr)},C{i})",
                          f"=SUMIFS({R(RR, 'G', nr)},{R(RR, 'A', nr)},B{i},{R(RR, 'C', nr)},C{i})",
                          f"=(ROUNDUP(F{i}/G{i},0)*H{i}+F{i}*I{i})/60",
                          f"=MONTH(INDEX({R(RP, 'F', nc)},MATCH(A{i},{R(RP, 'A', nc)},0)))",
                          f"=J{i}*N(INDEX({R('04_MAQUINAS', 'F', n['recursos'])},MATCH(E{i},{mq},0)))",
                          f"=J{i}*INDEX({R('04_MAQUINAS', 'G', n['recursos'])},MATCH(E{i},{mq},0))"
                          f"*INDEX({R('04_MAQUINAS', 'H', n['recursos'])},MATCH(E{i},{mq},0))", "CALCULADO"])
            i += 1
    assert len(filas) == 1152, len(filas)
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 10_MATERIALES_ORDEN (derivada)
    ws = hoja("10_MATERIALES_ORDEN", "DERIVADA (fórmulas)", "08_PEDIDOS × 03_DESPIECE", "der")
    cols = [Col("Pedido", "txt", 8), Col("Producto", "txt", 32), Col("Componente", "txt", 28), Col("Material", "txt", 44),
            Col("Cantidad pedida", "int", 9), Col("Piezas / unidad", "int", 9), Col("Piezas requeridas", "int", 10),
            Col("Costo componente por producto S/", "soles", 12), Col("Costo pedido S/", "soles", 12),
            Col("Masa de material del pedido kg", "dec2", 11), Col("Planchas equivalentes del pedido", "dec3", 11),
            Col("Costo disponible (1/0)", "bin", 9), Col("Naturaleza", "txt", 12)]
    nb = n["despiece"]
    D3 = "03_DESPIECE"
    sif = lambda col, i: (f"SUMIFS({R(D3, col, nb)},{R(D3, 'A', nb)},B{i},{R(D3, 'B', nb)},C{i})")  # noqa: E731
    filas, i = [], 2
    for p in N["pedidos"]:
        for b in [b for b in N["despiece"] if b["producto_id"] == p["producto_id"]]:
            filas.append([p["pedido_id"], pnom(p["producto_id"]), b["componente"], mat[b["material_id"]]["material"],
                          f"=INDEX({R(RP, 'D', nc)},MATCH(A{i},{R(RP, 'A', nc)},0))", f"={sif('D', i)}", f"=E{i}*F{i}",
                          f"={sif('J', i)}", f"=E{i}*H{i}", f"=E{i}*{sif('I', i)}", f"=E{i}*{sif('K', i)}",
                          f"={sif('L', i)}", "CALCULADO"])
            i += 1
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 11_CAPACIDAD (derivada)
    ws = hoja("11_CAPACIDAD", "DERIVADA (fórmulas)", "07_CALENDARIO × 04_MAQUINAS × 09_CARGA_ORDEN", "der")
    cols = [Col("Recurso", "txt", 44), Col("Mes", "int", 5), Col("Días", "int", 6), Col("Horas/día", "dec1", 8),
            Col("Cantidad", "int", 8), Col("Horas programadas", "dec1", 10), Col("Carga requerida h", "dec2", 11),
            Col("Carga / capacidad", "pct", 10), Col("Naturaleza", "txt", 12)]
    filas, i = [], 2
    for r in N["recursos"]:
        for mes in range(1, 13):
            filas.append([r["recurso"], mes, f"=SUMPRODUCT((MONTH({CA})=B{i})*(YEAR({CA})=p_anio_base)*{CD})",
                          "=p_horas_productivas_dia",
                          f"=INDEX({R('04_MAQUINAS', 'E', n['recursos'])},MATCH(A{i},{R('04_MAQUINAS', 'A', n['recursos'])},0))",
                          f"=C{i}*D{i}*E{i}",
                          f"=SUMIFS({R('09_CARGA_ORDEN', 'J', 1152)},{R('09_CARGA_ORDEN', 'E', 1152)},A{i},"
                          f"{R('09_CARGA_ORDEN', 'K', 1152)},B{i})", f'=IF(F{i}=0,"",G{i}/F{i})', "CALCULADO"])
            i += 1
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 12_PRECIOS
    ws = hoja("12_PRECIOS", "ENTRADA_PYTHON", "OPERATIVA 01_FAMILIAS", "oper")
    cols = [Col("Familia", "txt", 30), Col("Unidades agregadas 2025", "int", 12), Col("Ventas agregadas S/", "soles", 14),
            Col("Precio promedio S/", "soles", 12), Col("Penalidad fija %", "pct", 10), Col("Penalidad diaria %", "pct", 10),
            Col("Naturaleza", "txt", 34)]
    escribir_tabla(ws, cols, [[f["familia"], f["unidades_2025"], f["ventas_2025"], f"=C{i}/B{i}", f["penalidad_fija"],
                               f["penalidad_diaria"], "REAL (ventas); ESTIMADO (penalidades)"]
                              for i, f in enumerate(N["familias"], 2)], comentarios=False)

    # ---------------- 13_ELEGIBILIDAD
    ws = hoja("13_ELEGIBILIDAD", "ENTRADA_PYTHON", "OPERATIVA 07_ELEGIBILIDAD", "oper")
    cols = [Col("Producto", "txt", 32), Col("Paso", "int", 6), Col("Operación", "txt", 28), Col("Máquina elegible", "txt", 44),
            Col("Tipo de asignación", "txt", 34), Col("Elegible 1/0", "bin", 8), Col("Factor de tiempo relativo", "dec2", 9),
            Col("Naturaleza", "txt", 12)]
    escribir_tabla(ws, cols, [[pnom(e["producto_id"]), e["paso"], ruta_op[(e["producto_id"], e["paso"])]["operacion"],
                               rnom(e["recurso_id"]), TIPO_ASIG_INV[e["tipo_asignacion"]], e["elegible"], e["factor_tiempo"],
                               e["naturaleza"]] for e in N["elegibilidad"]], comentarios=False)

    # ---------------- 14_FALLAS_ESC
    ws = hoja("14_FALLAS_ESC", "ENTRADA_PYTHON", "SIMULACION 04_EVENTOS_SIM", "sim")
    cols = [Col("Evento", "txt", 8), Col("Inicio", "fechahora", 19), Col("Fin", "fechahora", 19), Col("Horas de reloj", "dec2", 8),
            Col("Equipo", "txt", 40), Col("Motivo hipotético", "txt", 44), Col("Pedido potencialmente afectado", "txt", 10),
            Col("Paso", "int", 6), Col("Tipo de dato", "txt", 26), Col("Horas de turno afectadas", "dec4", 10),
            Col("Naturaleza", "txt", 11)]
    escribir_tabla(ws, cols, [[e["evento_id"], e["inicio"], f"=B{i}+D{i}/24", e["horas_reloj"], rnom(e["recurso_id"]),
                               e["motivo"], e["pedido_ref"], e["paso_ref"],
                               "SIMULADO" + (" (mantenimiento planificado)" if e["tipo_evento"] != "FALLA" else ""),
                               e["horas_turno"], e["naturaleza"]] for i, e in enumerate(N["eventos"], 2)], comentarios=False)

    # ---------------- 15_RECUPERACION
    ws = hoja("15_RECUPERACION", "ENTRADA_PYTHON", "SIMULACION 06_RECUPERACION_SIM", "sim")
    cols = [Col("Producto", "txt", 34), Col("Paso", "int", 6), Col("Operación", "txt", 26), Col("Política", "txt", 18),
            Col("Conserva piezas", "bin", 9), Col("Reinicia lote", "bin", 9), Col("Repite setup", "bin", 9),
            Col("Piezas a repetir por interrupción", "int", 10), Col("Min de verificación tras reparación", "dec1", 11),
            Col("Naturaleza", "txt", 12)]
    filas = []
    for r in N["recuperacion"]:
        if r["ambito"] == "PRODUCTO_PASO":
            prod_txt, op_txt = pnom(r["producto_id"]), ruta_op[(r["producto_id"], r["paso"])]["operacion"]
        else:
            prod_txt, op_txt = GEN_REC[r["ambito"]]
        filas.append([prod_txt, r["paso"], op_txt, POLITICA[r["politica"]], r["conserva"], r["reinicia"], r["repite_setup"],
                      r["piezas_repetir"], r["min_verificacion"], r["naturaleza"]])
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 16_PROGRAMA_BASE
    ws = hoja("16_PROGRAMA_BASE", "ENTRADA_PYTHON (opcional, comparación E0)", "SIMULACION 09_PROGRAMA_FIFO_REF", "der")
    cols = [Col("Pedido", "txt", 8), Col("Producto", "txt", 32), Col("Paso", "int", 6), Col("Operación", "txt", 28),
            Col("Máquina", "txt", 40), Col("Inicio", "fechahora", 19), Col("Fin", "fechahora", 19), Col("Duración h", "dec6", 10),
            Col("Fecha compromiso", "fecha", 11), Col("Naturaleza", "txt", 12)]
    ped = {p["pedido_id"]: p for p in N["pedidos"]}
    escribir_tabla(ws, cols, [[f["pedido_id"], pnom(ped[f["pedido_id"]]["producto_id"]), f["paso"],
                               ruta_op[(ped[f["pedido_id"]]["producto_id"], f["paso"])]["operacion"], rnom(f["recurso_id"]),
                               f["inicio"], f["fin"], f["duracion_h"], ped[f["pedido_id"]]["compromiso"], f["naturaleza"]]
                              for f in N["fifo"]], comentarios=False)

    # ---------------- 18_ESCENARIOS
    ws = hoja("18_ESCENARIOS", "ENTRADA_PYTHON", "SIMULACION 05_ESCENARIOS_SIM", "sim")
    cols = [Col("Escenario", "txt", 9), Col("Caso", "txt", 46), Col("Momento de falla", "fechahora", 19),
            Col("Duración (h)", "dec2", 9), Col("Evento", "txt", 44), Col("Recuperación", "txt", 16), Col("Semilla", "int", 10),
            Col("Clasificación", "txt", 12), Col("Naturaleza", "txt", 13)]
    ev = {e["evento_id"]: e for e in N["eventos"]}
    filas = []
    for s in N["escenarios"]:
        evs = [x for x in (s["eventos"] or "").split("+") if x]
        mom = min(ev[x]["inicio"] for x in evs) if evs else None
        dur = sum(ev[x]["horas_reloj"] for x in evs) if evs else 0
        filas.append([s["escenario_id"], s["caso"], mom, dur, s["eventos"] or "Ninguna",
                      POLITICA[s["politica"]] if s["politica"] else "FIFO", s["semilla"],
                      "BASE" if not evs else "SIMULADO", s["naturaleza"]])
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 20_COSTOS (derivada)
    ws = hoja("20_COSTOS", "DERIVADA (fórmulas)", "09_CARGA_ORDEN + 10_MATERIALES_ORDEN", "der")
    cols = [Col("Pedido", "txt", 8), Col("Producto", "txt", 32), Col("Unidades", "int", 8), Col("Materiales S/", "soles", 12),
            Col("Mano de obra S/", "soles", 12), Col("Máquinas S/", "soles", 12), Col("Costo total S/", "soles", 12),
            Col("Valor pedido S/", "soles", 12), Col("Margen bruto %", "pct", 10), Col("Costo de materiales completo (1/0)", "bin", 11),
            Col("Naturaleza", "txt", 12)]
    nm = sum(1 for p in N["pedidos"] for b in N["despiece"] if b["producto_id"] == p["producto_id"])
    filas = [[p["pedido_id"], pnom(p["producto_id"]), f"=INDEX({R(RP, 'D', nc)},MATCH(A{i},{R(RP, 'A', nc)},0))",
              f"=SUMIF({R('10_MATERIALES_ORDEN', 'A', nm)},A{i},{R('10_MATERIALES_ORDEN', 'I', nm)})",
              f"=SUMIF({R('09_CARGA_ORDEN', 'A', 1152)},A{i},{R('09_CARGA_ORDEN', 'J', 1152)})*p_costo_hora_hombre",
              f"=SUMIF({R('09_CARGA_ORDEN', 'A', 1152)},A{i},{R('09_CARGA_ORDEN', 'L', 1152)})", f"=SUM(D{i}:F{i})",
              f"=INDEX({R(RP, 'K', nc)},MATCH(A{i},{R(RP, 'A', nc)},0))", f'=IF(H{i}=0,"",(H{i}-G{i})/H{i})',
              f"=IF(COUNTIFS({R('10_MATERIALES_ORDEN', 'A', nm)},A{i},{R('10_MATERIALES_ORDEN', 'L', nm)},0)=0,1,0)",
              "CALCULADO"] for i, p in enumerate(N["pedidos"], 2)]
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 25_PERDIDA_CAPACIDAD_ESC (derivada)
    ws = hoja("25_PERDIDA_CAPACIDAD_ESC", "DERIVADA (fórmulas)", "18_ESCENARIOS × 14_FALLAS_ESC", "der")
    rec_ev = [r for r in N["recursos"] if any(e["recurso_id"] == r["recurso_id"] for e in N["eventos"])]
    cols = [Col("Escenario", "txt", 9), Col("Eventos", "txt", 44)] + \
           [Col(r["recurso"], "dec2", 14) for r in rec_ev] + [Col("Total horas de turno perdidas", "dec2", 12),
                                                              Col("Naturaleza", "txt", 12)]
    ne = n["eventos"]
    filas = []
    for i, s in enumerate([s for s in N["escenarios"] if s["eventos"]], 2):
        fila = [s["escenario_id"], s["eventos"]]
        for j in range(len(rec_ev)):
            col = GL(j + 3)
            fila.append(f'=SUMPRODUCT(ISNUMBER(SEARCH("+"&{R("14_FALLAS_ESC", "A", ne)}&"+","+"&$B{i}&"+"))*'
                        f'({R("14_FALLAS_ESC", "E", ne)}={col}$1)*{R("14_FALLAS_ESC", "J", ne)})')
        fila += [f"=SUM(C{i}:{GL(len(rec_ev) + 2)}{i})", "CALCULADO"]
        filas.append(fila)
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 27_SENSIBILIDAD
    ws = hoja("27_SENSIBILIDAD", "ENTRADA_PYTHON (leída; niveles no usados por el código 1.4.0)", "SIMULACION 07_SENSIBILIDAD_SIM", "sim")
    escribir_tabla(ws, [Col("Factor", "txt", 36), Col("Niveles", "txt", 22), Col("Aplicado a", "txt", 26),
                        Col("Propósito", "txt", 46), Col("Naturaleza", "txt", 14)],
                   [[s["factor"], s["niveles"], s["aplicado_a"], s["proposito"], s["naturaleza"]] for s in N["sensibilidad"]],
                   comentarios=False)

    # ---------------- 24_VALIDACIONES (controles vivos; la fila final la actualiza el programa)
    ws = hoja("24_VALIDACIONES", "CONTROL (fórmulas)", "todas", "ctl")
    P8, C9 = "08_PEDIDOS", "09_CARGA_ORDEN"
    controles = [
        ("Pedidos", f"=COUNTA({R(P8, 'A', nc)})", nc, "Estructura (pedidos SIMULADOS)"),
        ("Unidades pedidas = unidades vendidas 2025 (12_PRECIOS)",
         f"=SUM({R(P8, 'D', nc)})-SUM({R('12_PRECIOS', 'B', n['familias'])})", 0,
         "Los pedidos simulados reparten exactamente las unidades REALES agregadas"),
        ("Operaciones pedido × ruta", f"=COUNTA({R(C9, 'A', 1152)})", 1152, "Instancia de cálculo"),
        ("Horas de proceso 08 = 09", f"=ROUND(SUM({R(P8, 'N', nc)})-SUM({R(C9, 'J', 1152)}),6)", 0, "Consistencia de cargas"),
        ("Carga 11 = 09", f"=ROUND(SUM({R('11_CAPACIDAD', 'G', 12 * n['recursos'])})-SUM({R(C9, 'J', 1152)}),6)", 0,
         "Consistencia de capacidad"),
        ("Días productivos del año base", f"=SUMPRODUCT((YEAR({CA})=p_anio_base)*{CD})", 261,
         "261 sin feriados (interruptor = 0); 249 con feriados"),
        ("Productos con familia inválida", f"=COUNTIF({R('01_PRODUCTOS', 'K', n['productos'])},0)", 0, "Integridad 01 ↔ 12"),
        ("Rutas con equipo inexistente", f"=COUNTIF({R('05_RUTAS', 'L', nr)},0)", 0, "Integridad 05 ↔ 04"),
        ("Pedidos con producto inexistente", f"=COUNTIF({R(P8, 'R', nc)},0)", 0, "Integridad 08 ↔ 01"),
        ("Recursos usados en rutas sin operario habilitado (simulado)",
         f"=COUNTIFS({R('04_MAQUINAS', 'N', n['recursos'])},1,{R('04_MAQUINAS', 'K', n['recursos'])},0)", 0,
         "Matriz de habilidades SIMULADA"),
        ("Pedidos cuya ruta excede la ventana r→d (ALERTA)", f"=SUM({R(P8, 'P', nc)})", "INFO",
         "Infactibles por construcción aun con planta vacía (precedencia lineal, sin traslape)"),
        ("Pedidos con liberación o compromiso en fin de semana (ALERTA)", f"=SUM({R(P8, 'Q', nc)})", "INFO",
         "El modelo los desplaza al día productivo; fechas SIMULADAS"),
        ("Pedidos con costo de materiales incompleto (ALERTA)", f"=COUNTIF({R('20_COSTOS', 'J', nc)},0)", "INFO",
         "Componentes comprados sin precio (MAT-OTR, MAT-MEL18)"),
        ("Margen bruto medio implícito (ALERTA)", f"=AVERAGE({R('20_COSTOS', 'I', nc)})", "INFO",
         "Un margen > 50 % sugiere costos incompletos (sin planilla completa de 25 trabajadores ni gastos)"),
        ("Python reproduce 16_PROGRAMA_BASE", "Ver entrega_v2_0/verificacion", "Sí",
         "Lo actualiza el comando 'reportes --escribir-excel'; la verificación v2 está en E5"),
    ]
    cols = [Col("Control", "txt", 58), Col("Resultado calculado", "gen", 18), Col("Objetivo", "gen", 10),
            Col("Estado", "txt", 10), Col("Interpretación", "txt", 70)]
    filas = []
    for i, (c, f, obj, txt) in enumerate(controles, 2):
        estado = '"INFO"' if obj == "INFO" else f'IF(ROUND(N(B{i})-N(C{i}),6)=0,"OK","REVISAR")'
        if obj == "Sí":
            estado = '"VER E5"'
        filas.append([c, f, obj, f"={estado}", txt])
    escribir_tabla(ws, cols, filas, comentarios=False)

    # ---------------- 00_INDICE
    ws = wb.create_sheet("00_INDICE", 0)
    ws.sheet_properties.tabColor = "7F7F7F"
    cols = [Col("Hoja", "txt", 28), Col("Rol", "txt", 44), Col("Origen", "txt", 58)]
    filas = sorted(indice, key=lambda x: x[0])
    filas += [["", "", ""],
              ["LIBRO GENERADO", "No editar a mano", f"Regenerar con herramientas/construir_bd_v2.py --paso modelo "
                                                     f"(BD v{VERSION_BD})"],
              ["ENTRADA_PYTHON", "Columnas leídas por Optimizacion_MILP_SAGAPI.py 1.4.0", "Valores literales (no fórmulas)"],
              ["DERIVADA", "Tablas de análisis con fórmulas vivas", "No las lee Python"],
              ["Naturaleza", "REAL / REFERENCIA / CALCULADO / ESTIMADO / SIMULADO / CONFIGURACION / PENDIENTE",
               "Definiciones en SAGAPI_BD_OPERATIVA_v2_0.xlsx, hoja 00_CODIGOS"]]
    escribir_tabla(ws, cols, filas, comentarios=False, filtro=False)
    wb._sheets.sort(key=lambda w: w.title)          # orden numérico de hojas; Python las busca por nombre
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "SAGAPI BD MODELO v2.0 (compatible con Python 1.4.0)"
    wb.save(ruta)


# =============================================================================== principal
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--paso", choices=["migrar", "modelo", "todo"], default="todo")
    ap.add_argument("--original", default="Base_Datos_SAGAPI_25_TRABAJADORES_SIMULADOS_v1_2.xlsx")
    ap.add_argument("--salida", default=os.path.join("entrega_v2_0", "base_datos"))
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    if a.paso in ("migrar", "todo"):
        o = leer_original(a.original)
        N = normalizar(o)
        par = {p["parametro_id"]: p["valor_num"] for p in N["parametros"]}
        calc_despiece(N, par)
        comprobar_reproduccion(N)
        construir_operativa(N, os.path.join(a.salida, ARCH_OPER))
        construir_simulacion(N, os.path.join(a.salida, ARCH_SIM))
        print("migración: OPERATIVA y SIMULACION escritas")
    if a.paso in ("modelo", "todo"):
        N = leer_fuentes(a.salida)
        construir_modelo(N, os.path.join(a.salida, ARCH_MOD))
        print("modelo compatible escrito")


if __name__ == "__main__":
    main()
