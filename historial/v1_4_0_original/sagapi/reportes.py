"""Tablas, figuras y hojas de resultados del Excel generadas desde los CSV de las corridas (sin transcripción manual)."""
from __future__ import annotations
import glob
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Patch
import numpy as np
import pandas as pd

GR = ["#1f1f1f", "#555555", "#8c8c8c", "#bdbdbd", "#e0e0e0"]
HATCH = ["", "///", "...", "xx", "\\\\"]
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 9, "axes.edgecolor": "#8c8c8c", "axes.linewidth": 0.6,
                     "axes.spines.top": False, "axes.spines.right": False, "axes.grid": True, "grid.color": "#e6e6e6",
                     "grid.linewidth": 0.6, "axes.axisbelow": True, "savefig.dpi": 200, "savefig.bbox": "tight"})
NOMBRE = {"FIFO+CONV": "FIFO +\nconvencional\n(situación actual)", "FIFO+MILP": "FIFO +\nMILP ante\nfallas",
          "EDD+EDD": "EDD +\nEDD", "MILP+CONV": "MILP +\nconvencional\nante fallas", "MILP+MILP": "MILP +\nMILP\n(propuesta)"}
ORDEN = ["FIFO+CONV", "FIFO+MILP", "EDD+EDD", "MILP+CONV", "MILP+MILP"]
METODOS = ["CONV", "EDD", "MILP-MONO", "MILP-LEX"]


def _f(x, d=1):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "—"
    s = f"{x:,.{d}f}"
    return s.replace(",", " ").replace(".", ",")


def md(df):
    cols = list(df.columns)
    out = ["| " + " | ".join(cols) + " |", "|" + "|".join(["---"] + ["---:"] * (len(cols) - 1)) + "|"]
    for _, r in df.iterrows():
        out.append("| " + " | ".join(str(v) for v in r.values) + " |")
    return "\n".join(out)


# ------------------------------------------------------------------------------------------------ tablas
def tablas(salida):
    T = {}
    e1 = pd.read_csv(os.path.join(salida, "E1_anual.csv"))
    e1["orden"] = e1.combinacion.map({k: i for i, k in enumerate(ORDEN)})
    e1 = e1.sort_values(["escenario", "orden"])
    T["T_ANUAL"] = md(pd.DataFrame({
        "Escenario": e1.escenario.map({"S00": "S00", "S20A": "S20A"}),
        "Combinación": e1.combinacion.str.replace("+", " + ", regex=False),
        "Atrasados": e1.atrasados.astype(int), "Cumpl. (%)": [_f(v) for v in e1.otd_pct],
        "Penalidad (S/)": [_f(v, 0) for v in e1.penalidad], "Tardanza total (días)": [_f(v) for v in e1.tardanza_dias],
        "Tardanza máx. (h)": [_f(v) for v in e1.tard_max_h],
        "MILP (óptimas)": [f"{int(a)} ({int(b)})" for a, b in zip(e1.decisiones_milp, e1.decisiones_optimas)],
        "Binarias máx.": e1.max_binarias.astype(int), "Tiempo (s)": [_f(v) for v in e1.tiempo_total_s]}))
    fam = pd.read_csv(os.path.join(salida, "E1_familias.csv"))
    piv = fam.pivot_table(index="familia", columns=["escenario", "combinacion"], values="atrasados")
    rows = []
    for f in piv.index:
        rows.append({"Familia": f,
                     "S00 FIFO+CONV": int(piv.loc[f, ("S00", "FIFO+CONV")]), "S00 MILP+CONV": int(piv.loc[f, ("S00", "MILP+CONV")]),
                     "S20A FIFO+CONV": int(piv.loc[f, ("S20A", "FIFO+CONV")]), "S20A FIFO+MILP": int(piv.loc[f, ("S20A", "FIFO+MILP")]),
                     "S20A MILP+MILP": int(piv.loc[f, ("S20A", "MILP+MILP")])})
    T["T_FAMILIAS"] = md(pd.DataFrame(rows))
    fp = fam.pivot_table(index="familia", columns=["escenario", "combinacion"], values="penalidad")
    T["T_FAMILIAS_PEN"] = md(pd.DataFrame([{"Familia": f, "S20A FIFO+CONV (S/)": _f(fp.loc[f, ("S20A", "FIFO+CONV")], 0),
                                            "S20A MILP+MILP (S/)": _f(fp.loc[f, ("S20A", "MILP+MILP")], 0)} for f in fp.index]))
    e2 = pd.read_csv(os.path.join(salida, "E2_escenarios.csv"))
    e2["programa_base"] = e2.programa_base.map(lambda b: "FIFO" if str(b).startswith("FIFO") else "MILP")
    res = []
    for (base, met), g in e2.groupby(["programa_base", "metodo"], sort=False):
        res.append({"Programa vigente": base, "Método": met,
                    "Atrasados (total)": int(g.atrasados.sum()), "Penalidad (S/)": _f(g.penalidad.sum(), 0),
                    "Tardanza pond. media (h)": _f(g.tard_pond.mean()), "D medio (h)": _f(g.D.mean()), "R medio": _f(g.R.mean(), 2),
                    "Cambios de secuencia": _f(g.SEQ.mean(), 1),
                    "Óptimos (%)": (_f(100 * (g.estado == "ÓPTIMO").mean()) if met.startswith("MILP") else "—"),
                    "Tiempo medio / máx. (s)": f"{_f(g.tiempo_s.mean(), 2)} / {_f(g.tiempo_s.max(), 1)}"})
    T["T_E2_RESUMEN"] = md(pd.DataFrame(res))
    filas = []
    for (sid, base), g in e2.groupby(["escenario", "programa_base"], sort=False):
        if base != "FIFO":
            continue
        c = g[g.metodo == "CONV"].iloc[0]; e = g[g.metodo == "EDD"].iloc[0]; l = g[g.metodo == "MILP-LEX"].iloc[0]
        filas.append({"Escenario": sid, "Máquina(s)": c.maquinas.replace("Plegadora hidráulica de chapa metálica", "Plegadora")
                      .replace("Guillotina hidráulica para corte de planchas de acero", "Guillotina")
                      .replace("Taladro de banco para perforación de acero", "Taladro")
                      .replace("Máquina de soldadura por arco MIG/MAG", "Soldadora MIG")
                      .replace("Horno de curado de pintura electrostática por lotes", "Horno"),
                      "Pedidos": int(c.pedidos_instancia), "Ops. libres": int(c.ops_libres),
                      "Atrasados CONV / EDD / LEX": f"{int(c.atrasados)} / {int(e.atrasados)} / {int(l.atrasados)}",
                      "Penalidad CONV (S/)": _f(c.penalidad, 0), "Penalidad LEX (S/)": _f(l.penalidad, 0),
                      "D LEX (h)": _f(l.D), "R LEX": int(l.R), "Estado LEX": l.estado, "Tiempo LEX (s)": _f(l.tiempo_s, 1)})
    T["T_E2_DETALLE"] = md(pd.DataFrame(filas))
    # prueba pareada (Wilcoxon) CONV vs LEX por escenario, por programa vigente
    from scipy.stats import wilcoxon
    tests = []
    for base, g in e2.groupby("programa_base", sort=False):
        a = g[g.metodo == "CONV"].set_index("escenario"); b = g[g.metodo == "MILP-LEX"].set_index("escenario")
        for var, nom in (("penalidad", "Penalidad (S/)"), ("atrasados", "Pedidos atrasados"), ("tard_pond", "Tardanza ponderada (h)"),
                         ("D", "Desplazamiento D (h)")):
            dlt = (a[var] - b[var]).dropna()
            nz = dlt[dlt.abs() > 1e-9]
            p = wilcoxon(nz).pvalue if len(nz) >= 6 else np.nan
            tests.append({"Programa vigente": base, "Variable": nom, "Media CONV − LEX": _f(dlt.mean(), 2),
                          "Mediana": _f(dlt.median(), 2), "LEX mejor / igual / peor": f"{int((dlt > 1e-9).sum())} / {int((dlt.abs() <= 1e-9).sum())} / {int((dlt < -1e-9).sum())}",
                          "p (Wilcoxon)": ("< 0,001" if p < 0.001 else _f(p, 3)) if not np.isnan(p) else "n. a."})
    T["T_E2_PRUEBAS"] = md(pd.DataFrame(tests))
    for pol in ("A", "B"):
        fn = os.path.join(salida, f"E2b_recurrentes_{pol}.csv")
        if os.path.exists(fn):
            r = pd.read_csv(fn)
            T[f"T_RECURRENTES_{pol}"] = md(pd.DataFrame({"Reacción": r.reaccion, "Pedidos": r.pedidos.astype(int),
                                                         "Atrasados": r.atrasados.astype(int), "Cumplimiento (%)": [_f(v) for v in r.otd_pct],
                                                         "Penalidad (S/)": [_f(v, 0) for v in r.penalidad],
                                                         "Tardanza ponderada (h)": [_f(v) for v in r.tard_pond],
                                                         "Estados MILP": r.get("estados", pd.Series([""] * len(r))).fillna("")}))
    fn = os.path.join(salida, "E3_pareto.csv")
    if os.path.exists(fn):
        p3 = pd.read_csv(fn)
        T["T_PARETO"] = md(pd.DataFrame({"Escenario": p3.escenario, "Punto": p3.k, "Penalidad (S/)": [_f(v, 0) for v in p3.penalidad],
                                         "Tardanza ponderada (h)": [_f(v) for v in p3.tard_pond], "D (h)": [_f(v) for v in p3.D],
                                         "R": [(_f(v, 0)) for v in p3.R], "f2 (h eq.)": [_f(v) for v in p3.f2],
                                         "Estado": p3.estado.fillna("regla")}))
    fn = os.path.join(salida, "E4_sensibilidad_eventos.csv")
    if os.path.exists(fn):
        s = pd.read_csv(fn)
        rows = []
        for (fac, niv, met), g in s.groupby(["factor", "nivel", "metodo"], sort=False):
            rows.append({"Factor": fac, "Nivel": niv, "Método": met, "Escenarios": len(g), "Atrasados (total)": int(g.atrasados.sum()),
                         "Penalidad total (S/)": _f(g.penalidad.sum(), 0), "Pen. diaria con tope (S/)": _f(g.pen_diaria_tope.sum(), 0),
                         "Tardanza pond. media (h)": _f(g.tard_pond.mean()), "D medio (h)": _f(g.D.mean()), "R medio": _f(g.R.mean(), 2),
                         "Óptimos (%)": _f(100 * (g.estado == "ÓPTIMO").mean()) if met != "CONV" else "—"})
        T["T_SENSIBILIDAD"] = md(pd.DataFrame(rows))
    fn = os.path.join(salida, "E4_sensibilidad_anual.csv")
    if os.path.exists(fn):
        s = pd.read_csv(fn)
        T["T_SENS_ANUAL"] = md(pd.DataFrame({"Factor": s.factor, "Escenario": s.escenario, "Combinación": s.combinacion,
                                             "Atrasados": s.atrasados.astype(int), "Penalidad fija (S/)": [_f(v, 0) for v in s.penalidad],
                                             "Pen. diaria con tope (S/)": [_f(v, 0) for v in s.pen_diaria_tope],
                                             "Tardanza (días háb.)": [_f(v) for v in s.tardanza_dias]}))
    # resiliencia, sostenibilidad y participación humana por método (E2)
    rr = []
    for (base, met), g in e2.groupby(["programa_base", "metodo"], sort=False):
        rr.append({"Programa vigente": base, "Método": met,
                   "Degradación media (S/)": _f((g.penalidad - g.penalidad_ref).mean(), 0),
                   "Pasan a atrasados": int((g.atrasados - g.atrasados_ref).clip(lower=0).sum()),
                   "Retraso propagado (h)": _f(g.retraso_propagado_h.mean()),
                   "Ops. movidas > 1 día": _f(g.ops_movidas_mas_1dia.mean()),
                   "Reasignaciones": _f(g.R.mean(), 2),
                   "Energía reproceso (kWh)": _f(g.energia_reproceso_kwh.sum(), 1),
                   "Material reproceso (kg)": _f(g.material_reproceso_kg.sum(), 1)})
    T["T_RESILIENCIA"] = md(pd.DataFrame(rr))
    fn = os.path.join(salida, "E5_montecarlo.csv")
    if os.path.exists(fn):
        mc = pd.read_csv(fn)
        from scipy.stats import wilcoxon
        rows = []
        rng = np.random.default_rng(2025)
        for pol, g in mc.groupby("politica"):
            a = g[g.metodo == "CONV"].set_index("replica"); b = g[g.metodo == "MILP-LEX"].set_index("replica")
            e = g[g.metodo == "EDD"].set_index("replica")
            for var, nom in (("penalidad", "Penalidad (S/)"), ("atrasados", "Pedidos atrasados"), ("tard_pond", "Tardanza ponderada (h)"),
                             ("D", "Desplazamiento D (h)")):
                dlt = (a[var] - b[var]).dropna()
                nz = dlt[dlt.abs() > 1e-9]
                p = wilcoxon(nz).pvalue if len(nz) >= 6 else np.nan
                boots = [rng.choice(dlt.values, len(dlt)).mean() for _ in range(10000)]
                rows.append({"Política": pol, "Variable": nom, "Réplicas": len(dlt),
                             "Media CONV": _f(a[var].mean(), 1), "Media EDD": _f(e[var].mean(), 1), "Media MILP-LEX": _f(b[var].mean(), 1),
                             "Diferencia media CONV − LEX": _f(dlt.mean(), 1),
                             "IC 95 % bootstrap": f"[{_f(np.percentile(boots, 2.5), 1)}; {_f(np.percentile(boots, 97.5), 1)}]",
                             "LEX mejor / igual / peor": f"{int((dlt > 1e-9).sum())} / {int((dlt.abs() <= 1e-9).sum())} / {int((dlt < -1e-9).sum())}",
                             "p (Wilcoxon)": ("< 0,001" if p < 0.001 else _f(p, 3)) if not np.isnan(p) else "n. a."})
        T["T_MC"] = md(pd.DataFrame(rows))
        lx = mc[mc.metodo == "MILP-LEX"]
        T["T_MC_COMPUTO"] = (f"{len(lx)} resoluciones MILP-LEX; óptimas {int((lx.estado == 'ÓPTIMO').sum())}; "
                             f"tiempo medio {_f(lx.tiempo_s.mean(), 2)} s, máximo {_f(lx.tiempo_s.max(), 1)} s; "
                             f"binarias medias {_f(lx.modelo_binarias.mean(), 0)}, máximo {_f(lx.modelo_binarias.max(), 0)}; "
                             f"violaciones {int(mc.violaciones.sum())}")
    ver = os.path.join(salida, "verificacion_enumeracion.csv")
    if os.path.exists(ver):
        v = pd.read_csv(ver)
        T["T_VERIF"] = (f"{len(v)} microinstancias; f1 coincide en {int((abs(v.f1_enum-v.f1_milp)<1e-4).sum())}, "
                        f"f1b en {int((abs(v.f1b_enum-v.f1b_milp)<1e-3).sum())}, f2 comprobable en {int(v.f2_enum.notna().sum())} "
                        f"y coincide en {int(((abs(v.f2_enum-v.f2_milp)<1e-3*np.maximum(1,v.f2_enum.abs())) & v.f2_enum.notna()).sum())}; "
                        f"violaciones {int(v.violaciones.sum())}")
    with open(os.path.join(salida, "tablas_capitulo4.md"), "w", encoding="utf-8") as fh:
        for k, v in T.items():
            fh.write(f"<!-- {k} -->\n{v}\n\n")
    return T


# ------------------------------------------------------------------------------------------------ figuras
def fig_anual(salida, fig_dir):
    e1 = pd.read_csv(os.path.join(salida, "E1_anual.csv"))
    s = e1[e1.escenario == "S20A"].set_index("combinacion").reindex(ORDEN)
    s0 = e1[e1.escenario == "S00"].set_index("combinacion")
    fig, axs = plt.subplots(1, 2, figsize=(10.5, 3.9))
    x = np.arange(len(ORDEN))
    for ax, var, tit, fmt in ((axs[0], "atrasados", "Pedidos atrasados (de 96)", "{:.0f}"),
                              (axs[1], "penalidad", "Penalidad económica (miles de S/)", "{:.1f}")):
        vals = s[var].values / (1000 if var == "penalidad" else 1)
        bars = ax.bar(x, vals, width=0.62, color=[GR[1] if c != "MILP+MILP" else GR[0] for c in ORDEN], edgecolor="white", linewidth=1.5)
        for b, v in zip(bars, vals):
            ax.text(b.get_x() + b.get_width() / 2, v, fmt.format(v), ha="center", va="bottom", fontsize=8, color=GR[0])
        ref = [s0[var].get(c, np.nan) for c in ORDEN]
        ax.scatter(x, np.array(ref, dtype=float) / (1000 if var == "penalidad" else 1), marker="_", s=380, color=GR[2], zorder=3,
                   label="mismo método sin fallas (S00)")
        ax.set_xticks(x); ax.set_xticklabels([NOMBRE[c] for c in ORDEN], fontsize=7.2)
        ax.set_title(tit, fontsize=9, loc="left")
    axs[0].legend(frameon=False, fontsize=7, loc="upper right")
    fig.suptitle("Año 2025 con las 12 fallas de la plegadora (S20A): programación + reacción", fontsize=9.5, x=0.01, ha="left")
    fig.savefig(os.path.join(fig_dir, "fig_anual.png")); plt.close(fig)


def fig_familias(salida, fig_dir):
    fam = pd.read_csv(os.path.join(salida, "E1_familias.csv"))
    fam = fam[fam.escenario == "S20A"]
    fams = ["Racks industriales", "Góndolas comerciales", "Lockers metálicos", "Otros productos metalmecánicos"]
    combos = ["FIFO+CONV", "FIFO+MILP", "MILP+MILP"]
    fig, ax = plt.subplots(figsize=(7.5, 3.0))
    w = 0.26
    for k, c in enumerate(combos):
        v = [fam[(fam.familia == f) & (fam.combinacion == c)].atrasados.sum() for f in fams]
        b = ax.bar(np.arange(4) + (k - 1) * w, v, width=w - 0.02, color=GR[k], hatch=HATCH[k], edgecolor="white", label=" ".join(NOMBRE[c].split()))
        for bb, vv in zip(b, v):
            ax.text(bb.get_x() + bb.get_width() / 2, vv, f"{vv:.0f}", ha="center", va="bottom", fontsize=7)
    ax.set_xticks(np.arange(4)); ax.set_xticklabels(["Racks", "Góndolas", "Lockers", "Otros"])
    ax.set_ylabel("Pedidos atrasados (de 24 por familia)")
    ax.legend(frameon=False, fontsize=7, ncol=3, loc="upper center", bbox_to_anchor=(0.5, 1.18))
    fig.savefig(os.path.join(fig_dir, "fig_familias.png")); plt.close(fig)


def fig_escenarios(salida, fig_dir):
    e2 = pd.read_csv(os.path.join(salida, "E2_escenarios.csv"))
    e2 = e2[e2.programa_base.str.startswith("FIFO")]
    sids = list(dict.fromkeys(e2.escenario))
    c = e2[e2.metodo == "CONV"].set_index("escenario").reindex(sids)
    l = e2[e2.metodo == "MILP-LEX"].set_index("escenario").reindex(sids)
    fig, ax = plt.subplots(figsize=(7.2, 7.5))
    y = np.arange(len(sids))[::-1]
    for yy, a, b in zip(y, c.penalidad / 1000, l.penalidad / 1000):
        ax.plot([b, a], [yy, yy], color=GR[3], lw=2, zorder=1)
    ax.scatter(c.penalidad / 1000, y, s=40, color="white", edgecolor=GR[0], lw=1.2, zorder=2, label="Reacción convencional")
    ax.scatter(l.penalidad / 1000, y, s=40, color=GR[0], zorder=3, label="MILP lexicográfico")
    ax.set_yticks(y); ax.set_yticklabels(sids, fontsize=7)
    ax.set_xlabel("Penalidad de los pedidos activos al fallar (miles de S/)")
    ax.legend(frameon=False, fontsize=8, loc="lower right")
    ax.set_title("Escenarios de falla sobre el programa FIFO vigente", fontsize=9, loc="left")
    fig.savefig(os.path.join(fig_dir, "fig_escenarios.png")); plt.close(fig)


def fig_pareto(salida, fig_dir):
    fn = os.path.join(salida, "E3_pareto.csv")
    if not os.path.exists(fn):
        return
    p = pd.read_csv(fn)
    esc = list(dict.fromkeys(p.escenario))
    fig, axs = plt.subplots(1, len(esc), figsize=(2.9 * len(esc), 2.9))
    axs = np.atleast_1d(axs)
    for ax, sid in zip(axs, esc):
        g = p[(p.escenario == sid) & (p.k != "CONV")].dropna(subset=["f2"])
        g = g.sort_values("tard_pond")
        ax.plot(g.tard_pond, g.f2, "-o", color=GR[0], ms=4, lw=1.2, label="Frente ε (penalidad = f1*)")
        cv = p[(p.escenario == sid) & (p.k == "CONV")].iloc[0]
        pen = g.penalidad.iloc[0] if len(g) else np.nan
        ax.set_title(f"{sid}: penalidad S/ {pen:,.0f}".replace(",", " "), fontsize=8.5, loc="left")
        ax.text(0.02, 0.03, f"Convencional: penalidad S/ {cv.penalidad:,.0f};\ntardanza pond. {cv.tard_pond:.1f} h; f2 = {cv.f2:.1f} h".replace(",", " "),
                transform=ax.transAxes, fontsize=6.3, color=GR[1], va="bottom")
        ax.set_xlabel("Tardanza ponderada (h)", fontsize=7.5)
    axs[0].set_ylabel("Inestabilidad f2 (h equivalentes)", fontsize=7.5)
    fig.tight_layout(); fig.savefig(os.path.join(fig_dir, "fig_pareto.png")); plt.close(fig)


def fig_sensibilidad(salida, fig_dir):
    fn = os.path.join(salida, "E4_sensibilidad_eventos.csv")
    if not os.path.exists(fn):
        return
    s = pd.read_csv(fn)
    s = s[s.factor == "Multiplicador de duración"].copy()
    s["nivel"] = s["nivel"].astype(float)
    g = s.groupby(["nivel", "metodo"]).agg(pen=("penalidad", "sum"), atr=("atrasados", "sum")).reset_index()
    fig, axs = plt.subplots(1, 2, figsize=(8, 2.8))
    for k, met in enumerate(["CONV", "MILP-LEX"]):
        h = g[g.metodo == met]
        axs[0].plot(h.nivel, h.atr, "-o", color=GR[k * 2], ms=4, label="Convencional" if met == "CONV" else "MILP lexicográfico",
                    markerfacecolor="white" if met == "CONV" else GR[0])
        axs[1].plot(h.nivel, h.pen / 1000, "-o", color=GR[k * 2], ms=4, markerfacecolor="white" if met == "CONV" else GR[0])
    for ax, t in zip(axs, ["Pedidos atrasados (suma F01–F12)", "Penalidad (miles de S/, suma F01–F12)"]):
        ax.set_xscale("log", base=2); ax.set_xticks([0.5, 1, 2, 4]); ax.set_xticklabels(["×0,5", "×1", "×2", "×4"])
        ax.set_xlabel("Multiplicador de la duración de la avería"); ax.set_title(t, fontsize=8.5, loc="left")
    axs[0].legend(frameon=False, fontsize=7)
    fig.tight_layout(); fig.savefig(os.path.join(fig_dir, "fig_sensibilidad.png")); plt.close(fig)


def fig_montecarlo(salida, fig_dir):
    fn = os.path.join(salida, "E5_montecarlo.csv")
    if not os.path.exists(fn):
        return
    mc = pd.read_csv(fn)
    fig, axs = plt.subplots(1, 2, figsize=(8, 2.9), sharey=False)
    for ax, pol in zip(axs, ["A", "B"]):
        g = mc[mc.politica == pol]
        a = g[g.metodo == "CONV"].set_index("replica").penalidad / 1000
        b = g[g.metodo == "MILP-LEX"].set_index("replica").penalidad / 1000
        lim = max(a.max(), b.max()) * 1.05
        ax.plot([0, lim], [0, lim], color=GR[3], lw=0.8, ls="--")
        ax.scatter(a, b.reindex(a.index), s=14, color=GR[0], alpha=0.8)
        ax.set_xlabel("Penalidad reacción convencional (miles de S/)")
        ax.set_title(f"Política {pol}: {'reanudar' if pol == 'A' else 'reiniciar lote'}", fontsize=8.5, loc="left")
        ax.set_xlim(0, lim); ax.set_ylim(0, lim)
    axs[0].set_ylabel("Penalidad MILP lexicográfico (miles de S/)")
    fig.tight_layout(); fig.savefig(os.path.join(fig_dir, "fig_montecarlo.png")); plt.close(fig)


CORTOS = [("Segundo taladro", "Taladro 2 (alternativo)"), ("Segunda estación soldadura", "Soldadora 2 (alternativa)"),
          ("Guillotina", "Guillotina"), ("Sierra", "Sierra de cinta"), ("Taladro", "Taladro"), ("Plegadora", "Plegadora"),
          ("soldadura", "Soldadora MIG/MAG"), ("desbarbado", "Desbarbado"), ("preparación superficial", "Preparación superficial"),
          ("Pistola", "Aplicación de pintura"), ("Horno", "Horno de curado"), ("enfriamiento", "Enfriamiento"),
          ("armado", "Armado"), ("inspección", "Inspección"), ("embalaje", "Embalaje")]


def nombre_corto(m):
    for clave, corto in CORTOS:
        if clave.lower() in m.lower():
            return corto
    return m[:24]


def gantt_caso(E, sub, base, sols, sid, ruta):
    """Gantt de las máquinas afectadas: programa vigente, reacción convencional y MILP-LEX."""
    inst, cal = E.I, E.d.calendario
    jobs = sub.jobs
    prog_base = {o: base[o] for j in jobs for o in inst.job_ops[j]}
    paneles = [("Programa vigente (FIFO) antes de la falla", prog_base)]
    for met, tit in (("CONV", "Reacción convencional (desplazamiento a la derecha)"), ("MILP-LEX", "Reprogramación MILP lexicográfica")):
        full = dict(sub.fixed)
        for o, (m, S) in sols[met].items():
            full[o] = (m, S, S + sub.pR[(o, m)])
        paneles.append((tit, full))
    maqs = sorted({m for _, p in paneles for (m, S, C) in p.values()}, key=lambda m: inst.machines.index(m))
    corto = {m: nombre_corto(m) for m in maqs}
    t0 = sub.t0 - 16
    t1 = max(C for _, p in paneles for (m, S, C) in p.values()) + 2
    jcol = {j: GR[k % 4] for k, j in enumerate(jobs)}
    fig, axs = plt.subplots(len(paneles), 1, figsize=(10, 2.0 + 0.32 * len(maqs) * len(paneles)), sharex=True)
    for ax, (tit, p) in zip(axs, paneles):
        for o, (m, S, C) in p.items():
            if C < t0:
                continue
            y = maqs.index(m)
            j = inst.ops[o]["job"]
            ax.barh(y, C - max(S, t0), left=max(S, t0), height=0.62, color=jcol[j], edgecolor="white", lw=0.8,
                    hatch="///" if (o in sub.interrupted) else "")
            if C - max(S, t0) > 3:
                ax.text(max(S, t0) + (C - max(S, t0)) / 2, y, inst.jobs[j]["id"], ha="center", va="center", fontsize=5.5,
                        color="white" if jcol[j] in (GR[0], GR[1]) else GR[0])
        for (m, a, b) in sub.all_windows:
            if m in maqs:
                ax.barh(maqs.index(m), b - a, left=a, height=0.9, color="none", edgecolor="black", hatch="xxx", lw=0.8)
        ax.axvline(sub.t0, color="black", lw=0.8, ls="--")
        for j in jobs:
            ax.axvline(inst.jobs[j]["d"], color=GR[2], lw=0.5, ls=":")
        ax.set_yticks(range(len(maqs))); ax.set_yticklabels([corto[m] for m in maqs], fontsize=6.5)
        ax.invert_yaxis(); ax.set_title(tit, fontsize=8.5, loc="left"); ax.grid(axis="y", visible=False)
        ax.set_xlim(t0, t1)
    ticks = np.arange(np.ceil(t0 / 8) * 8, t1, 8)
    axs[-1].set_xticks(ticks)
    axs[-1].set_xticklabels([cal.a_fecha(t).strftime("%d/%m") for t in ticks], fontsize=6.5)
    axs[-1].set_xlabel("Días hábiles (eje en horas productivas; línea discontinua = instante de la falla; punteadas = fechas comprometidas)")
    leg = [Patch(facecolor="none", edgecolor="black", hatch="xxx", label="Indisponibilidad"),
           Patch(facecolor=GR[2], hatch="///", edgecolor="white", label="Operación interrumpida")]
    axs[0].legend(handles=leg, frameon=False, fontsize=6.5, loc="upper right")
    fig.tight_layout(); fig.savefig(ruta); plt.close(fig)


def tabla_programa(E, sub, sols, ruta):
    inst, cal = E.I, E.d.calendario
    filas = []
    for met, sol in sols.items():
        for o, (m, S) in sorted(sol.items(), key=lambda kv: kv[1][1]):
            filas.append(dict(metodo=met, pedido=inst.jobs[inst.ops[o]["job"]]["id"], paso=inst.ops[o]["paso"],
                              operacion=inst.ops[o]["operacion"], maquina=m, inicio=cal.a_fecha(S), fin=cal.a_fecha(S + sub.pR[(o, m)], "fin"),
                              horas=sub.pR[(o, m)], inicio_vigente=(cal.a_fecha(sub.S0[o]) if o in sub.S0 else None),
                              maquina_vigente=sub.m0.get(o), reasignada=int(o in sub.m0 and sub.m0[o] != m)))
    pd.DataFrame(filas).to_csv(ruta, index=False)


# ------------------------------------------------------------------------------------------------ Excel
def escribir_excel(E, salida, img_dir=None):
    import openpyxl
    from openpyxl.chart import BarChart, Reference, ScatterChart, Series
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.workbook.properties import CalcProperties
    wb = openpyxl.load_workbook(E.archivo)
    ENC = PatternFill("solid", fgColor="4B4B4B"); AMA = PatternFill("solid", fgColor="FFF2B6")
    fB = Font(name="Carlito", bold=True, color="FFFFFF")

    def volcar(nombre, df, amarillas=(), nota=""):
        if nombre in wb.sheetnames:
            del wb[nombre]
        ws = wb.create_sheet(nombre)
        ws["A1"] = nota
        ws["A1"].font = Font(name="Carlito", bold=True)
        for j, c in enumerate(df.columns, 1):
            cell = ws.cell(2, j, c); cell.fill = ENC; cell.font = fB; cell.alignment = Alignment(wrap_text=True)
            ws.column_dimensions[openpyxl.utils.get_column_letter(j)].width = max(11, min(28, len(str(c)) + 2))
        for i, row in enumerate(df.itertuples(index=False), 3):
            for j, v in enumerate(row, 1):
                if isinstance(v, float) and np.isnan(v):
                    v = None
                if hasattr(v, "to_pydatetime"):
                    v = v.to_pydatetime()
                c = ws.cell(i, j, v)
                if df.columns[j - 1] in amarillas:
                    c.fill = AMA
                if isinstance(v, float):
                    c.number_format = "#,##0.00"
        ws.freeze_panes = "A3"
        return ws

    e1 = pd.read_csv(os.path.join(salida, "E1_anual.csv"))
    cols = ["escenario", "combinacion", "pedidos", "atrasados", "otd_pct", "penalidad", "pen_diaria_tope", "tardanza_h", "tardanza_dias",
            "tard_max_h", "tard_pond", "energia_kwh", "conflictos_personal_h", "decisiones_milp", "decisiones_optimas", "max_binarias", "tiempo_total_s"]
    ws = volcar("28_RESULTADOS_ANUALES", e1[cols], ("atrasados", "otd_pct", "penalidad"),
                "RESULTADO CALCULADO — año 2025 por eventos (datos sintéticos/estimados de la base); no son registros de SAGAPI")
    n = len(e1) + 2
    ch = BarChart(); ch.type = "col"; ch.title = "Pedidos atrasados por combinación y escenario"; ch.y_axis.title = "pedidos"
    ch.add_data(Reference(ws, min_col=4, min_row=2, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, max_col=2, min_row=3, max_row=n))
    ch.height, ch.width = 8, 18
    ch.series[0].graphicalProperties.solidFill = "4B4B4B"
    ws.add_chart(ch, "B" + str(n + 3))
    fam = pd.read_csv(os.path.join(salida, "E1_familias.csv"))
    w2 = wb["28_RESULTADOS_ANUALES"]
    start = n + 22
    w2.cell(start, 1, "Pedidos atrasados y penalidad por familia").font = Font(name="Carlito", bold=True)
    for j, c in enumerate(fam.columns, 1):
        cell = w2.cell(start + 1, j, c); cell.fill = ENC; cell.font = fB
    for i, row in enumerate(fam.itertuples(index=False), start + 2):
        for j, v in enumerate(row, 1):
            w2.cell(i, j, v)
    e2 = pd.read_csv(os.path.join(salida, "E2_escenarios.csv"))
    cols2 = ["escenario", "programa_base", "caso", "maquinas", "politica", "metodo", "fecha_falla", "pedidos_instancia", "ops_libres", "interrumpidas",
             "atrasados", "otd_pct", "penalidad", "tard_pond", "tardanza_h", "D", "R", "SEQ", "f2", "retraso_propagado_h",
             "energia_reproceso_kwh", "material_reproceso_kg", "estado", "tiempo_s", "violaciones", "modelo_binarias", "modelo_restricciones"]
    cols2 = [c for c in cols2 if c in e2.columns]
    volcar("29_RESULTADOS_ESCENARIOS", e2[cols2], ("atrasados", "penalidad", "D", "R"),
           "RESULTADO CALCULADO — escenarios de falla individuales y simultáneos (cada fila verificada: columna violaciones)")
    if os.path.exists(os.path.join(salida, "E3_pareto.csv")):
        p3 = pd.read_csv(os.path.join(salida, "E3_pareto.csv"))
        ws3 = volcar("30_FRENTE_PARETO", p3, ("tard_pond", "f2"), "RESULTADO CALCULADO — ε-restricción (penalidad fija en su óptimo)")
        sc = ScatterChart(); sc.title = "Frente tardanza ponderada – inestabilidad"; sc.x_axis.title = "tardanza ponderada (h)"; sc.y_axis.title = "f2 (h eq.)"
        cx = list(p3.columns).index("tard_pond") + 1; cy = list(p3.columns).index("f2") + 1
        for sid in dict.fromkeys(p3.escenario):
            idx = [i for i, (s, k) in enumerate(zip(p3.escenario, p3.k)) if s == sid and k != "CONV"]
            if not idx:
                continue
            r0, r1 = idx[0] + 3, idx[-1] + 3
            s = Series(Reference(ws3, min_col=cy, min_row=r0, max_row=r1), Reference(ws3, min_col=cx, min_row=r0, max_row=r1), title=sid)
            s.marker.symbol = "circle"
            sc.series.append(s)
        sc.height, sc.width = 9, 16
        ws3.add_chart(sc, "N3")
    sens = [f for f in ("E4_sensibilidad_eventos.csv", "E4_sensibilidad_anual.csv") if os.path.exists(os.path.join(salida, f))]
    if sens:
        s = pd.concat([pd.read_csv(os.path.join(salida, f)) for f in sens], ignore_index=True)
        keep = [c for c in ["factor", "nivel", "evento", "escenario", "combinacion", "metodo", "atrasados", "penalidad", "pen_diaria_tope",
                            "tard_pond", "D", "R", "estado", "tiempo_s", "violaciones"] if c in s.columns]
        volcar("31_SENSIBILIDAD_RES", s[keep], ("atrasados", "penalidad"), "RESULTADO CALCULADO — análisis de sensibilidad (27_SENSIBILIDAD)")
    if os.path.exists(os.path.join(salida, "E5_montecarlo.csv")):
        mc = pd.read_csv(os.path.join(salida, "E5_montecarlo.csv"))
        keep = [c for c in ["replica", "politica", "operacion", "fecha_falla", "duracion_h", "metodo", "pedidos_instancia", "ops_libres",
                            "atrasados", "penalidad", "tard_pond", "D", "R", "estado", "tiempo_s", "violaciones"] if c in mc.columns]
        volcar("33_ROBUSTEZ_MC", mc[keep], ("atrasados", "penalidad"),
               "RESULTADO CALCULADO — réplicas Monte Carlo de fallas de la plegadora (semilla en 19_PARAMETROS)")
    if "32_GRAFICOS" in wb.sheetnames:
        del wb["32_GRAFICOS"]
    wg = wb.create_sheet("32_GRAFICOS")
    wg["A1"] = "Gráficos generados por Optimizacion_MILP_SAGAPI.py (copias de la carpeta resultados/figuras)"
    wg["A1"].font = Font(name="Carlito", bold=True)
    fila = 3
    for png in ["fig_anual.png", "fig_familias.png", "fig_escenarios.png", "fig_pareto.png", "fig_sensibilidad.png", "fig_montecarlo.png"]:
        ruta = os.path.join(salida, "figuras", png)
        if os.path.exists(ruta):
            im = XLImage(ruta); im.width, im.height = int(im.width * 0.33), int(im.height * 0.33)
            wg.add_image(im, f"B{fila}"); fila += int(im.height / 19) + 3
    # openpyxl no conserva imágenes al reabrir: se reinsertan los esquemas de 01B
    img_dir = img_dir or os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "herramientas", "img")
    if "01B_IMAGENES" in wb.sheetnames and os.path.isdir(img_dir):
        wi = wb["01B_IMAGENES"]
        wi._images = []                      # evita duplicar los esquemas si el libro ya los conserva
        for r in range(2, wi.max_row + 1):
            v = wi.cell(r, 1).value
            if not v:
                continue
            import re
            m = re.search(r"!B(\d+)", str(v))
            prod = wb["01_PRODUCTOS"].cell(int(m.group(1)), 2).value if m else v
            ruta = os.path.join(img_dir, f"{prod}.png")
            if os.path.exists(ruta):
                im = XLImage(ruta); im.width, im.height = int(im.width * 0.62), int(im.height * 0.62)
                wi.add_image(im, f"G{r}")
    v = wb["24_VALIDACIONES"]
    for r in range(2, v.max_row + 1):
        if v.cell(r, 1).value == "Python reproduce 16_PROGRAMA_BASE":
            rep = json.load(open(os.path.join(salida, "E0_reproduccion_fifo.json")))
            v.cell(r, 2).value = "Sí" if rep["reproduce"] else "No"
    if "33_ROBUSTEZ_MC" in wb.sheetnames:
        k = wb.sheetnames.index("33_ROBUSTEZ_MC")
        wb._sheets.append(wb._sheets.pop(k))
    for w in wb.worksheets:
        if w.title[:2].isdigit() and int(w.title[:2]) >= 28:
            w.sheet_properties.tabColor = "FFD966"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.save(E.archivo)


def generar(E, salida, escribir_excel_flag=False, escribir_excel=False):
    fig_dir = os.path.join(salida, "figuras")
    os.makedirs(fig_dir, exist_ok=True)
    T = tablas(salida)
    for f in (fig_anual, fig_familias, fig_escenarios, fig_pareto, fig_sensibilidad, fig_montecarlo):
        try:
            f(salida, fig_dir)
        except FileNotFoundError:
            pass
    if escribir_excel or escribir_excel_flag:
        escribir_excel_fn = globals()["escribir_excel"]
        escribir_excel_fn(E, salida)
    return T
