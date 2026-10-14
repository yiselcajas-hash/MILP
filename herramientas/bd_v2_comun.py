"""Utilidades comunes de la base de datos SAGAPI v2.0 (estilos, escritura de tablas, códigos).

Lo usan ``construir_bd_v2.py`` (migración y exportación) y ``verificar_bd_v2.py``.
No forma parte del paquete ``sagapi`` ni modifica el modelo MILP.
"""
from __future__ import annotations

import datetime as dt

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

VERSION_BD = "2.0.0"

# ------------------------------------------------------------------ naturaleza del dato
NATURALEZA = {
    "REAL": "Dato declarado o respaldado por la empresa (p. ej. 25 trabajadores, 1 turno, ventas agregadas).",
    "REFERENCIA": "Dato público verificable ajeno a la empresa (ley, constante física, UIT, feriados oficiales).",
    "CALCULADO": "Resultado de una operación reproducible a partir de otros datos de la base.",
    "ESTIMADO": "Aproximación justificada (heredada, de catálogo o de ingeniería) que requiere validación en planta.",
    "SIMULADO": "Dato artificial creado para probar el modelo. No es historial ni registro de SAGAPI.",
    "CONFIGURACION": "Parámetro de ejecución del modelo o del experimento; no describe la planta.",
    "PENDIENTE": "Dato necesario que no existe en la base; se deja vacío a propósito (ver pendientes).",
}
LISTA_NATURALEZA = '"' + ",".join(NATURALEZA) + '"'

# ------------------------------------------------------------------ estilos
FUENTE = "Arial"
F_HDR = Font(name=FUENTE, bold=True, color="FFFFFF", size=10)
P_HDR = PatternFill("solid", fgColor="1F3864")
F_IN = Font(name=FUENTE, size=10, color="0000FF")       # dato introducido (literal)
F_FX = Font(name=FUENTE, size=10, color="000000")       # fórmula
F_TXT = Font(name=FUENTE, size=10, color="000000")
BORDE = Border(bottom=Side(style="thin", color="BFBFBF"))
TAB = {"dicc": "7F7F7F", "oper": "2E7D32", "reg": "1565C0", "sim": "E65100", "der": "595959", "ctl": "B8860B"}

FMT = {
    "txt": "@", "int": "0", "dec1": "0.0", "dec2": "#,##0.00", "dec3": "0.000", "dec4": "0.0000", "dec6": "0.000000",
    "soles": '"S/" #,##0.00', "pct": "0.0%", "fecha": "yyyy-mm-dd", "fechahora": "yyyy-mm-dd hh:mm:ss",
    "hora": "hh:mm", "bin": "0", "gen": "General",
}


class Col:
    """Definición de columna: nombre, formato, ancho, descripción y metadatos del diccionario."""

    def __init__(self, nombre, fmt="gen", ancho=None, desc="", unidad="", oblig=1, nat="", python="", nulo=""):
        self.nombre, self.fmt, self.ancho, self.desc = nombre, fmt, ancho, desc
        self.unidad, self.oblig, self.nat, self.python, self.nulo = unidad, oblig, nat, python, nulo


def escribir_tabla(ws, cols, filas, tab=None, congelar="A2", filtro=True, comentarios=True):
    """Escribe encabezado + filas. Literales en azul, fórmulas en negro. Devuelve nº de la última fila."""
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c.nombre)
        cell.font, cell.fill = F_HDR, P_HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[L(j)].width = c.ancho or max(10, min(40, len(str(c.nombre)) + 3))
        if comentarios and (c.desc or c.unidad or c.nat):
            partes = [c.desc] if c.desc else []
            if c.unidad:
                partes.append(f"Unidad: {c.unidad}")
            if c.nat:
                partes.append(f"Naturaleza: {c.nat}")
            partes.append("Obligatorio" if c.oblig else "Opcional (vacío válido)")
            cm = Comment("\n".join(partes), "SAGAPI BD v2")
            cm.width, cm.height = 260, 110
            cell.comment = cm
    ws.row_dimensions[1].height = 32
    for i, fila in enumerate(filas, 2):
        for j, (c, v) in enumerate(zip(cols, fila), 1):
            if v is None or v == "":
                continue
            cell = ws.cell(i, j, v)
            es_fx = isinstance(v, str) and v.startswith("=")
            cell.font = F_FX if es_fx else F_IN
            if c.fmt != "gen":
                cell.number_format = FMT[c.fmt]
    for j, c in enumerate(cols, 1):          # formato también en celdas vacías (plantillas de captura)
        if c.fmt != "gen" and not filas:
            for i in range(2, 202):
                ws.cell(i, j).number_format = FMT[c.fmt]
    if congelar:
        ws.freeze_panes = congelar
    ultima = max(1, len(filas) + 1)
    if filtro and cols:
        ws.auto_filter.ref = f"A1:{L(len(cols))}{ultima}"
    if tab:
        ws.sheet_properties.tabColor = TAB[tab]
    return ultima


def validacion_lista(ws, formula, rango, mensaje=""):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorTitle="Valor no válido", error=mensaje or "Elija un valor de la lista.")
    ws.add_data_validation(dv)
    dv.add(rango)
    return dv


def validacion_entero(ws, rango, minimo=0, maximo=None, permitir_vacio=True):
    if maximo is None:
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=str(minimo), allow_blank=permitir_vacio)
    else:
        dv = DataValidation(type="whole", operator="between", formula1=str(minimo), formula2=str(maximo),
                            allow_blank=permitir_vacio)
    dv.showErrorMessage = True
    dv.error = "Debe ser un número entero en el rango permitido."
    ws.add_data_validation(dv)
    dv.add(rango)
    return dv


def validacion_decimal(ws, rango, minimo=0.0, permitir_vacio=True):
    dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(minimo), allow_blank=permitir_vacio)
    dv.showErrorMessage = True
    dv.error = "Debe ser un número mayor o igual que el mínimo."
    ws.add_data_validation(dv)
    dv.add(rango)
    return dv


def nombre_definido(wb, nombre, hoja, celda):
    col, fila = celda.rstrip("0123456789"), celda[len(celda.rstrip("0123456789")):]
    ref = f"'{hoja}'!${col}${fila}"
    wb.defined_names[nombre] = DefinedName(nombre, attr_text=ref)


def hoja_diccionario(wb, especificacion, titulo_libro):
    """especificacion: lista (hoja, [Col...]) -> hoja 00_DICCIONARIO."""
    ws = wb.create_sheet("00_DICCIONARIO", 0)
    cols = [Col("hoja", ancho=26), Col("columna", ancho=34), Col("tipo", ancho=12), Col("unidad", ancho=14),
            Col("obligatorio_1_0", "int", 12), Col("naturaleza", ancho=16), Col("descripcion", ancho=70),
            Col("usado_por_python", ancho=40)]
    tipo = {"txt": "texto", "int": "entero", "bin": "binario 0/1", "fecha": "fecha", "fechahora": "fecha-hora",
            "hora": "hora", "soles": "moneda S/", "pct": "fracción", "gen": "texto/número"}
    filas = []
    for hoja, columnas in especificacion:
        for c in columnas:
            desc = ("Fórmula: " + c.desc[1:].strip()) if c.desc.startswith("=") else c.desc
            filas.append([hoja, c.nombre, tipo.get(c.fmt, "decimal"), c.unidad, c.oblig, c.nat, desc, c.python])
    escribir_tabla(ws, cols, filas, tab="dicc", comentarios=False)
    ws2 = wb.create_sheet("00_CODIGOS", 1)
    escribir_tabla(ws2, [Col("codigo_naturaleza", ancho=18), Col("definicion", ancho=110)],
                   [[k, v] for k, v in NATURALEZA.items()] +
                   [["", ""], ["LIBRO", titulo_libro], ["VERSION_BD", VERSION_BD],
                    ["COLORES", "Texto azul = dato introducido; texto negro = fórmula; encabezado azul oscuro."],
                    ["FECHAS", "AAAA-MM-DD (fecha) y AAAA-MM-DD hh:mm:ss (fecha-hora)."]],
                   tab="dicc", filtro=False, comentarios=False)
    return ws


def a_fecha(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime.combine(v, dt.time())
    return v


def num_o_none(v):
    """Número o None. Textos de relleno ('No aplica', 'Según especificación') se convierten en vacío."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def entero(v):
    return None if v is None else int(round(float(v)))
