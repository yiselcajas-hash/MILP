#!/usr/bin/env python3
"""Verificación independiente de la base SAGAPI v2.0 (Entregable 5).

Comprueba, sobre los libros RECALCULADOS (valores en caché de LibreOffice/Excel):
  V01 claves únicas y no nulas            V06 integridad referencial entre hojas
  V02 campos obligatorios                 V07 consistencia de tiempos, rutas, capacidades y escenarios
  V03 duplicados                          V08 compatibilidad con la lectura de Python (sagapi.datos.leer)
  V04 formatos numéricos y fechas         V09 parámetros del MILP idénticos al original (r, d, w, P, p_om, ...)
  V05 unidades y fórmulas (recalculo)     V10 comparación con el Excel original (cambios intencionados)
Opcional --pruebas-python: ejecuta 'validar', las 50 pruebas unitarias, la enumeración (40) y los
pilotos rodantes S08A/S13A con el libro original y con el MODELO v2.0, y compara sus salidas.

Uso:  python herramientas/verificar_bd_v2.py [--pruebas-python] [--salida entrega_v2_0/verificacion]
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter, defaultdict

import openpyxl

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
ORIGINAL = os.path.join(RAIZ, "Base_Datos_SAGAPI_25_TRABAJADORES_SIMULADOS_v1_2.xlsx")
BD = os.path.join(RAIZ, "entrega_v2_0", "base_datos")
OPER = os.path.join(BD, "SAGAPI_BD_OPERATIVA_v2_0.xlsx")
SIM = os.path.join(BD, "SAGAPI_BD_SIMULACION_v2_0.xlsx")
MOD = os.path.join(BD, "SAGAPI_BD_MODELO_v2_0.xlsx")
TOL = 1e-6


class Informe:
    def __init__(self):
        self.res = []

    def add(self, cod, nombre, ok, detalle="", nivel="ERROR", n=None):
        self.res.append(dict(codigo=cod, prueba=nombre, resultado="PASA" if ok else ("ALERTA" if nivel == "ALERTA" else "FALLA"),
                             detalle=detalle if isinstance(detalle, str) else json.dumps(detalle, ensure_ascii=False,
                                                                                         default=str)[:1500],
                             casos=n))
        marca = "OK " if ok else ("!! " if nivel == "ALERTA" else "XX ")
        print(f"{marca}{cod} {nombre}" + (f" — {self.res[-1]['detalle'][:160]}" if not ok or detalle else ""))


def leer_libro(ruta, valores=True):
    wb = openpyxl.load_workbook(ruta, data_only=valores)
    out = {}
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        enc = list(filas[0])
        out[ws.title] = [dict(zip(enc, f)) for f in filas[1:] if f and any(v is not None for v in f)]
        out["_enc_" + ws.title] = enc
    return out


def es_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(v)


def cerca(a, b, tol=TOL):
    if a is None or b is None:
        return a is None and b is None
    return abs(float(a) - float(b)) <= tol * max(1.0, abs(float(b)))


# ============================================================================ V01–V07 estructura
def verificar_estructura(inf, O, S, M):
    LIB = {"OPERATIVA": O, "SIMULACION": S}
    claves = {
        ("OPERATIVA", "01_FAMILIAS"): ["familia_id"], ("OPERATIVA", "02_PRODUCTOS"): ["producto_id"],
        ("OPERATIVA", "03_MATERIALES"): ["material_id"], ("OPERATIVA", "04_DESPIECE"): ["bom_id"],
        ("OPERATIVA", "05_RECURSOS"): ["recurso_id"], ("OPERATIVA", "06_RUTAS"): ["producto_id", "paso"],
        ("OPERATIVA", "07_ELEGIBILIDAD"): ["producto_id", "paso", "recurso_id"], ("OPERATIVA", "08_PERSONAL"): ["operario_id"],
        ("OPERATIVA", "09_CALENDARIO"): ["fecha"], ("OPERATIVA", "10_CAPACIDAD_MENSUAL"): ["recurso_id", "mes"],
        ("OPERATIVA", "11_PARAMETROS"): ["parametro_id"], ("SIMULACION", "01_PEDIDOS_SIM"): ["pedido_id"],
        ("SIMULACION", "02_PERFILES_SIM"): ["operario_id"], ("SIMULACION", "03_HABILIDADES_SIM"): ["operario_id"],
        ("SIMULACION", "04_EVENTOS_SIM"): ["evento_id"], ("SIMULACION", "05_ESCENARIOS_SIM"): ["escenario_id"],
        ("SIMULACION", "06_RECUPERACION_SIM"): ["ambito", "producto_id", "paso", "politica"],
        ("SIMULACION", "08_CONFIG_MODELO"): ["parametro_id"], ("SIMULACION", "09_PROGRAMA_FIFO_REF"): ["pedido_id", "paso"],
    }
    # V01
    malos = []
    for (libro, hoja), ks in claves.items():
        L = LIB[libro]
        vistos = Counter(tuple(f.get(k) for k in ks) for f in L[hoja])
        nulos = [f for f in L[hoja] if any(f.get(k) is None for k in ks) and not (hoja == "06_RECUPERACION_SIM"
                                                                              and f.get("ambito") != "PRODUCTO_PASO")]
        dup = [k for k, c in vistos.items() if c > 1]
        if dup or nulos:
            malos.append(f"{libro}/{hoja}: {len(dup)} duplicadas, {len(nulos)} nulas")
    for L, hoja, col in [(O, "02_PRODUCTOS", "producto"), (O, "05_RECURSOS", "recurso"), (O, "03_MATERIALES", "material")]:
        c = Counter(f[col] for f in L[hoja])
        if any(v > 1 for v in c.values()):
            malos.append(f"{hoja}.{col} con nombres repetidos")
    inf.add("V01", "Claves primarias únicas y no nulas (19 tablas)", not malos, "; ".join(malos) or "19 tablas sin duplicados ni nulos",
            n=sum(len(LIB[lb][h]) for (lb, h) in claves))
    # V02 obligatorios según diccionario
    faltan = []
    for L, nombre in ((O, "OPERATIVA"), (S, "SIMULACION")):
        for d in L["00_DICCIONARIO"]:
            hoja, col, ob = d["hoja"], d["columna"], d["obligatorio_1_0"]
            if not ob or hoja not in L or not L[hoja]:
                continue
            nul = [i for i, f in enumerate(L[hoja], 2) if f.get(col) in (None, "")]
            if hoja == "06_RECUPERACION_SIM" and col in ("producto_id", "paso"):
                nul = [i for i, f in enumerate(L[hoja], 2) if f["ambito"] == "PRODUCTO_PASO" and f.get(col) is None]
            if nul:
                faltan.append(f"{nombre}/{hoja}.{col}: filas {nul[:6]}")
    inf.add("V02", "Campos obligatorios completos (según 00_DICCIONARIO)", not faltan, "; ".join(faltan) or
            "Todos los campos marcados obligatorios tienen valor")
    # V03 duplicados de fila completa
    dups = []
    for L, nombre in ((O, "OPERATIVA"), (S, "SIMULACION"), (M, "MODELO")):
        for hoja, filas in L.items():
            if hoja.startswith("_") or hoja.startswith("00_") or hoja == "99_LISTAS":
                continue
            c = Counter(tuple(sorted((k, str(v)) for k, v in f.items())) for f in filas)
            n = sum(v - 1 for v in c.values() if v > 1)
            if n:
                dups.append(f"{nombre}/{hoja}: {n}")
    inf.add("V03", "Sin filas duplicadas", not dups, "; ".join(dups) or "Ninguna hoja tiene filas idénticas")
    # V04 formatos
    tipos_malos = []
    for L, nombre in ((O, "OPERATIVA"), (S, "SIMULACION")):
        for d in L["00_DICCIONARIO"]:
            hoja, col, tipo = d["hoja"], d["columna"], d["tipo"]
            if hoja not in L:
                continue
            for i, f in enumerate(L[hoja], 2):
                v = f.get(col)
                if v is None or v == "":
                    continue
                if tipo in ("entero", "binario 0/1", "decimal", "moneda S/", "fracción"):
                    if not es_num(v):
                        tipos_malos.append(f"{nombre}/{hoja}.{col} fila {i}: {v!r}")
                    elif tipo == "entero" and abs(v - round(v)) > 1e-9:
                        tipos_malos.append(f"{nombre}/{hoja}.{col} fila {i}: no entero {v}")
                    elif tipo == "binario 0/1" and v not in (0, 1):
                        tipos_malos.append(f"{nombre}/{hoja}.{col} fila {i}: no binario {v}")
                    elif v < 0:
                        tipos_malos.append(f"{nombre}/{hoja}.{col} fila {i}: negativo {v}")
                elif tipo in ("fecha", "fecha-hora") and not isinstance(v, (dt.datetime, dt.date)):
                    tipos_malos.append(f"{nombre}/{hoja}.{col} fila {i}: no es fecha {v!r}")
                elif tipo == "hora" and not isinstance(v, (dt.time, dt.datetime)):
                    tipos_malos.append(f"{nombre}/{hoja}.{col} fila {i}: no es hora {v!r}")
    err = [f"{h}!{c}" for L in (O, S, M) for h, filas in L.items() if not h.startswith("_")
           for i, fl in enumerate(filas) for c, v in fl.items() if isinstance(v, str) and v.startswith("#")
           and v.rstrip("!?0/").upper() in ("#DIV", "#N/A", "#NAME", "#REF", "#VALUE", "#NUM", "#NULL")]
    inf.add("V04", "Tipos de dato (sin texto en campos numéricos; enteros, binarios, fechas)", not tipos_malos and not err,
            "; ".join(tipos_malos[:10] + err[:10]) or "Todos los valores respetan el tipo declarado en el diccionario")

    # V05 fórmulas y unidades: recálculo independiente
    par = {p["parametro_id"]: p["valor_num"] for p in O["11_PARAMETROS"]}
    mat = {m["material_id"]: m for m in O["03_MATERIALES"]}
    dif = []
    for b in O["04_DESPIECE"]:
        m = mat[b["material_id"]]
        area = b["area_chapa_m2_pieza"]
        if b["ancho_pieza_mm"] is not None and b["alto_pieza_mm"] is not None and not cerca(area, b["ancho_pieza_mm"] * b["alto_pieza_mm"] / 1e6):
            dif.append((b["bom_id"], "area"))
        if b["tipo_calculo"] == "CHAPA_AREA":
            masa = b["piezas_por_unidad"] * area * m["espesor_mm"] / 1000 * par["densidad_acero"]
            costo = masa * m["costo_unitario_soles"] / par["aprovechamiento_chapa"]
        elif b["tipo_calculo"] == "PERFIL_LONGITUD":
            masa = b["piezas_por_unidad"] * b["longitud_pieza_m"] * par["kg_por_m_perfil"]
            costo = masa * m["costo_unitario_soles"] / par["aprovechamiento_perfil"]
        elif b["tipo_calculo"] == "UNIDAD":
            masa, costo = None, b["piezas_por_unidad"] * m["costo_unitario_soles"]
        else:
            masa, costo = None, None
        if not cerca(b["masa_kg_producto"] if b["masa_kg_producto"] != "" else None, masa):
            dif.append((b["bom_id"], "masa", b["masa_kg_producto"], masa))
        if not cerca(b["costo_soles_producto"] if b["costo_soles_producto"] != "" else None, costo):
            dif.append((b["bom_id"], "costo", b["costo_soles_producto"], costo))
        pl = b["piezas_por_unidad"] * area / par["area_plancha_m2"] if area else None
        if not cerca(b["planchas_equiv_producto"] if b["planchas_equiv_producto"] != "" else None, pl):
            dif.append((b["bom_id"], "planchas"))
    fam = {f["familia_id"]: f for f in O["01_FAMILIAS"]}
    for f in O["01_FAMILIAS"]:
        if not cerca(f["precio_promedio_soles"], f["ventas_2025_soles"] / f["unidades_vendidas_2025"]):
            dif.append((f["familia_id"], "precio"))
    for c in O["09_CALENDARIO"]:
        esp = 0 if (par["aplicar_feriados"] == 1 and c["feriado_nacional"]) else c["laborable_base"]
        if c["productivo_efectivo"] != esp or not cerca(c["horas_productivas"], esp * par["horas_productivas_dia"]) \
                or c["dia_semana"] != c["fecha"].isoweekday():
            dif.append((str(c["fecha"].date()), "calendario"))
    dias = defaultdict(int)
    for c in O["09_CALENDARIO"]:
        if c["fecha"].year == par["anio_base"]:
            dias[c["fecha"].month] += c["productivo_efectivo"]
    cant = {r["recurso_id"]: r["cantidad"] for r in O["05_RECURSOS"]}
    for c in O["10_CAPACIDAD_MENSUAL"]:
        if c["dias_productivos"] != dias[c["mes"]] or not cerca(c["horas_disponibles"], dias[c["mes"]] * 8 * cant[c["recurso_id"]]):
            dif.append((c["recurso_id"], c["mes"], "capacidad"))
    for p in O["08_PERSONAL"]:
        if not cerca(p["control_horario_h"], p["horas_netas_dia"]):
            dif.append((p["operario_id"], "jornada"))
    for p in S["01_PEDIDOS_SIM"]:
        if p["plazo_dias"] != (p["fecha_compromiso"] - p["fecha_recepcion"]).days:
            dif.append((p["pedido_id"], "plazo"))
    for e in S["04_EVENTOS_SIM"]:
        if abs((e["fin"] - e["inicio"]).total_seconds() / 3600 - e["horas_reloj"]) > 1e-3:
            dif.append((e["evento_id"], "fin"))
    # MODELO: cargas, costos y capacidad
    from sagapi import datos as D
    ruta = {(r["Producto"], r["Paso"]): r for r in M["05_RUTAS"]}
    ped = {p["Pedido"]: p for p in M["08_PEDIDOS"]}
    carga = defaultdict(float)
    for k in M["09_CARGA_ORDEN"]:
        r, p = ruta[(k["Producto"], k["Paso"])], ped[k["Pedido"]]
        h = D.duracion_h(dict(setup_min=r["Min fijos por lote"], min_unidad=r["Minutos por unidad"]), p["Unidades"], p["Lote"])
        if not cerca(k["Horas requeridas"], h):
            dif.append((k["Pedido"], k["Paso"], "horas 09"))
        carga[(k["Equipo"], p["Liberación 2025"].month)] += h
    for c in M["11_CAPACIDAD"]:
        if not cerca(c["Carga requerida h"], carga.get((c["Recurso"], c["Mes"]), 0.0)):
            dif.append((c["Recurso"], c["Mes"], "carga 11"))
    prec = {f["Familia"]: f for f in M["12_PRECIOS"]}
    for p in M["08_PEDIDOS"]:
        v = prec[p["Familia"]]["Ventas agregadas S/"] / prec[p["Familia"]]["Unidades agregadas 2025"] * p["Unidades"]
        if not cerca(p["Valor pedido S/"], v) or not cerca(p["Penalidad fija S/"], v * prec[p["Familia"]]["Penalidad fija %"]):
            dif.append((p["Pedido"], "valor/penalidad"))
    inf.add("V05", "Fórmulas y unidades: recálculo independiente en Python = valores de la hoja", not dif,
            dif[:12] or "Área, masa, costo, planchas, precio, calendario, capacidad, jornada, plazo, fin de evento, "
            "horas por operación (= sagapi.datos.duracion_h), carga mensual y valor/penalidad coinciden (tol. 1e-6)")

    # V06 integridad referencial
    ids = lambda L, h, c: {f[c] for f in L[h]}  # noqa: E731
    P, R, Mt, W = ids(O, "02_PRODUCTOS", "producto_id"), ids(O, "05_RECURSOS", "recurso_id"), ids(O, "03_MATERIALES", "material_id"), \
        ids(O, "08_PERSONAL", "operario_id")
    rutas = {(r["producto_id"], r["paso"]) for r in O["06_RUTAS"]}
    ev = ids(S, "04_EVENTOS_SIM", "evento_id")
    huerf = []
    chk = [("02_PRODUCTOS.familia_id", [f["familia_id"] for f in O["02_PRODUCTOS"]], set(fam)),
           ("04_DESPIECE.producto_id", [f["producto_id"] for f in O["04_DESPIECE"]], P),
           ("04_DESPIECE.material_id", [f["material_id"] for f in O["04_DESPIECE"]], Mt),
           ("06_RUTAS.producto_id", [f["producto_id"] for f in O["06_RUTAS"]], P),
           ("06_RUTAS.recurso_id", [f["recurso_id"] for f in O["06_RUTAS"]], R),
           ("07_ELEGIBILIDAD.(producto,paso)", [(f["producto_id"], f["paso"]) for f in O["07_ELEGIBILIDAD"]], rutas),
           ("07_ELEGIBILIDAD.recurso_id", [f["recurso_id"] for f in O["07_ELEGIBILIDAD"]], R),
           ("10_CAPACIDAD_MENSUAL.recurso_id", [f["recurso_id"] for f in O["10_CAPACIDAD_MENSUAL"]], R),
           ("SIM 01_PEDIDOS_SIM.producto_id", [f["producto_id"] for f in S["01_PEDIDOS_SIM"]], P),
           ("SIM 02_PERFILES_SIM.operario_id", [f["operario_id"] for f in S["02_PERFILES_SIM"]], W),
           ("SIM 03_HABILIDADES_SIM.operario_id", [f["operario_id"] for f in S["03_HABILIDADES_SIM"]], W),
           ("SIM 03_HABILIDADES_SIM.columnas", [c for c in S["_enc_03_HABILIDADES_SIM"][1:] if c != "naturaleza"], R),
           ("SIM 04_EVENTOS_SIM.recurso_id", [f["recurso_id"] for f in S["04_EVENTOS_SIM"]], R),
           ("SIM 04_EVENTOS_SIM.pedido_ref", [f["pedido_ref"] for f in S["04_EVENTOS_SIM"]], ids(S, "01_PEDIDOS_SIM", "pedido_id")),
           ("SIM 05_ESCENARIOS_SIM.eventos", [x for f in S["05_ESCENARIOS_SIM"] for x in (f["eventos"] or "").split("+") if x], ev),
           ("SIM 06_RECUPERACION_SIM.(producto,paso)", [(f["producto_id"], f["paso"]) for f in S["06_RECUPERACION_SIM"]
                                                        if f["ambito"] == "PRODUCTO_PASO"], rutas),
           ("SIM 09_PROGRAMA_FIFO_REF.recurso_id", [f["recurso_id"] for f in S["09_PROGRAMA_FIFO_REF"]], R),
           ("SIM 99_LISTAS = OPERATIVA", [f["producto_id"] for f in S["99_LISTAS"] if f["producto_id"]] +
            [f["recurso_id"] for f in S["99_LISTAS"] if f["recurso_id"]], P | R)]
    for nombre, vals, dom in chk:
        malos = sorted({str(v) for v in vals if v not in dom})
        if malos:
            huerf.append(f"{nombre}: {malos[:5]}")
    pedidos_prod = {p["producto_id"] for p in S["01_PEDIDOS_SIM"]}
    sin_ruta = [p for p in pedidos_prod if not any(r[0] == p for r in rutas)]
    sin_bom = [p for p in pedidos_prod if not any(b["producto_id"] == p for b in O["04_DESPIECE"])]
    if sin_ruta or sin_bom:
        huerf.append(f"productos con pedido sin ruta {sin_ruta} o sin despiece {sin_bom}")
    inf.add("V06", "Integridad referencial (18 relaciones + producto con pedido ⇒ ruta y despiece)", not huerf,
            "; ".join(huerf) or "Sin códigos huérfanos")

    # V07 consistencia operativa
    alertas, errores = [], []
    por_prod = defaultdict(list)
    for r in O["06_RUTAS"]:
        por_prod[r["producto_id"]].append(r)
    for p, pasos in por_prod.items():
        pasos.sort(key=lambda r: r["paso"])
        if [r["paso"] for r in pasos] != list(range(1, len(pasos) + 1)):
            errores.append(f"{p}: pasos no correlativos")
        for k, r in enumerate(pasos):
            if r["paso_predecesor"] != (None if k == 0 else pasos[k - 1]["paso"]):
                errores.append(f"{p} paso {r['paso']}: predecesor incoherente")
            if r["min_fijos_lote"] + r["min_por_unidad"] <= 0:
                errores.append(f"{p} paso {r['paso']}: tiempo nulo")
            if r["tratamiento_tiempo"] == "POR_LOTE" and r["min_por_unidad"] != 0:
                errores.append(f"{p} paso {r['paso']}: POR_LOTE con tiempo unitario")
    elig = defaultdict(list)
    for e in O["07_ELEGIBILIDAD"]:
        elig[(e["producto_id"], e["paso"])].append(e)
    for r in O["06_RUTAS"]:
        pr = [e for e in elig[(r["producto_id"], r["paso"])] if e["tipo_asignacion"] == "PRINCIPAL"]
        if len(pr) != 1 or pr[0]["recurso_id"] != r["recurso_id"] or pr[0]["factor_tiempo"] != 1:
            errores.append(f"{r['producto_id']}/{r['paso']}: principal de elegibilidad ≠ ruta")
    for r in O["05_RECURSOS"]:
        usado = any(x["recurso_id"] == r["recurso_id"] for x in O["06_RUTAS"] + O["07_ELEGIBILIDAD"])
        if usado and (r["cantidad"] or 0) < 1:
            errores.append(f"{r['recurso_id']}: usado en rutas con cantidad 0")
    hab = S["03_HABILIDADES_SIM"]
    for rid in {x["recurso_id"] for x in O["06_RUTAS"] if x["producto_id"] in pedidos_prod}:
        if sum(h.get(rid) or 0 for h in hab) < 1:
            errores.append(f"{rid}: sin trabajador habilitado (simulado)")
    if len(O["08_PERSONAL"]) != par["trabajadores_produccion"] or len(hab) != par["trabajadores_produccion"]:
        errores.append("dotación ≠ 25")
    fechas = [c["fecha"] for c in O["09_CALENDARIO"]]
    if any((b - a).days != 1 for a, b in zip(fechas, fechas[1:])):
        errores.append("calendario con días faltantes")
    for p in S["01_PEDIDOS_SIM"]:
        if not (p["fecha_recepcion"] <= p["fecha_liberacion"] <= p["fecha_compromiso"]):
            errores.append(f"{p['pedido_id']}: fechas desordenadas")
        if not (fechas[0] <= p["fecha_liberacion"] and p["fecha_compromiso"] <= fechas[-1]):
            errores.append(f"{p['pedido_id']}: fuera del calendario")
    eventos = {e["evento_id"]: e for e in S["04_EVENTOS_SIM"]}
    fifo = {(f["pedido_id"], f["paso"]): f for f in S["09_PROGRAMA_FIFO_REF"]}
    ped_prod = {p["pedido_id"]: p["producto_id"] for p in S["01_PEDIDOS_SIM"]}
    ruta_rec = {(r["producto_id"], r["paso"]): r["recurso_id"] for r in O["06_RUTAS"]}
    for e in eventos.values():
        if not (fechas[0] <= e["inicio"] and e["fin"] <= fechas[-1] + dt.timedelta(days=1)):
            errores.append(f"{e['evento_id']}: fuera del calendario")
        if ruta_rec.get((ped_prod[e["pedido_ref"]], e["paso_ref"])) != e["recurso_id"]:
            errores.append(f"{e['evento_id']}: el paso de referencia no se procesa en el recurso del evento")
        f = fifo[(e["pedido_ref"], e["paso_ref"])]
        if not (f["inicio"] <= e["inicio"] < f["fin"]):
            alertas.append(f"{e['evento_id']}: no cae dentro de la operación FIFO de referencia")
    for s in S["05_ESCENARIOS_SIM"]:
        evs = [eventos[x] for x in (s["eventos"] or "").split("+") if x]
        for i, a in enumerate(evs):
            for b in evs[i + 1:]:
                if a["recurso_id"] == b["recurso_id"] and a["inicio"] < b["fin"] and b["inicio"] < a["fin"]:
                    errores.append(f"{s['escenario_id']}: {a['evento_id']} y {b['evento_id']} se solapan en {a['recurso_id']}")
        if evs and s["politica_recuperacion"] not in ("A", "B"):
            errores.append(f"{s['escenario_id']}: política inválida")
    sobrecarga = [(c["Recurso"], c["Mes"], round(c["Carga / capacidad"], 3)) for c in M["11_CAPACIDAD"]
                  if isinstance(c["Carga / capacidad"], (int, float)) and c["Carga / capacidad"] > 1]
    if sobrecarga:
        errores.append(f"sobrecarga mensual {sobrecarga[:5]}")
    inf.add("V07", "Consistencia de rutas, elegibilidad, dotación, calendario, pedidos, eventos y carga ≤ capacidad",
            not errores, "; ".join(errores[:12]) or "Rutas correlativas con predecesor lineal; principal = ruta; ningún "
            "recurso usado con cantidad 0; 25 trabajadores; calendario continuo; fechas ordenadas; eventos dentro del "
            "calendario y sin solapes por escenario; carga ≤ capacidad en los 204 recurso-mes")
    rutas_exceden = [p["Pedido"] for p in M["08_PEDIDOS"] if p["Ruta excede ventana (1/0)"] == 1]
    finde = sum(p["Fecha en fin de semana (1/0)"] for p in M["08_PEDIDOS"])
    inf.add("V07b", "Alertas de realismo (no bloquean; requieren validación empresarial)", False,
            dict(pedidos_ruta_excede_ventana=rutas_exceden, pedidos_con_fecha_fin_de_semana=finde,
                 eventos_fuera_de_operacion_fifo=alertas,
                 carga_max_mensual={c["Recurso"]: round(max(x["Carga / capacidad"] for x in M["11_CAPACIDAD"]
                                                            if x["Recurso"] == c["Recurso"] and
                                                            isinstance(x["Carga / capacidad"], (int, float))), 3)
                                    for c in M["11_CAPACIDAD"] if c["Mes"] == 1 and isinstance(c["Carga / capacidad"], (int, float))
                                    and c["Carga / capacidad"] > 0}),
            nivel="ALERTA")


# ============================================================================ V08–V10 compatibilidad y comparación
def _resumen_datos(d):
    from sagapi import instancia
    I = instancia.construir(d)
    cal = d.calendario
    out = dict(
        productos={k: (v["familia"], v["lote"]) for k, v in d.productos.items()},
        maquinas={k: (m.cantidad, m.costo_h, m.potencia_kw, m.factor_carga, m.grupo_operario, m.clave)
                  for k, m in d.maquinas.items()},
        rutas={p: [(r["paso"], r["operacion"], r["maquina"], round(r["setup_min"], 9), round(r["min_unidad"], 9), r["trato"])
                   for r in rs]
               for p, rs in d.rutas.items()},
        elegibilidad={f"{k[0]}|{k[1]}": sorted(v) for k, v in d.elegibilidad.items()},
        calendario=(cal.tramos, cal.h_dia, [str(x) for x in cal.fechas], cal.prod, cal.total),
        fallas={e: (str(f["inicio"].replace(microsecond=round(f["inicio"].microsecond, -3))), f["horas_reloj"], f["maquina"],
                    f["pedido"], f["paso"], round(cal.a_horas(f["inicio"]), 6), round(cal.a_horas(f["fin"]), 6))
                for e, f in d.fallas.items()},
        recuperacion=sorted((str(r["producto"]), r["paso"], r["politica"], r["conserva"], r["reinicia"], r["repite_setup"],
                             (int(r["repetir"]) if str(r["repetir"]).isdigit() else 0), r["verif_min"]) for r in d.recuperacion),
        escenarios={s["id"]: (s["eventos"], s["politica"], s["semilla"]) for s in d.escenarios},
        parametros={k: d.parametros.get(k) for k in ("Aplicar feriados nacionales (1/0)", "Jornada productiva",
                                                       "Máquina de la avería", "Costo equivalente reasignación",
                                                       "Réplicas Monte Carlo de fallas de la plegadora", "Semilla Monte Carlo")},
        habilidades={w: sorted(s) for w, s in d.habilidades.items()}, operarios=d.operarios,
        masa_chapa={k: round(v, 9) for k, v in d.masa_chapa.items()},
        programa_base=[(b["pedido"], b["paso"], b["maquina"], round(cal.a_horas(b["inicio"]), 5), round(cal.a_horas(b["fin"]), 5))
                       for b in d.programa_base],
        jobs=[(j["id"], j["producto"], j["familia"], j["q"], round(j["r"], 9), round(j["d"], 9), round(j["w"], 9),
               round(j["P"], 6), round(j["pdia"], 6), round(j["valor"], 6)) for j in I.jobs],
        ops=[(o["job"], o["paso"], o["principal"], o["trato"], round(o["setup_min"], 9), round(o["min_unidad"], 9), o["lote"],
              o["q"]) for o in I.ops],
        elig_inst={str(o): v for o, v in I.elig.items()},
        p={f"{o}|{m}": round(v, 9) for (o, m), v in I.p.items()},
        machines=I.machines)
    return out, I


def verificar_compatibilidad(inf):
    from sagapi import datos as D
    try:
        dm = D.leer(MOD)
        ok, det = True, f"{len(dm.pedidos)} pedidos, {sum(len(v) for v in dm.rutas.values())} pasos de ruta, " \
                        f"{len(dm.maquinas)} recursos, {len(dm.fallas)} eventos, {len(dm.escenarios)} escenarios"
    except Exception as ex:      # noqa: BLE001
        inf.add("V08", "sagapi.datos.leer + validar sobre el MODELO v2.0", False, f"{type(ex).__name__}: {ex}")
        return
    inf.add("V08", "sagapi.datos.leer + validar (integridad del programa) sobre el MODELO v2.0", ok, det)
    do = D.leer(ORIGINAL)
    ro, Io = _resumen_datos(do)
    rm, Im = _resumen_datos(dm)
    dif = {k: "difiere" for k in ro if ro[k] != rm[k]}
    claves_milp = ["jobs", "ops", "elig_inst", "p", "machines", "calendario", "fallas", "recuperacion", "habilidades",
                   "operarios", "parametros", "programa_base", "masa_chapa"]
    milp_dif = [k for k in claves_milp if k in dif]
    inf.add("V09", "Parámetros del MILP idénticos al original (r_j, d_j, w_j, P_j, q, lote, p_om, M_o, ventanas de falla, "
                   "calendario, recuperación, habilidades, parámetros leídos)", not milp_dif,
            f"diferencias: {milp_dif}" if milp_dif else
            f"{len(Im.jobs)} pedidos, {len(Im.ops)} operaciones, {len(Im.p)} tiempos p_om y {len(dm.fallas)} ventanas idénticos",
            n=len(Im.p))
    otros = {k: v for k, v in dif.items() if k not in claves_milp}
    return otros


def comparar_original(inf, O, S, otros_dif):
    """V10: compara la base nueva con el original y separa cambios intencionados de no intencionados."""
    o = openpyxl.load_workbook(ORIGINAL, data_only=True)
    t = {}
    for ws in o.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if filas:
            enc = [str(c).strip() if c is not None else "" for c in filas[0]]
            t[ws.title] = [dict(zip(enc, f)) for f in filas[1:] if f and any(v is not None for v in f)]
    cuentas = {
        "productos": (len(t["01_PRODUCTOS"]), len(O["02_PRODUCTOS"])), "materiales": (len(t["02_MATERIALES"]), len(O["03_MATERIALES"])),
        "despiece": (len(t["03_DESPIECE"]), len(O["04_DESPIECE"])), "recursos": (len(t["04_MAQUINAS"]), len(O["05_RECURSOS"])),
        "rutas": (len(t["05_RUTAS"]), len(O["06_RUTAS"])), "elegibilidad": (len(t["13_ELEGIBILIDAD"]), len(O["07_ELEGIBILIDAD"])),
        "operarios": (len(t["06_OPERARIOS"]), len(O["08_PERSONAL"])), "calendario": (len(t["07_CALENDARIO"]), len(O["09_CALENDARIO"])),
        "pedidos": (len(t["08_PEDIDOS"]), len(S["01_PEDIDOS_SIM"])), "eventos": (len(t["14_FALLAS_ESC"]), len(S["04_EVENTOS_SIM"])),
        "escenarios": (len(t["18_ESCENARIOS"]), len(S["05_ESCENARIOS_SIM"])),
        "recuperacion": (len(t["15_RECUPERACION"]), len(S["06_RECUPERACION_SIM"])),
        "programa_fifo": (len(t["16_PROGRAMA_BASE"]), len(S["09_PROGRAMA_FIFO_REF"])),
        "habilidades": (25, len(S["03_HABILIDADES_SIM"]))}
    perdidas = {k: v for k, v in cuentas.items() if v[0] != v[1]}
    # valores de despiece: solo cambian (a) costo de herrajes antes vacío y (b) planchas (corrección)
    intencionados = Counter()
    no_intencionados = []
    for a, b in zip(t["03_DESPIECE"], O["04_DESPIECE"]):
        def n(x):
            return x if isinstance(x, (int, float)) and not isinstance(x, bool) else None
        if n(a["Masa kg / producto"]) is not None and not cerca(b["masa_kg_producto"], a["Masa kg / producto"]):
            no_intencionados.append((b["bom_id"], "masa"))
        ca, cb = n(a["Costo S/ / producto"]), (b["costo_soles_producto"] if b["costo_soles_producto"] != "" else None)
        if ca is None and cb is not None:
            intencionados["costo de herraje calculado (antes vacío)"] += 1
        elif ca is not None and not cerca(cb, ca):
            no_intencionados.append((b["bom_id"], "costo", ca, cb))
        pa = n(a["Planchas 1.20×2.40 equivalentes / producto"])
        pb = b["planchas_equiv_producto"] if b["planchas_equiv_producto"] != "" else None
        if (pa or 0) > 0 and not cerca(pb, pa):
            intencionados["planchas corregidas (× piezas por unidad)"] += 1
        for col_o, col_n in (("Ancho pieza mm", "ancho_pieza_mm"), ("Alto pieza mm", "alto_pieza_mm"),
                             ("Longitud pieza m", "longitud_pieza_m")):
            if isinstance(a[col_o], str) and b[col_n] is None:
                intencionados[f"texto de relleno eliminado en {col_o}"] += 1
    for a, b in zip(t["02_MATERIALES"], O["03_MATERIALES"]):
        if (a["Costo unitario S/"] or None) != b["costo_unitario_soles"]:
            no_intencionados.append((b["material_id"], "precio"))
    for a, b in zip(t["04_MAQUINAS"], O["05_RECURSOS"]):
        if (a["Cantidad de escenario"], a["Costo de uso S//h"]) != (b["cantidad"], b["costo_uso_soles_h"]):
            no_intencionados.append((b["recurso_id"], "recurso"))
    for a, b in zip(t["01_PRODUCTOS"], O["02_PRODUCTOS"]):
        if (a["Alto cm"], a["Ancho cm"], a["Fondo cm"], a["Lote guía"]) != (b["alto_cm"], b["ancho_cm"], b["fondo_cm"], b["lote_guia_und"]):
            no_intencionados.append((b["producto_id"], "producto"))
    if otros_dif:
        no_intencionados.append(("lectura Python", list(otros_dif)))
    inf.add("V10", "Comparación con el Excel original (sin pérdida de registros ni cambios no intencionados)",
            not perdidas and not no_intencionados,
            dict(registros_original_vs_nuevo=cuentas, perdidas=perdidas, cambios_intencionados=dict(intencionados),
                 no_intencionados=no_intencionados[:10]))
    return cuentas, dict(intencionados)


def verificar_reproducibilidad(inf):
    """V11: el MODELO se regenera desde OPERATIVA + SIMULACION sin diferencias (fórmulas sin comillas, tol. 1e-9)."""
    sys.path.insert(0, os.path.join(RAIZ, "herramientas"))
    import construir_bd_v2 as C
    tmp = tempfile.mkdtemp(prefix="sagapi_regen_")
    try:
        for a in (C.ARCH_OPER, C.ARCH_SIM):
            shutil.copy2(os.path.join(BD, a), tmp)
        C.construir_modelo(C.leer_fuentes(tmp), os.path.join(tmp, C.ARCH_MOD))
        a = openpyxl.load_workbook(os.path.join(tmp, C.ARCH_MOD))
        b = openpyxl.load_workbook(MOD)
        dif, n = [], 0
        for ws in a.worksheets:
            for fila in ws.iter_rows():
                for c in fila:
                    v1, v2 = c.value, b[ws.title][c.coordinate].value
                    n += 1
                    if isinstance(v1, str) and isinstance(v2, str):
                        v1, v2 = v1.replace("'", ""), v2.replace("'", "")
                    if v1 == v2 or (es_num(v1) and es_num(v2) and cerca(v1, v2, 1e-9)) or \
                            (hasattr(v1, "isoformat") and hasattr(v2, "isoformat") and abs((v1 - v2).total_seconds()) < 1e-3):
                        continue
                    dif.append(f"{ws.title}!{c.coordinate}")
        inf.add("V11", "Reproducibilidad: el MODELO se regenera desde las dos bases fuente sin diferencias",
                not dif and a.sheetnames == b.sheetnames, dif[:10] or f"{n} celdas comparadas; 0 diferencias")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


# ============================================================================ pruebas del programa
def correr(cmd, cwd, log):
    t = time.monotonic()
    r = subprocess.run(cmd, cwd=cwd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8",
                       errors="replace")
    with open(log, "w", encoding="utf-8") as fh:
        fh.write("COMANDO: " + " ".join(cmd) + "\n\n" + r.stdout)
    return r.returncode, round(time.monotonic() - t, 2), r.stdout


def pruebas_python(inf, salida, tl=60):
    """Ejecuta el programa sin modificar con el libro original y con el MODELO v2.0 en copias temporales."""
    nombre = "Base_Datos_SAGAPI_25_TRABAJADORES_SIMULADOS_v1_2.xlsx"
    resultados = {}
    logs = os.path.join(salida, "logs")
    os.makedirs(logs, exist_ok=True)
    for etiqueta, libro in (("original_v1_2", ORIGINAL), ("modelo_v2_0", MOD)):
        tmp = tempfile.mkdtemp(prefix=f"sagapi_{etiqueta}_")
        for item in ("sagapi", "tests", "herramientas", "Optimizacion_MILP_SAGAPI.py"):
            src = os.path.join(RAIZ, item)
            (shutil.copytree if os.path.isdir(src) else shutil.copy2)(src, os.path.join(tmp, item))
        shutil.copy2(libro, os.path.join(tmp, nombre))
        py = [sys.executable, "-X", "utf8"]
        res = {}
        pasos = [("validar", py + ["Optimizacion_MILP_SAGAPI.py", "validar", "--salida", "out_validar"]),
                 ("unitarias", py + ["-m", "unittest", "discover", "-s", "tests", "-p", "test_*.py"]),
                 ("enumeracion_40", py + ["-m", "tests.test_enumeracion", "40"]),
                 ("rodante_S08A", py + ["Optimizacion_MILP_SAGAPI.py", "rodante-personal", "--escenario", "S08A", "--pedidos", "3",
                                        "--ventanas", "5", "--paso-rodante", "8", "--max-operaciones", "36",
                                        "--horizonte-externos", "120", "--tl", str(tl), "--salida", "out_S08A"]),
                 ("rodante_S13A", py + ["Optimizacion_MILP_SAGAPI.py", "rodante-personal", "--escenario", "S13A", "--pedidos", "2",
                                        "--ventanas", "3", "--paso-rodante", "8", "--max-operaciones", "36",
                                        "--horizonte-externos", "120", "--tl", str(tl), "--salida", "out_S13A"]),
                 ("integrar_S13A", py + ["Optimizacion_MILP_SAGAPI.py", "integrar-personal", "--escenario", "S13A",
                                         "--pedidos", "2", "--tl", str(tl), "--salida", "out_integrar"]),
                 ("caso_S06A", py + ["Optimizacion_MILP_SAGAPI.py", "caso", "--escenario", "S06A", "--tl", str(tl),
                                     "--salida", "out_caso"])]
        for paso, cmd in pasos:
            rc, seg, out = correr(cmd, tmp, os.path.join(logs, f"{etiqueta}_{paso}.txt"))
            res[paso] = dict(codigo_salida=rc, segundos=seg, resumen=[x for x in out.splitlines() if x.strip()][-3:])
            print(f"   [{etiqueta}] {paso}: {'OK' if rc == 0 else 'ERROR'} ({seg}s)")
        # salidas comparables
        comp = {}
        for carpeta, archivos in (("out_S08A", ["rodante_S08A_compromisos.csv", "rodante_S08A_decisiones.csv"]),
                                  ("out_S13A", ["rodante_S13A_compromisos.csv", "rodante_S13A_decisiones.csv"]),
                                  ("out_caso", ["programa_S06A.csv", "cambios_S06A.csv"]),
                                  ("out_integrar", [])):
            for a in archivos:
                ruta = os.path.join(tmp, carpeta, a)
                comp[a] = open(ruta, encoding="utf-8-sig").read() if os.path.exists(ruta) else None
        integ = [x for x in os.listdir(os.path.join(tmp, "out_integrar")) if x.endswith(".csv") and "segmentos" not in x] \
            if os.path.isdir(os.path.join(tmp, "out_integrar")) else []
        for a in integ:
            comp[a] = open(os.path.join(tmp, "out_integrar", a), encoding="utf-8-sig").read()
        val = os.path.join(tmp, "out_validar", "validacion_datos.json")
        if os.path.exists(val):
            j = json.load(open(val, encoding="utf-8"))
            j.pop("archivo", None)
            j.pop("highs", None)
            comp["validacion_datos.json"] = json.dumps(j, sort_keys=True, ensure_ascii=False)
        resultados[etiqueta] = dict(pasos=res, salidas=comp)
        shutil.rmtree(tmp, ignore_errors=True)
    o, m = resultados["original_v1_2"], resultados["modelo_v2_0"]
    for paso in o["pasos"]:
        ok = m["pasos"][paso]["codigo_salida"] == 0 and o["pasos"][paso]["codigo_salida"] == 0
        inf.add(f"P-{paso}", f"Programa Python 1.4.0 sin cambios: {paso} (original y v2.0)", ok,
                f"original rc={o['pasos'][paso]['codigo_salida']} ({o['pasos'][paso]['segundos']} s); "
                f"v2.0 rc={m['pasos'][paso]['codigo_salida']} ({m['pasos'][paso]['segundos']} s); "
                f"v2.0: {' | '.join(m['pasos'][paso]['resumen'])[:300]}")
    iguales = {a: o["salidas"].get(a) == m["salidas"].get(a) for a in set(o["salidas"]) | set(m["salidas"])}
    inf.add("P-salidas", "Salidas idénticas original vs v2.0 (compromisos, decisiones, asignaciones, validación)",
            all(iguales.values()), iguales)
    return {k: v["pasos"] for k, v in resultados.items()}


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--pruebas-python", action="store_true")
    ap.add_argument("--tl", type=float, default=60.0)
    ap.add_argument("--salida", default=os.path.join(RAIZ, "entrega_v2_0", "verificacion"))
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    inf = Informe()
    O, S, M = leer_libro(OPER), leer_libro(SIM), leer_libro(MOD)
    print("== Estructura e integridad (V01–V07)")
    verificar_estructura(inf, O, S, M)
    print("== Compatibilidad con Python y parámetros MILP (V08–V09)")
    otros = verificar_compatibilidad(inf) or {}
    print("== Comparación con el original (V10)")
    cuentas, intenc = comparar_original(inf, O, S, otros)
    verificar_reproducibilidad(inf)
    pasos = None
    if a.pruebas_python:
        print("== Ejecución del programa Python (original vs v2.0)")
        pasos = pruebas_python(inf, a.salida, a.tl)
    import scipy
    import numpy
    meta = dict(fecha=time.strftime("%Y-%m-%d %H:%M:%S"), python=sys.version.split()[0], scipy=scipy.__version__,
                numpy=numpy.__version__, openpyxl=openpyxl.__version__,
                libros={os.path.basename(p): os.path.getsize(p) for p in (OPER, SIM, MOD)})
    resumen = dict(meta=meta, pruebas=inf.res, pasos_python=pasos,
                   totales=Counter(r["resultado"] for r in inf.res))
    with open(os.path.join(a.salida, "verificacion_bd_v2.json"), "w", encoding="utf-8") as fh:
        json.dump(resumen, fh, ensure_ascii=False, indent=1, default=str)
    print("TOTALES:", dict(resumen["totales"]))
    return 0 if not any(r["resultado"] == "FALLA" for r in inf.res) else 1


if __name__ == "__main__":
    sys.exit(main())
