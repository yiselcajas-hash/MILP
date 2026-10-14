"""Añade a un libro guardado con openpyxl los valores calculados de sus fórmulas (caché <v>), tomados de una
copia recalculada con LibreOffice, sin alterar ninguna fórmula.

Por qué: openpyxl guarda las fórmulas sin su último valor; Excel las recalcula al abrir, pero otros lectores
(vistas previas, pandas, visores web) verían celdas vacías. LibreOffice sí calcula, pero al guardar reescribe
las referencias a hojas cuyo nombre empieza con un dígito ('08_PEDIDOS'!A1 -> 08_PEDIDOS!A1), forma que Excel
no acepta. Por eso se conserva el XML de openpyxl y solo se le agrega el valor de cada fórmula.

Uso: python inyectar_cache.py libro_openpyxl.xlsx copia_recalculada.xlsx salida.xlsx
"""
import re
import sys
import zipfile

import openpyxl
from lxml import etree

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PNS = "http://schemas.openxmlformats.org/package/2006/relationships"


def hojas_xml(z):
    wb = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.findall(f"{{{PNS}}}Relationship")}
    out = {}
    for s in wb.find(f"{{{NS}}}sheets"):
        t = target[s.get(f"{{{RNS}}}id")]
        t = t.lstrip("/")
        out[s.get("name")] = t if t.startswith("xl/") else "xl/" + t
    return out


def main(src, recalc, dst):
    vals = openpyxl.load_workbook(recalc, data_only=True)
    zin = zipfile.ZipFile(src)
    mapa = hojas_xml(zin)
    inv = {v: k for k, v in mapa.items()}
    zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    n_f = n_v = 0
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in inv and inv[item.filename] in vals.sheetnames:
            ws = vals[inv[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter(f"{{{NS}}}c"):
                f = c.find(f"{{{NS}}}f")
                if f is None:
                    continue
                n_f += 1
                v_old = c.find(f"{{{NS}}}v")
                if v_old is not None:
                    c.remove(v_old)
                val = ws[c.get("r")].value
                if val is None:
                    continue
                v = etree.SubElement(c, f"{{{NS}}}v")
                if isinstance(val, bool):
                    c.set("t", "b"); v.text = "1" if val else "0"
                elif isinstance(val, (int, float)):
                    if "t" in c.attrib:
                        del c.attrib["t"]
                    v.text = repr(float(val)) if isinstance(val, float) else str(val)
                elif hasattr(val, "toordinal"):          # fecha: número de serie de Excel
                    import datetime as dt
                    base = dt.datetime(1899, 12, 30)
                    x = val if isinstance(val, dt.datetime) else dt.datetime.combine(val, dt.time())
                    if "t" in c.attrib:
                        del c.attrib["t"]
                    v.text = repr((x - base).total_seconds() / 86400.0)
                else:
                    s = str(val)
                    if re.match(r"^#(DIV/0!|N/A|NAME\?|NULL!|NUM!|REF!|VALUE!)$", s):
                        c.set("t", "e")
                    else:
                        c.set("t", "str")
                    v.text = s
                n_v += 1
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(item, data)
    zout.close()
    print(f"fórmulas: {n_f}; valores añadidos: {n_v}")


if __name__ == "__main__":
    main(*sys.argv[1:4])
