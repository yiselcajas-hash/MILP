"""Esquemas referenciales (vista frontal a escala) de los productos con pedidos, dibujados a partir
de las dimensiones de 01_PRODUCTOS. No son fotografías: sirven para reconocer el producto y su tamaño."""
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle

G1, G2, G3, G4 = "#3a3a3a", "#7a7a7a", "#bdbdbd", "#e6e6e6"


def _base(ax, W, H):
    ax.set_xlim(-25, W + 25); ax.set_ylim(-30, H + 25); ax.set_aspect("equal"); ax.axis("off")
    ax.annotate("", xy=(0, -15), xytext=(W, -15), arrowprops=dict(arrowstyle="<->", color=G2, lw=0.8))
    ax.text(W / 2, -27, f"{W:.0f} cm", ha="center", va="bottom", fontsize=7, color=G1)
    ax.annotate("", xy=(-12, 0), xytext=(-12, H), arrowprops=dict(arrowstyle="<->", color=G2, lw=0.8))
    ax.text(-22, H / 2, f"{H:.0f} cm", ha="center", va="center", rotation=90, fontsize=7, color=G1)


def estanteria(ax, W, H, n=5, reforzada=False):
    _base(ax, W, H)
    t = 6 if reforzada else 4
    for x in (0, W - t):
        ax.add_patch(Rectangle((x, 0), t, H, fc=G2, ec=G1, lw=0.8))
    for k in range(n):
        y = 8 + k * (H - 14) / (n - 1)
        ax.add_patch(Rectangle((t, y), W - 2 * t, 3 if not reforzada else 4, fc=G3, ec=G1, lw=0.8))
    if reforzada:
        ax.plot([t, W - t], [8, H - 8], color=G2, lw=0.6); ax.plot([t, W - t], [H - 8, 8], color=G2, lw=0.6)


def gondola(ax, W, H, central=True):
    _base(ax, W, H)
    ax.add_patch(Rectangle((0, 0), W, 12, fc=G2, ec=G1, lw=0.8))
    ax.add_patch(Rectangle((W / 2 - 2 if central else 0, 12), 4 if central else W, H - 12, fc=G4, ec=G1, lw=0.6))
    for x in range(8, int(W), 8):
        for y in range(20, int(H), 8):
            ax.plot(x, y, marker=".", ms=1.2, color=G2)
    for k in range(4):
        y = 12 + (k + 1) * (H - 16) / 4.4
        ax.add_patch(Rectangle((2, y), W - 4, 2.5, fc=G3, ec=G1, lw=0.7))


def locker(ax, W, H, puertas):
    _base(ax, W, H)
    ax.add_patch(Rectangle((0, 0), W, H, fc=G4, ec=G1, lw=1))
    cols = 2
    filas = puertas // cols
    w = (W - 6) / cols; h = (H - 12) / filas
    for c in range(cols):
        for f in range(filas):
            x = 3 + c * w; y = 9 + f * h
            ax.add_patch(Rectangle((x + 1, y + 1), w - 2, h - 2, fc=G3, ec=G1, lw=0.8))
            for k in range(3):
                ax.plot([x + w * 0.25, x + w * 0.75], [y + h * 0.75 + k * 2, y + h * 0.75 + k * 2], color=G2, lw=0.5)
            ax.plot(x + w - 5, y + h / 2, marker="s", ms=2, color=G1)


def armario(ax, W, H):
    _base(ax, W, H)
    ax.add_patch(Rectangle((0, 0), W, H, fc=G4, ec=G1, lw=1))
    for x in (2, W / 2):
        ax.add_patch(Rectangle((x, 6), W / 2 - 2, H - 8, fc=G3, ec=G1, lw=0.8))
    ax.plot([W / 2 - 4, W / 2 - 4], [H / 2 - 5, H / 2 + 5], color=G1, lw=1.5)
    ax.plot([W / 2 + 4, W / 2 + 4], [H / 2 - 5, H / 2 + 5], color=G1, lw=1.5)


def dibujar(nombre, tipo, alto, ancho, ruta):
    fig, ax = plt.subplots(figsize=(2.6, 2.6 * min(1.4, max(0.75, alto / ancho))), dpi=110)
    t = tipo.lower()
    if "locker" in t:
        locker(ax, ancho, alto, 6 if "6" in t else 4)
    elif "góndola" in t or "gondola" in t:
        gondola(ax, ancho, alto, central="central" in t)
    elif "gabinete" in t or "armario" in t:
        armario(ax, ancho, alto)
    else:
        estanteria(ax, ancho, alto, n=4 if alto < 120 else 5, reforzada="reforzada" in t)
    ax.set_title(nombre, fontsize=7.5, color=G1)
    fig.tight_layout()
    fig.savefig(ruta, facecolor="white")
    plt.close(fig)


if __name__ == "__main__":
    import openpyxl, os
    wb = openpyxl.load_workbook(sys.argv[1], read_only=True)
    out = sys.argv[2]
    pedidos = {r[1] for r in wb["08_PEDIDOS"].iter_rows(min_row=2, values_only=True) if r[1]}
    for r in wb["01_PRODUCTOS"].iter_rows(min_row=2, values_only=True):
        if r[1] in pedidos:
            dibujar(r[1], r[2], r[3], r[4], os.path.join(out, f"{r[1]}.png"))
    print("ok")
