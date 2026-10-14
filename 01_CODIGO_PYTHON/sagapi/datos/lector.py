"""Lectura estricta del libro Excel según el esquema declarado.

* Lee los valores con ``data_only=True`` y, en una segunda pasada, las fórmulas: una columna de
  ENTRADA que contenga una fórmula es un ERROR (en v1.4 se leía como vacía sin aviso).
* Convierte cada celda al tipo declarado y comprueba obligatoriedad, rangos y dominios.
* Nunca corrige datos: informa cada problema con hoja, fila y columna.
"""
from __future__ import annotations

import datetime as dt
import math

import openpyxl

from .esquema import ESQUEMA, HOJAS
from .problemas import InformeCalidad

TEXTOS_RELLENO = {"no aplica", "según especificación", "sin registro", "todo el lote", "n/a", "na", "-"}


def _convertir(valor, tipo):
    """Devuelve (valor_convertido, mensaje_error)."""
    if valor is None or (isinstance(valor, str) and valor.strip() == ""):
        return None, None
    if tipo == "texto":
        if isinstance(valor, float) and valor.is_integer():
            valor = int(valor)
        return str(valor).strip(), None
    if tipo in ("entero", "binario"):
        if isinstance(valor, bool) or not isinstance(valor, (int, float)):
            return None, f"se esperaba un {tipo} y hay texto {valor!r}"
        if not math.isfinite(valor) or abs(valor - round(valor)) > 1e-9:
            return None, f"se esperaba un {tipo} y hay {valor!r}"
        v = int(round(valor))
        if tipo == "binario" and v not in (0, 1):
            return None, f"binario fuera de {{0,1}}: {v}"
        return v, None
    if tipo in ("decimal", "moneda", "fraccion"):
        if isinstance(valor, bool) or not isinstance(valor, (int, float)) or not math.isfinite(valor):
            return None, f"se esperaba un número y hay {valor!r}"
        return float(valor), None
    if tipo in ("fecha", "fechahora"):
        if isinstance(valor, dt.datetime):
            return valor, None
        if isinstance(valor, dt.date):
            return dt.datetime.combine(valor, dt.time()), None
        return None, f"se esperaba una fecha y hay {valor!r}"
    if tipo == "hora":
        if isinstance(valor, dt.time):
            return valor, None
        if isinstance(valor, dt.datetime):
            return valor.time(), None
        if isinstance(valor, (int, float)) and 0 <= valor < 1:
            seg = round(valor * 86400)
            return dt.time(seg // 3600, (seg % 3600) // 60), None
        return None, f"se esperaba una hora y hay {valor!r}"
    return valor, None


def leer_libro(ruta, informe: InformeCalidad | None = None):
    """Devuelve ({hoja: [fila dict]}, InformeCalidad). Las filas conservan el número de fila Excel en '_fila'."""
    informe = informe or InformeCalidad()
    try:
        wv = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        wf = openpyxl.load_workbook(ruta, data_only=False, read_only=True)
    except Exception as ex:  # noqa: BLE001
        informe.error("Q00", "-", None, None, f"No se pudo abrir el libro: {type(ex).__name__}: {ex}")
        return {}, informe
    tablas = {}
    for h in HOJAS:
        if h.nombre not in wv.sheetnames:
            if h.obligatoria and h.lee_python:
                informe.error("Q00", h.nombre, None, None, "falta la hoja obligatoria")
            tablas[h.nombre] = []
            continue
        if not h.lee_python:
            tablas[h.nombre] = []
            continue
        filas_v = list(wv[h.nombre].iter_rows(values_only=True))
        filas_f = list(wf[h.nombre].iter_rows(values_only=True))
        if not filas_v:
            informe.error("Q00", h.nombre, 1, None, "hoja sin encabezado")
            tablas[h.nombre] = []
            continue
        enc = [str(c).strip() if c is not None else None for c in filas_v[0]]
        pos = {c: j for j, c in enumerate(enc) if c}
        dinamica = h.nombre == "22_HABILIDADES_SIM"
        for c in h.columnas:
            if c.nombre not in pos:
                if c.derivada:
                    informe.advertencia("Q00", h.nombre, 1, c.nombre, "columna calculada ausente")
                elif c.obligatorio:
                    informe.error("Q00", h.nombre, 1, c.nombre, "falta la columna obligatoria")
                else:
                    informe.advertencia("Q00", h.nombre, 1, c.nombre, "falta una columna opcional")
        declaradas = {c.nombre for c in h.columnas}
        extras = [c for c in pos if c not in declaradas]
        if extras and not dinamica:
            informe.advertencia("Q00", h.nombre, 1, ", ".join(extras), "columnas no declaradas en el esquema (se ignoran)")
        meta = {c.nombre: c for c in h.columnas}
        filas = []
        j_nat = pos.get("naturaleza")
        for i, (fv, ff) in enumerate(zip(filas_v[1:], filas_f[1:]), 2):
            if fv is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in fv):
                continue
            fila = {"_fila": i}
            # Una fila declarada CALCULADO (p. ej. un parámetro derivado) puede contener fórmulas: se lee su valor.
            fila_calculada = j_nat is not None and j_nat < len(fv) and fv[j_nat] == "CALCULADO"
            for nombre, j in pos.items():
                v = fv[j] if j < len(fv) else None
                f = ff[j] if j < len(ff) else None
                c = meta.get(nombre)
                if c is None:
                    if dinamica and nombre != "naturaleza":
                        if isinstance(f, str) and f.startswith("="):
                            informe.error("Q01", h.nombre, i, nombre, "celda de entrada con fórmula")
                        conv, msg = _convertir(v, "binario")
                        if msg:
                            informe.error("Q04", h.nombre, i, nombre, msg)
                        fila[nombre] = conv
                    elif nombre == "naturaleza":
                        fila[nombre] = v
                    continue
                if c.derivada:
                    continue
                if isinstance(f, str) and f.startswith("=") and not fila_calculada:
                    informe.error("Q01", h.nombre, i, nombre, "columna de ENTRADA con fórmula; escriba el valor")
                    fila[nombre] = None
                    continue
                if isinstance(v, str) and v.strip().lower() in TEXTOS_RELLENO and c.tipo != "texto":
                    informe.error("Q04", h.nombre, i, nombre, f"texto de relleno {v!r} en campo numérico; deje la celda vacía")
                    fila[nombre] = None
                    continue
                conv, msg = _convertir(v, c.tipo)
                if msg:
                    informe.error("Q04", h.nombre, i, nombre, msg)
                if conv is None and c.obligatorio and not msg:
                    informe.error("Q03", h.nombre, i, nombre, "campo obligatorio vacío")
                if conv is not None and c.tipo in ("entero", "decimal", "moneda", "fraccion"):
                    if c.minimo is not None and conv < c.minimo - 1e-12:
                        informe.error("Q05", h.nombre, i, nombre, f"valor {conv} menor que el mínimo {c.minimo}")
                    if c.maximo is not None and conv > c.maximo + 1e-12:
                        informe.error("Q05", h.nombre, i, nombre, f"valor {conv} mayor que el máximo {c.maximo}")
                if conv is not None and c.dominio and conv not in c.dominio:
                    informe.error("Q05", h.nombre, i, nombre, f"valor {conv!r} fuera del dominio {c.dominio}")
                fila[nombre] = conv
            filas.append(fila)
        tablas[h.nombre] = filas
    wv.close()
    wf.close()
    return tablas, informe


def esquema_hoja(nombre):
    return ESQUEMA[nombre]
