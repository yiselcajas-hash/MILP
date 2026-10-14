"""Libro de resultados (separado de la base de datos): reúne las tablas de 04_RESULTADOS con su procedencia.

No contiene fórmulas: todos los valores provienen de las corridas. Los textos que empiezan con '=' se escriben
como texto para que Excel no los interprete como fórmulas.
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

F_HDR = Font(name="Arial", bold=True, color="FFFFFF", size=10)
P_HDR = PatternFill("solid", fgColor="1F3864")
F_TXT = Font(name="Arial", size=10)

HOJAS = [
    ("E1_anual", "Comparacion_MILP_FIFO/E1_anual.csv", "Año completo por eventos (S00, S20A) por combinación"),
    ("E1_familias", "Comparacion_MILP_FIFO/E1_por_familia.csv", "E1 por familia de producto"),
    ("E2_fallas", "Comparacion_MILP_FIFO/E2_escenarios_falla.csv", "Escenarios de avería: KPI del sistema y Δ vs S00"),
    ("E2b_evento", "Comparacion_MILP_FIFO/E2b_reaccion_evento.csv", "Reacción CONV/EDD/MILP sobre el mismo subproblema"),
    ("E3_operativos", "Escenarios/E3_escenarios_operativos.csv", "Escenarios operativos simulados"),
    ("E4_sensibilidad", "Escenarios/E4_sensibilidad.csv", "Sensibilidad (26_SENSIBILIDAD_SIM)"),
    ("E5_rendimiento", "Resultados_MILP/E5_rendimiento.csv", "Tamaño y tiempo del MILP: v2.0 vs v1.4"),
    ("E0_cuello_botella", "Indicadores_KPI/E0_diagnostico_cuello_botella.csv", "Espera y utilización por unidad (FIFO)"),
    ("E0_carga_teorica", "Indicadores_KPI/E0_carga_teorica_recursos.csv", "Carga de trabajo por recurso"),
    ("E0_cota_tardanza", "Indicadores_KPI/E0_cota_inferior_tardanza.csv", "Tardanza inevitable por pedido"),
    ("KPI_definiciones", "Indicadores_KPI/KPI_definiciones.csv", "Fórmula, unidad y procedencia de cada KPI"),
    ("Pedidos_S20A_MILP", "Indicadores_KPI/pedidos_S20A_MILP_MILP.csv", "Detalle por pedido · S20A · MILP+MILP"),
    ("Pedidos_S20A_FIFO", "Indicadores_KPI/pedidos_S20A_FIFO_CONV.csv", "Detalle por pedido · S20A · FIFO+CONV"),
    ("Util_S20A_MILP", "Indicadores_KPI/utilizacion_S20A_MILP_MILP.csv", "Utilización por unidad · S20A · MILP+MILP"),
    ("Personal_S20A_MILP", "Indicadores_KPI/carga_trabajadores_S20A_MILP_MILP.csv", "Carga por trabajador · S20A · MILP+MILP"),
    ("Eventos_S20A_MILP", "Resultados_MILP/eventos_S20A_MILP_MILP.csv", "Registro de eventos del horizonte rodante"),
    ("Programa_S20A_MILP", "Resultados_MILP/programa_S20A_MILP_MILP.csv", "Programa ejecutado · S20A · MILP+MILP"),
]


def _seguro(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") and not _es_numero(v):
        return "'" + v
    return v


def _es_numero(s):
    try:
        float(s)
        return True
    except ValueError:
        return False


def _hoja(wb, nombre, df, descripcion):
    ws = wb.create_sheet(nombre[:31])
    ws.cell(1, 1, descripcion).font = Font(name="Arial", bold=True, size=11)
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(3, j, str(c))
        cell.font, cell.fill = F_HDR, P_HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = max(10, min(40, len(str(c)) + 2))
    for i, fila in enumerate(df.itertuples(index=False), 4):
        for j, v in enumerate(fila, 1):
            if pd.isna(v) if not isinstance(v, (list, dict)) else False:
                continue
            if hasattr(v, "item"):
                v = v.item()
            ws.cell(i, j, _seguro(v)).font = F_TXT
    ws.freeze_panes = "A4"
    if len(df):
        ws.auto_filter.ref = f"A3:{get_column_letter(len(df.columns))}{3 + len(df)}"
    return ws


def libro_resultados(resultados, ruta, log=print):
    R = Path(resultados)
    wb = Workbook()
    ws = wb.active
    ws.title = "LEEME"
    e0 = R / "Resultados_MILP" / "E0_validacion.json"
    info = json.loads(e0.read_text(encoding="utf-8")) if e0.exists() else {}
    lineas = [
        "SAGAPI v2.0 — Libro de resultados computacionales",
        "",
        "Todos los valores provienen de corridas reales del programa (ver 04_RESULTADOS/registro_experimentos.log).",
        "Pedidos, averías y habilidades del personal son SIMULADOS; tiempos de ruta ESTIMADOS; montos en S/ ESTIMADOS.",
        "Ningún valor es un resultado histórico de SAGAPI ni un ahorro verificado en planta.",
        "Un programa 'factible' aquí lo es respecto del modelo y del verificador independiente; no está validado en planta.",
        "",
        f"Versión del programa: {info.get('version', '?')} · HiGHS: {info.get('highs', '?')} · base: {info.get('base', '?')}",
        f"Regla de objetivo: {info.get('regla_objetivo', '?')} · continuidad: {info.get('regla_continuidad', '?')} · "
        f"β = {info.get('beta', '?')} · horizonte congelado = {info.get('horizonte_congelado_h', '?')} h",
        "",
        "Hojas:",
    ]
    presentes = []
    for nombre, rel, desc in HOJAS:
        p = R / rel
        if not p.exists():
            log(f"Libro de resultados: falta {rel} (hoja {nombre} omitida)")
            continue
        df = pd.read_csv(p)
        _hoja(wb, nombre, df, desc + f"  [origen: 04_RESULTADOS/{rel}]")
        presentes.append(f"  {nombre}: {desc}")
    for i, t in enumerate(lineas + presentes, 1):
        ws.cell(i, 1, _seguro(t)).font = Font(name="Arial", size=10, bold=(i == 1))
    ws.column_dimensions["A"].width = 120
    Path(ruta).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    log(f"Libro de resultados: {ruta} ({len(presentes)} hojas de datos)")
    return ruta
