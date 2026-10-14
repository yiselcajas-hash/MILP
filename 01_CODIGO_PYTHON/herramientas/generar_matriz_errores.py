#!/usr/bin/env python3
"""Genera 06_AUDITORIA_Y_VALIDACION/Matriz_Errores.xlsx: matriz de seguimiento con estado inicial y final.

Fuentes: (1) problemas del código y del modelo (C-xx, N-xx) declarados abajo con su evidencia;
(2) los 84 errores del Excel de la etapa 1 (E-xxx) con la acción aplicada entonces;
(3) los datos pendientes de la empresa (PV-xx). Estados: Pendiente · Parcialmente corregido ·
Corregido y validado · No verificable. Un estado «Corregido y validado» cita la prueba o el archivo que lo respalda.

Uso: python 01_CODIGO_PYTHON/herramientas/generar_matriz_errores.py [--etapa1 RUTA] [--salida RUTA]
"""
from __future__ import annotations

import argparse
import csv
import os
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CV, PC, NV, PE = "Corregido y validado", "Parcialmente corregido", "No verificable", "Pendiente"

# id, problema, archivo (v1.4 → v2.0), prioridad, estado inicial, estado final, acción realizada, evidencia
CODIGO = [
    ("C-01", "R11 no restringía la programación (sin variables S ni x): cobertura de personal aparente",
     "modelo.py → modelo/milp.py, personal.py", "P0", PE, CV,
     "R11 eliminada del MILP. Certificado de cobertura por flujo máximo (condición suficiente) + MILP de asignación "
     "de trabajadores con cliques de solape + reparación laboral cuando el certificado falla.",
     "tests/test_personal.py (8 pruebas, incl. un trabajador para dos máquinas obliga a secuenciar)"),
    ("C-04", "Curado POR_LOTE partido entre jornadas (físicamente imposible)", "modelo.py, reglas.py, verificacion.py",
     "P0", PE, CV, "R12 (k_o entero, tramo L = 4 h o 8 h) en el MILP, en todas las reglas y en el verificador.",
     "tests/test_restricciones.py::test_R12_*; tests/test_enumeracion.py (25 de 40 casos con R12)"),
    ("C-16", "Lectura silenciosa de fórmulas en columnas de entrada", "datos.py → datos/lector.py", "P0", PE, CV,
     "Lectura doble (valores + fórmulas); fórmula en columna de ENTRADA = error Q01 que bloquea.",
     "tests/test_datos_calidad.py::test_defecto_detectado_y_bloquea[Q01]"),
    ("C-17", "Esquema implícito y frágil", "datos.py → datos/esquema.py, calidad.py", "P0", PE, CV,
     "Esquema declarado (tipos, rangos, dominios, claves); controles Q00–Q17; ERROR detiene con mensaje y código 2.",
     "tests/test_datos_calidad.py (10 pruebas); Programa_principal.py validar"),
    ("C-19", "Los reportes sobrescribían el Excel de entrada", "reportes.py → reportes/excel.py", "P0", PE, CV,
     "Libro de resultados separado (04_RESULTADOS/Resultados_SAGAPI_v2.xlsx); la base nunca se escribe.",
     "tests/test_cli_reportes.py::test_base_no_se_modifica"),
    ("C-02", "Personal integrado solo en un piloto (2 % / 0,17 % de operaciones)", "personal_integrado.py, rodante_personal.py",
     "P1", PE, CV, "Asignación de los 25 trabajadores a TODAS las operaciones del año (y tramos parciales) con verificación "
     "L1–L5. Validado con habilidades SIMULADAS: no sustituye la matriz real (PV-02).",
     "04_RESULTADOS/Indicadores_KPI/carga_trabajadores_*.csv; tests/test_personal.py::test_asignacion_milp_anual_fifo"),
    ("C-03", "Capacidad unitaria de todo recurso", "modelo.py, instancia.py, verificacion.py", "P1", PE, CV,
     "Un recurso con cantidad k se expande en k unidades idénticas (m#1..m#k). La base actual solo tiene k = 1 "
     "(PIN-EXT k = 0): la capacidad k > 1 se valida con microinstancias.",
     "tests/test_restricciones.py::test_R4_R5_*; tests/test_enumeracion.py (33 de 40 casos con k = 2)"),
    ("C-05", "La penalidad fija como objetivo primario sacrificaba la tardanza", "modelo.py, optimizacion.py", "P1", PE, CV,
     "Jerarquía configurable (28_CONFIG_MODELO): por defecto P1 Σw·T, P2 Σ P·U, P3 estabilidad (solo al reprogramar). "
     "La alternativa PENALIDAD_FIJA se conserva y se compara (E4).",
     "tests/test_restricciones.py::test_R7_R8_*; 04_RESULTADOS/Escenarios/E4_sensibilidad.csv (regla_objetivo)"),
    ("C-07", "Big-M de R6 inválido cuando a_w < p_om", "modelo.py:113 → modelo/milp.py", "P1", PE, CV,
     "M = UB + p_om − a_w; prueba de validez de la fila en la esquina S = UB.",
     "tests/test_restricciones.py::test_R6_ventana_futura_y_bigM_valido; enumeración con ventanas futuras (13 casos)"),
    ("C-08", "Miopía de evento_unico (solo pedidos del subproblema)", "experimentos.py", "P1", PE, CV,
     "E2 evalúa cada escenario sobre el año completo (KPI del sistema y Δ frente a S00); E2b conserva la comparación "
     "sobre el mismo subproblema.", "04_RESULTADOS/Comparacion_MILP_FIFO/E2_escenarios_falla.csv"),
    ("C-11", "Comparador FIFO débil; efectos de secuenciación y reprogramación mezclados", "experimentos.py", "P1", PE, CV,
     "FIFO+CONV, FIFO+MILP, EDD+EDD, MILP+CONV, MILP+MILP con los mismos datos, restricciones y verificador.",
     "04_RESULTADOS/Comparacion_MILP_FIFO/E1_anual.csv"),
    ("C-12", "Selección del piloto por índice", "Optimizacion_MILP_SAGAPI.py", "P1", PE, CV,
     "El horizonte rodante incluye en cada evento a TODOS los pedidos activos.",
     "tests/test_rodante.py::test_anio_con_averias_sin_duplicados"),
    ("C-18", "Feriados leídos de dos formas", "datos.py, experimentos.py", "P1", PE, CV,
     "Una sola lectura (09_CALENDARIO + aplicar_feriados); sensibilidad 0/1 en E4.",
     "04_RESULTADOS/Escenarios/E4_sensibilidad.csv (aplicar_feriados)"),
    ("C-09", "Horas de liberación y vencimiento fijas en el código", "instancia.py", "P2", PE, CV,
     "Se leen de 11_PARAMETROS (hora_liberacion, hora_vencimiento).", "sagapi/datos/carga.py"),
    ("C-10", "MILP monolítico anual inviable (~54 000 binarias)", "arquitectura", "P2", PE, CV,
     "Descomposición por eventos (horizonte rodante) con tamaño medido por evento; todos los eventos MILP de E1 "
     "con óptimo demostrado.", "04_RESULTADOS/Resultados_MILP/eventos_S20A_MILP_MILP.csv; E5_rendimiento.csv"),
    ("C-13", "Wilcoxon sobre escenarios dependientes", "reportes.py", "P2", PE, CV,
     "Retirado; se reportan diferencias por escenario sin inferencia estadística.", "Registro_Cambios.md"),
    ("C-15", "La sensibilidad ignoraba el Excel", "experimentos.py", "P2", PE, CV,
     "E4 recorre los factores y niveles de 26_SENSIBILIDAD_SIM.", "04_RESULTADOS/Escenarios/E4_sensibilidad.csv"),
    ("C-21", "Tres implementaciones del lexicográfico", "optimizacion.py, rodante_personal.py, principal", "P2", PE, CV,
     "Un único procedimiento (optimizacion.optimizar) con estado global y respaldo heurístico rotulado.",
     "tests/test_solver_estados.py; tests/test_enumeracion.py"),
    ("C-22", "Programa principal monolítico", "Optimizacion_MILP_SAGAPI.py → Programa_principal.py + sagapi/cli.py",
     "P2", PE, CV, "Capas: datos, instancia, modelo, solver, optimización, simulación, verificación, personal, KPI, "
     "reportes, CLI; registro (logging) y códigos de salida.", "tests/test_cli_reportes.py"),
    ("C-23", "Dependencias sin fijar", "requirements.txt", "P2", PE, CV, "Versiones exactas; pruebas en entorno limpio.",
     "06_AUDITORIA_Y_VALIDACION/Reporte_Pruebas/"),
    ("C-24", "Pruebas sin cobertura de riesgos", "tests/", "P2", PE, CV,
     "Enumeración exhaustiva, una prueba por restricción, calidad de datos, personal, horizonte rodante, regresión "
     "v1.4, estados del solucionador, CLI.", "06_AUDITORIA_Y_VALIDACION/Reporte_Pruebas/"),
    ("C-25", "Indicadores incompletos", "indicadores.py", "P2", PE, CV,
     "KPI con fórmula, unidad y procedencia; cota inferior de tardanza.", "04_RESULTADOS/Indicadores_KPI/KPI_definiciones.csv"),
    ("C-06", "Precedencia lineal sin lotes de transferencia", "datos y modelo", "P2", NV, NV,
     "Sin cambio: requiere datos de flujo real (qué operaciones pueden solaparse).", "PV (flujo de planta)"),
    ("C-14", "Monte Carlo circular sobre averías sintéticas", "experimentos.py", "P3", NV, PC,
     "Análisis retirado (no informaba sobre la planta). La robustez frente a averías reales sigue sin poder evaluarse.",
     "PV-07 (historial de paradas)"),
    ("C-20", "El generador de datos no reproducía la base", "herramientas/", "P2", CV, CV,
     "Etapa 1 (regeneración sin diferencias); etapa 2: construir_base_final.py reproduce la base final.",
     "Evidencias_Validacion/etapa1_auditoria (verificación de la etapa 1)"),
    # hallazgos de la etapa 2
    ("N-01", "El verificador se detenía con KeyError ante una unidad no elegible en lugar de informar la violación",
     "verificacion.py", "P1", "Detectado en etapa 2", CV, "Se informa 'R1: unidad no elegible' y se detiene la evaluación.",
     "tests/test_restricciones.py::test_R1_R2_asignacion_y_alternativa"),
    ("N-02", "Personal insuficiente producía un error interno poco claro", "simulacion.py, personal.py", "P1",
     "Detectado en etapa 2", CV, "Comprobación previa: operaciones sin personal habilitado → ErrorSimulacion con mensaje.",
     "tests/test_personal.py::test_personal_insuficiente_detiene_con_mensaje_claro"),
    ("N-03", "El mantenimiento planificado se trataba como avería imprevista", "simulacion.py", "P2", "Detectado en etapa 2",
     CV, "Las ventanas de MANTENIMIENTO_PLANIFICADO se conocen desde el inicio; todos los métodos las respetan.",
     "tests/test_rodante.py::test_mantenimiento_planificado_es_anunciado"),
    ("N-04", "Área de chapa mezclaba fórmulas y valores en una columna de entrada", "base de datos (04_DESPIECE)", "P1",
     "Detectado en etapa 2", CV, "Se escribe el valor y un control Q08 verifica área = ancho × alto / 10⁶.",
     "Programa_principal.py validar (0 errores)"),
    ("N-05", "MILP+CONV S20A en v2.0 con configuración v1.4: 22 atrasados frente a 23 de v1.4", "simulacion.py", "P3",
     "Detectado en etapa 2", NV, "Todos los eventos son óptimos en ambas versiones; la diferencia surge de soluciones "
     "óptimas alternativas en eventos previos que cambian el estado posterior (no se puede atribuir a un error).",
     "06_AUDITORIA_Y_VALIDACION/Evidencias_Validacion/regresion_v14.md"),
]

MAPA_E = {"CORREGIDO": CV, "ELIMINADO": CV, "REEMPLAZADO": CV, "RECONSTRUIDO": CV, "REUBICADO": CV, "SEPARADO": CV,
          "CONSERVADO": CV, "PENDIENTE": NV, "DOCUMENTADO": NV}


def estado_e(accion):
    a = accion.upper()
    if "PENDIENTE" in a and any(k in a for k in ("CORREGIDO", "SEPARADO")):
        return PC
    if "PENDIENTE" in a or a.startswith("DOCUMENTADO"):
        return NV
    for k, v in MAPA_E.items():
        if a.startswith(k):
            return v
    return PE


def escribir(ws, enc, filas, anchos):
    for j, (c, w) in enumerate(zip(enc, anchos), 1):
        cell = ws.cell(1, j, c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, f in enumerate(filas, 2):
        for j, v in enumerate(f, 1):
            v = ("'" + v) if isinstance(v, str) and v[:1] in "=+-@" else v
            c = ws.cell(i, j, v)
            c.font = Font(name="Arial", size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(enc))}{len(filas) + 1}"


def main(argv=None):
    ap = argparse.ArgumentParser()
    cands = [os.path.join(RAIZ, "06_AUDITORIA_Y_VALIDACION", "Evidencias_Validacion", "etapa1_auditoria", "auditoria"),
             os.path.join(RAIZ, "historial", "etapa1_auditoria_v2_0", "auditoria")]
    ap.add_argument("--etapa1", default=next((c for c in cands if os.path.isdir(c)), cands[-1]))
    ap.add_argument("--salida", default=os.path.join(RAIZ, "06_AUDITORIA_Y_VALIDACION", "Matriz_Errores.xlsx"))
    a = ap.parse_args(argv)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    e_rows = list(csv.DictReader(open(os.path.join(a.etapa1, "matriz_errores_excel.csv"), encoding="utf-8-sig")))
    pv = list(csv.DictReader(open(os.path.join(a.etapa1, "pendientes_validacion.csv"), encoding="utf-8-sig")))
    seg = []
    for (i, prob, arch, pri, ini, fin, acc, ev) in CODIGO:
        seg.append((i, "Código / modelo", prob, arch, pri, ini, fin, acc, ev))
    for r in e_rows:
        seg.append((r["id"], "Base de datos (Excel)", r["problema"], f"{r['hoja_original']} {r['celda_o_rango']}",
                    {"ALTA": "P0", "MEDIA": "P1", "BAJA": "P2"}.get(r["severidad"].upper(), "P2"),
                    "Corregido en etapa 1" if estado_e(r["accion"]) == CV else "Pendiente (etapa 1)",
                    estado_e(r["accion"]), f"[{r['accion']}] {r['correccion_aplicada_o_motivo']}",
                    f"{r['ubicacion_v2']}; carga estricta de la base final sin errores (validar)"))
    for r in pv:
        seg.append((r["id"], "Dato de la empresa", r["dato_requerido"], r["ubicacion_v2"], r["prioridad"], NV, NV,
                    f"Solicitar a: {r['solicitar_a']}. {r['por_que_no_se_reconstruye']}", r["impacto_en_el_modelo"]))
    enc = ["ID", "Tipo", "Problema identificado", "Archivo / ubicación", "Prioridad", "Estado inicial", "Estado final",
           "Acción realizada", "Evidencia / prueba"]
    escribir(wb.create_sheet("Matriz_Seguimiento"), enc, seg, [8, 16, 60, 32, 9, 18, 20, 70, 50])
    cnt = Counter((s[1], s[6]) for s in seg)
    tipos = ["Código / modelo", "Base de datos (Excel)", "Dato de la empresa"]
    estados = [CV, PC, NV, PE]
    filas = [[t] + [cnt.get((t, e), 0) for e in estados] + [sum(cnt.get((t, e), 0) for e in estados)] for t in tipos]
    filas.append(["Total"] + [sum(cnt.get((t, e), 0) for t in tipos) for e in estados] + [len(seg)])
    ws.cell(1, 1, "Matriz de seguimiento SAGAPI v2.0 — estado final").font = Font(name="Arial", bold=True, size=12)
    ws.cell(2, 1, "Estados: Corregido y validado (con prueba o archivo que lo respalda) · Parcialmente corregido · "
                  "No verificable (requiere datos de la empresa) · Pendiente.").font = Font(name="Arial", size=9)
    for j, c in enumerate(["Tipo"] + estados + ["Total"], 1):
        cell = ws.cell(4, j, c)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        ws.column_dimensions[get_column_letter(j)].width = 24
    for i, f in enumerate(filas, 5):
        for j, v in enumerate(f, 1):
            ws.cell(i, j, v).font = Font(name="Arial", size=10, bold=(f[0] == "Total"))
    wb.save(a.salida)
    print(f"{a.salida}: {len(seg)} filas;", {e: sum(cnt.get((t, e), 0) for t in tipos) for e in estados})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
