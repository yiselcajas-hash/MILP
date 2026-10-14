#!/usr/bin/env python3
"""Recalcula las fórmulas de los libros v2.0 con LibreOffice (sin interfaz) y normaliza las referencias.

1. LibreOffice ejecuta calculateAll() y guarda: así los valores quedan en caché y pandas,
   openpyxl(data_only=True) o cualquier visor los leen sin abrir Excel.
2. LibreOffice escribe las referencias a hojas cuyo nombre empieza con un dígito sin comillas
   (03_DESPIECE!A2). Excel exige '03_DESPIECE'!A2, por lo que se restituyen las comillas en las
   fórmulas, validaciones y nombres definidos del XML, sin tocar los valores calculados.
3. Informa cualquier error de fórmula (#REF!, #NAME?, #DIV/0!, #VALUE!, #N/A, #NUM!, #NULL!).

Si LibreOffice no está instalado basta con abrir y guardar el libro en Excel (recalcula al abrir).
Uso: python herramientas/recalcular_bd_v2.py entrega_v2_0/base_datos/*.xlsx
"""
from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalcularYGuardar()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
ERRORES = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def recalcular_libreoffice(ruta, timeout=300):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return "LibreOffice no está instalado: abrir y guardar el libro en Excel para recalcular."
    env = dict(os.environ, SAL_USE_VCLPLUGIN="svp")
    with tempfile.TemporaryDirectory(prefix="sagapi_lo_") as perfil:
        url = Path(perfil).as_uri()
        subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                       capture_output=True, timeout=timeout, env=env)
        destino = Path(perfil) / "user" / "basic" / "Standard"
        destino.mkdir(parents=True, exist_ok=True)
        (destino / "Module1.xba").write_text(MACRO, encoding="utf-8")
        antes = os.stat(ruta).st_mtime_ns
        r = subprocess.run([soffice, "--headless", "--norestore", f"-env:UserInstallation={url}",
                            "vnd.sun.star.script:Standard.Module1.RecalcularYGuardar?language=Basic&location=application",
                            str(Path(ruta).absolute())], capture_output=True, text=True, timeout=timeout, env=env)
        if r.returncode != 0 or os.stat(ruta).st_mtime_ns == antes:
            return f"LibreOffice no recalculó el libro (código {r.returncode}): {r.stderr.strip()[:300]}"
    return None


def restituir_comillas(ruta):
    """Pone entre comillas simples las referencias a hojas que lo necesitan, dentro del XML."""
    with zipfile.ZipFile(ruta) as z:
        partes = {n: z.read(n) for n in z.namelist()}
        info = {n: z.getinfo(n) for n in z.namelist()}
    libro = partes["xl/workbook.xml"].decode("utf-8")
    hojas = re.findall(r'<sheet [^>]*name="([^"]+)"', libro)
    necesitan = sorted([h for h in hojas if re.match(r"^\d", h) or re.search(r"[^\w.]", h)], key=len, reverse=True)
    if not necesitan:
        return 0
    patron = re.compile(r"(?<![\w'\]])(" + "|".join(re.escape(h) for h in necesitan) + r")(?=!)")
    etiqueta = re.compile(r"(<(?:\w+:)?(?:f|formula1|formula2|formula|definedName)\b[^>]*>)([^<]*)(</)")
    cambios = 0

    def arreglar(m):
        nonlocal cambios
        nuevo, k = patron.subn(lambda x: f"'{x.group(1)}'", m.group(2))
        cambios += k
        return m.group(1) + nuevo + m.group(3)

    for nombre in list(partes):
        if nombre == "xl/workbook.xml" or (nombre.startswith("xl/worksheets/") and nombre.endswith(".xml")):
            partes[nombre] = etiqueta.sub(arreglar, partes[nombre].decode("utf-8")).encode("utf-8")
    tmp = str(ruta) + ".tmp"
    with zipfile.ZipFile(tmp, "w") as z:
        for nombre, datos in partes.items():
            z.writestr(info[nombre], datos, compress_type=zipfile.ZIP_DEFLATED)
    os.replace(tmp, ruta)
    return cambios


def revisar(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    errores = [f"{ws.title}!{c.coordinate}" for ws in wb.worksheets for fila in ws.iter_rows() for c in fila
               if isinstance(c.value, str) and c.value in ERRORES]
    sin_valor = 0
    wf = openpyxl.load_workbook(ruta, data_only=False)
    formulas = 0
    for ws in wf.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    if wb[ws.title][c.coordinate].value is None:
                        sin_valor += 1
    return dict(formulas=formulas, errores=len(errores), muestra_errores=errores[:20], formulas_sin_valor_en_cache=sin_valor)


def main(rutas):
    salida = {}
    for ruta in rutas:
        aviso = recalcular_libreoffice(ruta)
        n = restituir_comillas(ruta)
        salida[os.path.basename(ruta)] = dict(recalculo=aviso or "LibreOffice: recalculado", referencias_corregidas=n,
                                              **revisar(ruta))
    print(json.dumps(salida, indent=1, ensure_ascii=False))
    return 0 if all(v["errores"] == 0 for v in salida.values()) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
