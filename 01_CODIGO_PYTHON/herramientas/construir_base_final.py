#!/usr/bin/env python3
"""Construye SAGAPI_BASE_CORREGIDA_FINAL.xlsx y Diccionario_de_datos.xlsx.

Fuente: la base corregida de la etapa 1 (02_BASE_DATOS/version_anterior/SAGAPI_BD_OPERATIVA_v2_0.xlsx y
SAGAPI_BD_SIMULACION_v2_0.xlsx). Se integran en UN libro con el esquema declarado en
`sagapi/datos/esquema.py`. Los datos no cambian; las adiciones (documentadas en el registro de cambios) son:
  * 06_RUTAS: trabajadores_requeridos (=1, el mismo supuesto de v1.4) e interrumpible_en_pausas (0 para POR_LOTE);
  * 11_PARAMETROS: lote_no_cruza, horizonte_congelado_h, fuente_habilidades, fuente_pedidos;
  * 28_CONFIG_MODELO: regla_objetivo (TARDANZA_PONDERADA), limite_asignacion_personal_s; máquina crítica por clave;
  * 26_SENSIBILIDAD_SIM: claves de factor que el programa lee;
  * 27_ESCENARIOS_OPERATIVOS_SIM: escenarios de carga, plazos, personal e indisponibilidad (SIMULADOS).
Se eliminan las hojas auxiliares 99_LISTAS y 09_PROGRAMA_FIFO_REF (el FIFO lo calcula el programa).

Uso (desde la raíz del proyecto):
  python 01_CODIGO_PYTHON/herramientas/construir_base_final.py
  python 01_CODIGO_PYTHON/herramientas/recalcular_excel.py 02_BASE_DATOS/SAGAPI_BASE_CORREGIDA_FINAL.xlsx
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as GL

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
from sagapi.datos.esquema import ESQUEMA, HOJAS, NATURALEZA_DEF, NATURALEZAS  # noqa: E402
from estilo_excel import (F_FX, F_HDR, F_IN, FMT, P_HDR, nombre_definido, validacion_decimal,  # noqa: E402
                          validacion_entero, validacion_lista)

RAIZ = os.path.dirname(os.path.dirname(AQUI))
ORIGEN = os.path.join(RAIZ, "02_BASE_DATOS", "version_anterior")
VERSION_BASE = "3.0.0"
TIPO_FMT = {"texto": "txt", "entero": "int", "binario": "int", "decimal": "gen", "moneda": "soles", "fraccion": "pct",
            "fecha": "fecha", "fechahora": "fechahora", "hora": "hora"}
TAB = {"OPERATIVA": "2E7D32", "REGISTRO": "1565C0", "SIMULACION": "E65100", "DERIVADA": "595959", "INDICE": "7F7F7F"}
LISTA_NAT = '"' + ",".join(NATURALEZAS) + '"'


def leer(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=False)
    out = {}
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if filas:
            enc = list(filas[0])
            out[ws.title] = [dict(zip(enc, f)) for f in filas[1:] if f and any(v is not None for v in f)]
            out["_enc_" + ws.title] = enc
    return out


def lit(v):
    return None if isinstance(v, str) and v.startswith("=") else v


# ----------------------------------------------------------------------------------------- escritura
def escribir(wb, nombre, filas, columnas=None, comentarios=True):
    """filas: lista de dict {columna: valor | '=fórmula con {r}'}; columnas: orden (por defecto el esquema)."""
    h = ESQUEMA.get(nombre)
    cols = columnas or [c.nombre for c in h.columnas]
    ws = wb.create_sheet(nombre)
    meta = {c.nombre: c for c in h.columnas} if h else {}
    for j, cn in enumerate(cols, 1):
        cell = ws.cell(1, j, cn)
        cell.font, cell.fill = F_HDR, P_HDR
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[GL(j)].width = max(10, min(42, len(cn) + 3))
        c = meta.get(cn)
        if comentarios and c is not None:
            partes = [c.descripcion] if c.descripcion else []
            partes += [f"Tipo: {c.tipo}"] + ([f"Unidad: {c.unidad}"] if c.unidad else [])
            if c.naturaleza:
                partes.append(f"Naturaleza: {c.naturaleza}")
            if c.minimo is not None or c.maximo is not None:
                partes.append(f"Rango: [{c.minimo if c.minimo is not None else '-∞'}, {c.maximo if c.maximo is not None else '∞'}]")
            partes.append("Calculada (fórmula)" if c.derivada else ("Obligatorio" if c.obligatorio else "Opcional (vacío válido)"))
            cm = Comment("\n".join(partes), "SAGAPI BD")
            cm.width, cm.height = 280, 120
            cell.comment = cm
    ws.row_dimensions[1].height = 32
    for i, fila in enumerate(filas, 2):
        for j, cn in enumerate(cols, 1):
            v = fila.get(cn)
            if v is None or v == "":
                continue
            if isinstance(v, str) and v.startswith("="):
                v = v.replace("{r}", str(i))
            cell = ws.cell(i, j, v)
            cell.font = F_FX if isinstance(v, str) and v.startswith("=") else F_IN
            c = meta.get(cn)
            tipo = c.tipo if c else ("binario" if isinstance(v, int) else "texto")
            if TIPO_FMT.get(tipo, "gen") != "gen":
                cell.number_format = FMT[TIPO_FMT[tipo]]
    if not filas:
        for j, cn in enumerate(cols, 1):
            c = meta.get(cn)
            if c and TIPO_FMT.get(c.tipo, "gen") != "gen":
                for i in range(2, 202):
                    ws.cell(i, j).number_format = FMT[TIPO_FMT[c.tipo]]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{GL(len(cols))}{max(1, len(filas) + 1)}"
    if h:
        ws.sheet_properties.tabColor = TAB[h.grupo]
        n = max(len(filas) + 1, 201 if not filas else len(filas) + 1)
        for j, cn in enumerate(cols, 1):
            c = meta.get(cn)
            if c is None or c.derivada:
                continue
            rng = f"{GL(j)}2:{GL(j)}{n}"
            if c.dominio:
                validacion_lista(ws, '"' + ",".join(c.dominio) + '"', rng)
            elif c.tipo in ("entero", "binario"):
                validacion_entero(ws, rng, int(c.minimo if c.minimo is not None else 0),
                                  1 if c.tipo == "binario" else (int(c.maximo) if c.maximo is not None else None))
            elif c.tipo in ("decimal", "moneda", "fraccion") and c.minimo is not None:
                validacion_decimal(ws, rng, c.minimo)
    return ws


def col_letra(nombre, columna, columnas=None):
    cols = columnas or [c.nombre for c in ESQUEMA[nombre].columnas]
    return GL(cols.index(columna) + 1)


def rango(nombre, columna, n, columnas=None):
    L = col_letra(nombre, columna, columnas)
    return f"'{nombre}'!${L}$2:${L}${n + 1}"


# ----------------------------------------------------------------------------------------- construcción
def construir(salida_base, salida_dicc):
    O = leer(os.path.join(ORIGEN, "SAGAPI_BD_OPERATIVA_v2_0.xlsx"))
    S = leer(os.path.join(ORIGEN, "SAGAPI_BD_SIMULACION_v2_0.xlsx"))
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    n = {}

    # ---------- 11_PARAMETROS (primero: nombres definidos)
    pars = [dict(r) for r in O["11_PARAMETROS"] if r["parametro_id"] != "categoria_empresarial"]
    for r in pars:
        if r["parametro_id"] == "hora_liberacion":
            r["nota"] = "Hora en que el pedido liberado puede iniciar (lo lee el programa)."
            r["usado_por_python"] = 1
        if r["parametro_id"] == "hora_vencimiento":
            r["nota"] = "Hora límite del día comprometido; si no es productivo, cierre del último día productivo previo."
            r["usado_por_python"] = 1
        if r["parametro_id"] in ("trabajadores_produccion", "costo_hora_hombre", "anio_base", "horas_productivas_dia"):
            r["usado_por_python"] = 1
    pars += [
        dict(parametro_id="lote_no_cruza", parametro="Operaciones no interrumpibles no cruzan", valor_txt="TRAMO",
             unidad="TRAMO | JORNADA", naturaleza="ESTIMADO", usado_por_python=1,
             nota="TRAMO: no cruzan refrigerio ni cierre (ciclo de horno continuo). JORNADA: solo no cruzan el cierre."),
        dict(parametro_id="horizonte_congelado_h", parametro="Horizonte congelado al reprogramar", valor_num=0,
             unidad="h productivas", naturaleza="CONFIGURACION", usado_por_python=1,
             nota="Operaciones que empiezan antes de t0 + H no se reprograman (0 = solo se fijan las iniciadas)."),
        dict(parametro_id="fuente_habilidades", parametro="Fuente de la matriz de habilidades", valor_txt="SIMULADA",
             unidad="SIMULADA | REAL", naturaleza="CONFIGURACION", usado_por_python=1,
             nota="REAL usa 16_REG_HABILIDADES (nivel ≥ 2 y opera solo = 1) cuando la empresa la registre."),
        dict(parametro_id="fuente_pedidos", parametro="Fuente de pedidos", valor_txt="SIMULADA", unidad="SIMULADA | REAL",
             naturaleza="CONFIGURACION", usado_por_python=1, nota="REAL usa 12_REG_ENTREGAS cuando existan registros."),
    ]
    fv = [p["parametro_id"] for p in pars].index("ventas_anuales") + 2
    fu = [p["parametro_id"] for p in pars].index("uit_2025") + 2
    pars.append(dict(parametro_id="categoria_empresarial", parametro="Categoría empresarial (Ley 30056)",
                     valor_txt=f'=IF(C{fv}/C{fu}<=150,"Microempresa",IF(C{fv}/C{fu}<=1700,"Pequeña empresa",'
                               f'IF(C{fv}/C{fu}<=2300,"Mediana empresa","Gran empresa")))',
                     unidad="-", naturaleza="CALCULADO", nota="Ventas anuales / UIT.", usado_por_python=0))
    escribir(wb, "11_PARAMETROS", pars)
    for i, p in enumerate(pars, 2):
        if p.get("valor_num") is not None:
            nombre_definido(wb, f"p_{p['parametro_id']}", "11_PARAMETROS", f"C{i}")
    n["par"] = len(pars)

    # ---------- 01..10 operativos
    fam = [dict(r, precio_promedio_soles="=D{r}/C{r}") for r in O["01_FAMILIAS"]]
    escribir(wb, "01_FAMILIAS", fam)
    nb, nr, nprod = len(O["04_DESPIECE"]), len(O["06_RUTAS"]), len(O["02_PRODUCTOS"])
    prods = []
    for r in O["02_PRODUCTOS"]:
        prods.append({k: lit(v) for k, v in r.items()} | dict(
            n_componentes=f"=COUNTIF({rango('04_DESPIECE', 'producto_id', nb)},A{{r}})",
            n_pasos_ruta=f"=COUNTIF({rango('06_RUTAS', 'producto_id', nr)},A{{r}})",
            dimensiones_genericas=f"=IF(COUNTIFS($E$2:$E${nprod + 1},E{{r}},$F$2:$F${nprod + 1},F{{r}},"
                                  f"$G$2:$G${nprod + 1},G{{r}})>=5,1,0)"))
    ws = escribir(wb, "02_PRODUCTOS", prods)
    validacion_lista(ws, f"='01_FAMILIAS'!$A$2:$A${len(fam) + 1}", f"B2:B{nprod + 1}")
    mats = [{k: lit(v) for k, v in r.items()} for r in O["03_MATERIALES"]]
    escribir(wb, "03_MATERIALES", mats)
    nm = len(mats)
    MID, PRE, ESP = (rango("03_MATERIALES", c, nm) for c in ("material_id", "costo_unitario_soles", "espesor_mm"))
    bom = []
    for r in O["04_DESPIECE"]:
        d = {k: lit(v) for k, v in r.items()}
        area = r["area_chapa_m2_pieza"]
        if isinstance(area, str) and area.startswith("="):
            # Área por pieza = ancho × alto: se escribe como DATO (la columna es de entrada); el control Q08
            # comprueba la coherencia con las medidas.
            d["area_chapa_m2_pieza"] = round(r["ancho_pieza_mm"] * r["alto_pieza_mm"] / 1e6, 12)
        precio = f"INDEX({PRE},MATCH(D{{r}},{MID},0))"
        d["masa_kg_producto"] = (f'=IF(E{{r}}="CHAPA_AREA",F{{r}}*J{{r}}*INDEX({ESP},MATCH(D{{r}},{MID},0))/1000*p_densidad_acero,'
                                 f'IF(E{{r}}="PERFIL_LONGITUD",F{{r}}*I{{r}}*p_kg_por_m_perfil,""))')
        d["costo_soles_producto"] = (f'=IF(E{{r}}="CHAPA_AREA",K{{r}}*{precio}/p_aprovechamiento_chapa,IF(E{{r}}="PERFIL_LONGITUD",'
                                     f'K{{r}}*{precio}/p_aprovechamiento_perfil,IF(E{{r}}="UNIDAD",F{{r}}*{precio},"")))')
        d["planchas_equiv_producto"] = '=IF(N(J{r})>0,F{r}*J{r}/p_area_plancha_m2,"")'
        bom.append(d)
    ws = escribir(wb, "04_DESPIECE", bom)
    validacion_lista(ws, f"='02_PRODUCTOS'!$A$2:$A${nprod + 1}", f"B2:B{nb + 1}")
    validacion_lista(ws, f"='03_MATERIALES'!$A$2:$A${nm + 1}", f"D2:D{nb + 1}")
    recs = [{k: lit(v) for k, v in r.items()} for r in O["05_RECURSOS"]]
    escribir(wb, "05_RECURSOS", recs)
    nrec = len(recs)
    rutas = []
    for r in O["06_RUTAS"]:
        d = {k: lit(v) for k, v in r.items()}
        d["trabajadores_requeridos"] = 1
        d["interrumpible_en_pausas"] = 0 if r["tratamiento_tiempo"] == "POR_LOTE" else 1
        rutas.append(d)
    ws = escribir(wb, "06_RUTAS", rutas)
    validacion_lista(ws, f"='02_PRODUCTOS'!$A$2:$A${nprod + 1}", f"A2:A{nr + 1}")
    validacion_lista(ws, f"='05_RECURSOS'!$A$2:$A${nrec + 1}", f"D2:D{nr + 1}")
    ele = [{k: lit(v) for k, v in r.items()} for r in O["07_ELEGIBILIDAD"]]
    ws = escribir(wb, "07_ELEGIBILIDAD", ele)
    validacion_lista(ws, f"='05_RECURSOS'!$A$2:$A${nrec + 1}", f"C2:C{len(ele) + 1}")
    per = [{k: lit(v) for k, v in r.items()} | dict(control_horario_h="=(D{r}-C{r})*24-E{r}") for r in O["08_PERSONAL"]]
    escribir(wb, "08_PERSONAL", per)
    cal = [{k: lit(v) for k, v in r.items()} | dict(dia_semana="=WEEKDAY(A{r},2)",
                                                  productivo_efectivo='=IF(AND(p_aplicar_feriados=1,D{r}<>""),0,C{r})',
                                                  horas_productivas="=E{r}*p_horas_productivas_dia") for r in O["09_CALENDARIO"]]
    escribir(wb, "09_CALENDARIO", cal)
    ncal = len(cal)
    CF, CE = rango("09_CALENDARIO", "fecha", ncal), rango("09_CALENDARIO", "productivo_efectivo", ncal)
    capm = []
    for rr in recs:
        for mes in range(1, 13):
            capm.append(dict(recurso_id=rr["recurso_id"], mes=mes,
                             dias_productivos=f"=SUMPRODUCT((MONTH({CF})=B{{r}})*(YEAR({CF})=p_anio_base)*{CE})",
                             horas_dia="=p_horas_productivas_dia",
                             cantidad=f"=INDEX({rango('05_RECURSOS', 'cantidad', nrec)},MATCH(A{{r}},"
                                      f"{rango('05_RECURSOS', 'recurso_id', nrec)},0))",
                             horas_disponibles="=C{r}*D{r}*E{r}", naturaleza="CALCULADO"))
    escribir(wb, "10_CAPACIDAD_MENSUAL", capm)
    for nombre in ("12_REG_ENTREGAS", "13_REG_PARADAS", "14_REG_CALIDAD", "15_REG_TIEMPOS", "16_REG_HABILIDADES"):
        ws = escribir(wb, nombre, [])
        cols = [c.nombre for c in ESQUEMA[nombre].columnas]
        for cn, hoja, k in (("producto_id", "02_PRODUCTOS", nprod), ("recurso_id", "05_RECURSOS", nrec),
                            ("operario_id", "08_PERSONAL", len(per))):
            if cn in cols:
                L = GL(cols.index(cn) + 1)
                validacion_lista(ws, f"='{hoja}'!$A$2:$A${k + 1}", f"{L}2:{L}201")

    # ---------- simulación
    peds = [{k: lit(v) for k, v in r.items()} | dict(plazo_dias="=G{r}-E{r}", liberacion_fin_semana="=IF(WEEKDAY(F{r},2)>5,1,0)",
                                                   compromiso_fin_semana="=IF(WEEKDAY(G{r},2)>5,1,0)")
            for r in S["01_PEDIDOS_SIM"]]
    ws = escribir(wb, "20_PEDIDOS_SIM", peds)
    validacion_lista(ws, f"='02_PRODUCTOS'!$A$2:$A${nprod + 1}", f"B2:B{len(peds) + 1}")
    escribir(wb, "21_PERFILES_SIM", [dict(r) for r in S["02_PERFILES_SIM"]])
    enc = [c for c in S["_enc_03_HABILIDADES_SIM"] if c]
    ws = escribir(wb, "22_HABILIDADES_SIM", [dict(r) for r in S["03_HABILIDADES_SIM"]], columnas=enc)
    validacion_entero(ws, f"B2:{GL(len(enc) - 1)}{len(S['03_HABILIDADES_SIM']) + 1}", 0, 1, permitir_vacio=False)
    evs = [{k: lit(v) for k, v in r.items()} | dict(fin="=D{r}+E{r}/24") for r in S["04_EVENTOS_SIM"]]
    escribir(wb, "23_EVENTOS_SIM", evs)
    escribir(wb, "24_ESCENARIOS_FALLA_SIM", [dict(r, politica_recuperacion=r["politica_recuperacion"])
                                             for r in S["05_ESCENARIOS_SIM"]])
    escribir(wb, "25_RECUPERACION_SIM", [dict(r) for r in S["06_RECUPERACION_SIM"]])
    sens = [
        ("duracion_falla_mult", "Multiplicador de la duración de las averías", "0.5;1;2;4", "S20A (F01–F12, política A)",
         "Robustez ante averías más cortas o más largas"),
        ("regla_objetivo", "Regla del objetivo primario", "TARDANZA_PONDERADA;PENALIDAD_FIJA", "S00 y S20A",
         "Compensación entre tardanza y número de pedidos atrasados"),
        ("beta_reasignacion", "Peso de reasignación β (h equivalentes)", "0;8;24", "S20A",
         "Compensación entre tardanza y estabilidad del programa"),
        ("horizonte_congelado_h", "Horizonte congelado al reprogramar (h)", "0;8", "S20A", "Efecto de congelar el turno en curso"),
        ("factor_alternativa", "Factor de tiempo en máquina alternativa", "1.0;1.1;1.25", "S15A y S16A",
         "Valor de la flexibilidad de asignación"),
        ("aplicar_feriados", "Aplicar feriados nacionales", "0;1", "S00", "Calendario sin y con feriados"),
        ("lote_no_cruza", "Regla de continuidad de operaciones no interrumpibles", "TRAMO;JORNADA", "S00",
         "Efecto de permitir que el horno opere durante el refrigerio"),
    ]
    escribir(wb, "26_SENSIBILIDAD_SIM", [dict(factor_id=a, factor=b, niveles=c, aplicado_a=d, proposito=e,
                                               naturaleza="CONFIGURACION") for a, b, c, d, e in sens])
    ausentes20 = ";".join(f"OP-{k:02d}" for k in range(21, 26))
    ausentes15 = ";".join(f"OP-{k:02d}" for k in range(16, 26))
    ope = [
        ("E-BASE", "Carga, plazos y personal de la base (96 pedidos simulados)", 1.0, 1.0, None, None, None),
        ("E-CARGA-BAJA", "Carga baja: cantidades × 0,75", 0.75, 1.0, None, None, None),
        ("E-CARGA-ALTA", "Carga alta: cantidades × 1,25", 1.25, 1.0, None, None, None),
        ("E-CARGA-MUY-ALTA", "Carga muy alta: cantidades × 1,5 (prueba de cuello de botella)", 1.5, 1.0, None, None, None),
        ("E-PLAZO-AJUSTADO", "Plazos de entrega × 0,8", 1.0, 0.8, None, None, None),
        ("E-PLAZO-HOLGADO", "Plazos de entrega × 1,25", 1.0, 1.25, None, None, None),
        ("E-PERSONAL-20", "5 trabajadores ausentes todo el año (OP-21 a OP-25)", 1.0, 1.0, ausentes20, None, None),
        ("E-PERSONAL-15", "10 trabajadores ausentes todo el año (OP-16 a OP-25)", 1.0, 1.0, ausentes15, None, None),
        ("E-MANT-PLE", "Mantenimiento planificado de la plegadora (F21)", 1.0, 1.0, None, "F21", "A"),
        ("E-FALLAS-PLE", "12 fallas de la plegadora 2025 (F01–F12)", 1.0, 1.0, None,
         "+".join(f"F{k:02d}" for k in range(1, 13)), "A"),
        ("E-FALLA-LARGA", "Falla prolongada de la plegadora (F13, 29 h de reloj)", 1.0, 1.0, None, "F13", "A"),
        ("E-ALTA+FALLAS", "Carga alta con las 12 fallas de la plegadora", 1.25, 1.0, None,
         "+".join(f"F{k:02d}" for k in range(1, 13)), "A"),
    ]
    escribir(wb, "27_ESCENARIOS_OPERATIVOS_SIM", [dict(escenario_id=a, descripcion=b, factor_cantidad=c, factor_plazo=d,
                                                       operarios_ausentes=e, eventos=f, politica_recuperacion=g,
                                                       naturaleza="SIMULADO") for a, b, c, d, e, f, g in ope])
    conf = []
    for r in S["08_CONFIG_MODELO"]:
        d = dict(r)
        pid = d["parametro_id"]
        if pid in ("replicas_mc", "semilla_mc", "niveles_epsilon", "rho_epsilon", "penalidad_diaria_formula"):
            continue
        if pid == "regla_penalidad":
            d = dict(parametro_id="regla_objetivo", parametro="Regla del objetivo primario", valor_txt="TARDANZA_PONDERADA",
                     unidad="TARDANZA_PONDERADA | PENALIDAD_FIJA | PENALIDAD_DIARIA_CON_TOPE", naturaleza="CONFIGURACION",
                     usado_por_python=1, nota="P1 = Σ w_j·T_j (h); P2 = penalidad estimada; P3 = estabilidad. Ver modelo matemático.")
        if pid == "maquina_averia":
            d = dict(parametro_id="maquina_critica", parametro="Recurso crítico de estudio", valor_txt="PLE-01", unidad="recurso_id",
                     naturaleza="CONFIGURACION", usado_por_python=1, nota="Plegadora hidráulica de chapa metálica.")
        if pid in ("tolerancia_lexicografica", "limite_computo_s", "brecha_relativa", "beta_reasignacion", "tolerancia_verificacion"):
            d["usado_por_python"] = 1
            d["nota"] = {"limite_computo_s": "Por fase lexicográfica; --tl lo reemplaza.",
                         "tolerancia_verificacion": "Holgura numérica del verificador independiente."}.get(pid, d.get("nota"))
        conf.append(d)
    conf.append(dict(parametro_id="limite_asignacion_personal_s", parametro="Límite de cómputo de la asignación de personal",
                     valor_num=120, unidad="s", naturaleza="CONFIGURACION", usado_por_python=1,
                     nota="MILP de balance de carga sobre el programa final."))
    escribir(wb, "28_CONFIG_MODELO", conf)

    # ---------- índice
    ws = wb.create_sheet("00_INDICE", 0)
    filas = [dict(hoja=h.nombre, grupo=h.grupo, naturaleza=h.naturaleza or "ver columna naturaleza",
                  lee_python=1 if h.lee_python else 0, descripcion=h.descripcion) for h in HOJAS]
    filas += [dict(hoja=""), dict(hoja="VERSION_BASE", grupo=VERSION_BASE, descripcion="SAGAPI_BASE_CORREGIDA_FINAL"),
              dict(hoja="COLORES", descripcion="Azul = dato introducido; negro = fórmula. Pestañas: verde operativa, azul "
                                               "registros reales, naranja simulación, gris derivadas.")]
    filas += [dict(hoja=k, grupo="NATURALEZA", descripcion=v) for k, v in NATURALEZA_DEF.items()]
    cols = ["hoja", "grupo", "naturaleza", "lee_python", "descripcion"]
    for j, cn in enumerate(cols, 1):
        c = ws.cell(1, j, cn)
        c.font, c.fill = F_HDR, P_HDR
        ws.column_dimensions[GL(j)].width = [30, 16, 24, 11, 100][j - 1]
    for i, f in enumerate(filas, 2):
        for j, cn in enumerate(cols, 1):
            if f.get(cn) not in (None, ""):
                ws.cell(i, j, f[cn]).font = F_IN
    ws.sheet_properties.tabColor = TAB["INDICE"]
    ws.freeze_panes = "A2"
    orden = ["00_INDICE"] + [h.nombre for h in HOJAS]
    wb._sheets.sort(key=lambda w: orden.index(w.title))
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = f"SAGAPI_BASE_CORREGIDA_FINAL v{VERSION_BASE}"
    wb.save(salida_base)

    # ---------- diccionario de datos (separado)
    dic = openpyxl.Workbook()
    ws = dic.active
    ws.title = "DICCIONARIO"
    cols = ["hoja", "grupo", "columna", "tipo", "unidad", "obligatorio", "calculada", "naturaleza", "minimo", "maximo",
            "dominio", "descripcion", "uso_en_python"]
    for j, cn in enumerate(cols, 1):
        c = ws.cell(1, j, cn)
        c.font, c.fill = F_HDR, P_HDR
        ws.column_dimensions[GL(j)].width = [26, 12, 28, 10, 14, 11, 10, 15, 8, 8, 30, 70, 40][j - 1]
    i = 2
    for h in HOJAS:
        cs = list(h.columnas)
        if h.nombre == "22_HABILIDADES_SIM":
            cs += [type(cs[0])("<recurso_id>", "binario", descripcion="Una columna por recurso: 1 = habilitado (SIMULADO).",
                               uso_python="habilidades"), type(cs[0])("naturaleza", descripcion="SIMULADO")]
        for c in cs:
            vals = [h.nombre, h.grupo, c.nombre, c.tipo, c.unidad, "Sí" if c.obligatorio and not c.derivada else "No",
                    "Sí" if c.derivada else "No", c.naturaleza, c.minimo, c.maximo, " | ".join(c.dominio),
                    c.descripcion, c.uso_python]
            for j, v in enumerate(vals, 1):
                if v not in (None, ""):
                    ws.cell(i, j, v).font = F_IN
            i += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{GL(len(cols))}{i - 1}"
    w2 = dic.create_sheet("NATURALEZA")
    for j, cn in enumerate(["codigo", "definicion"], 1):
        c = w2.cell(1, j, cn)
        c.font, c.fill = F_HDR, P_HDR
    w2.column_dimensions["A"].width, w2.column_dimensions["B"].width = 18, 110
    for k, (a, b) in enumerate(NATURALEZA_DEF.items(), 2):
        w2.cell(k, 1, a).font = F_IN
        w2.cell(k, 2, b).font = F_IN
    dic.save(salida_dicc)
    return salida_base, salida_dicc


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--salida", default=os.path.join(RAIZ, "02_BASE_DATOS"))
    a = ap.parse_args()
    os.makedirs(a.salida, exist_ok=True)
    print(construir(os.path.join(a.salida, "SAGAPI_BASE_CORREGIDA_FINAL.xlsx"),
                    os.path.join(a.salida, "Diccionario_de_datos.xlsx")))
